"""مراجعةُ كلِّ تقارير لوحة المصنع — قراءةٌ محضة، صفرُ كتابة، صفرُ إنفاق.

طلبُ المالك (2026-09-19): «راجع كلَّ تقرير في لوحة المصنع» بعد أن خرجت كلُّ
دراسةٍ بلافتة «لم يجتز فحصَ الجودة النهائي». الأداةُ تُجيب ذلك بأثرٍ لا بحدس:
لكلِّ دراسةٍ لها `analysis_id` تبني العرضَ وتُشغّل بوّابةَ التسليم **مرّتين**
— بالنصّ كما هو، ثم بعد تطبيعِ العناوين (الدرس ٢٦١) — فيظهر بالضبط أيُّ
حجبٍ كان سببُه شكلَ سطرِ العنوان وأيُّه عيبُ محتوىً حقيقيّ.

لا تُعدَّل قاعدةٌ ولا يُستدعى نموذجٌ ولا مزوّدٌ مدفوع: `build_view` و
`silk_export_gate.evaluate` حسابٌ محضٌ فوق المخزَّن. القواعدُ تُفتَح
`mode=ro` (نمط `audit_unit_gaps.py` / `export_captured.py`).

التشغيل على Railway (قشرة الخدمة):
    python3 tools/audit_study_reports.py
    python3 tools/audit_study_reports.py --json          # مخرَج مهيكل
    python3 tools/audit_study_reports.py --study 6       # دراسةٌ بعينها
    python3 tools/audit_study_reports.py --headings      # عيّنةُ أسطر العناوين
    python3 tools/audit_study_reports.py --client        # مستندُ العميل الفعليّ + سجلّ بلاغات المالك
    python3 tools/audit_study_reports.py --study 11 --dump   # نصُّ مستند العميل كاملاً
    python3 tools/audit_study_reports.py --blob-file a.json --client  # نتيجةٌ نُزِّلت من GET /analyses/{id}
    curl -H "X-API-Key: …" "$HOST/ops/audit-studies?study=11&dump=1"  # من خارج الخادم

Read-only review of every factory study report: rebuilds each view and runs
the delivery gate twice (raw vs. heading-canonicalized) so a block caused by
heading shape is told apart from a real content defect. No writes, no paid
calls.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sqlite3
import sys

_TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.dirname(_TOOLS_DIR)
for p in (_REPO_ROOT, _TOOLS_DIR):
    if p not in sys.path:
        sys.path.insert(0, p)

_HEADING_RE = re.compile(r"^##\s+(\d+)\.\s*(.+?)\s*$", re.M)


def _platform_db(cli: "str | None") -> str:
    if cli:
        return cli
    if os.environ.get("SILK_PLATFORM_DB"):
        return os.environ["SILK_PLATFORM_DB"]
    return os.path.join(os.environ.get("SILK_DATA_DIR", "data"), "platform.db")


def _engine_db(cli: "str | None") -> str:
    if cli:
        return cli
    if os.environ.get("SILK_DB"):
        return os.environ["SILK_DB"]
    return os.path.join(os.environ.get("SILK_DATA_DIR", "data"), "silk.db")


def _ro(path: str) -> sqlite3.Connection:
    con = sqlite3.connect(f"file:{os.path.abspath(path)}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    return con


def _columns(con: sqlite3.Connection, table: str) -> set:
    return {r["name"] for r in con.execute(f"PRAGMA table_info({table})")}


def _studies(path: str, only: "int | None") -> list:
    con = _ro(path)
    try:
        cols = _columns(con, "studies")
        wanted = [c for c in ("id", "analysis_id", "product", "market_pref",
                              "state", "report_language", "run_error")
                  if c in cols]
        sql = f"SELECT {', '.join(wanted)} FROM studies"
        if only is not None:
            sql += " WHERE id = ?"
            rows = con.execute(sql + " ORDER BY id", (only,)).fetchall()
        else:
            rows = con.execute(sql + " ORDER BY id").fetchall()
    finally:
        con.close()
    return [dict(r) for r in rows]


def _analysis_blob(path: str, analysis_id: int) -> "dict | None":
    con = _ro(path)
    try:
        row = con.execute("SELECT json_blob FROM analyses WHERE id = ?",
                          (analysis_id,)).fetchone()
    finally:
        con.close()
    if not row:
        return None
    try:
        return json.loads(row["json_blob"])
    except (TypeError, ValueError):
        return None


def _gate(view: dict) -> dict:
    """قرارُ التسليم كما تقرؤه نقاطُ المنصّة نفسُها — لا منطقَ ثانٍ هنا.

    مراجعة §58 #5: `evaluate` لا يُعيد مفتاح `blocked_checks` إطلاقاً
    (مفاتيحُه: blocking/digest/fail_drivers/findings/gate/verdict) — فكانت
    القائمةُ تُطبَع فارغةً دائماً وتُقرَأ «لا فحصَ حاجب».
    """
    import silk_export_gate as eg
    decision = eg.evaluate(view) or {}
    return {"verdict": decision.get("verdict"),
            "fail_drivers": list(decision.get("fail_drivers") or []),
            "blocking": [str(c) for c in (decision.get("blocking") or [])],
            "blocked": bool(eg.is_blocked(decision))}


def _sections_seen(text: str) -> int:
    return len(set(_HEADING_RE.findall(text or "")))


def _heading_lines(text: str, limit: int = 12) -> list:
    out = []
    for ln in str(text or "").split("\n"):
        s = ln.strip()
        if s.startswith("#") or s.startswith("**") or re.match(
                r"^[0-9٠-٩]{1,2}[ \t]*[.)\-–—:]", s):
            out.append(s[:80])
        if len(out) >= limit:
            break
    return out


def factory_view(blob: dict, lang: str, analysis_id: "int | None" = None) -> dict:
    """عرضُ المصنع نفسُه + النثرُ البديل المخزَّن — بلا أيّ نداءٍ مدفوع."""
    from silk_platform.engine_bridge import factory_client_view
    # النثرُ البديل المخزَّن (`client_fallback_prose`) يرافق النتيجةَ أصلاً؛ لا
    # تحضيرَ هنا (بلا ذكاءٍ اصطناعيّ يكون التحضيرُ بلا أثر).
    return factory_client_view(blob, lang)


def client_text(view: dict) -> tuple:
    """(نصُّ مستند العميل كما يُبنى الآن، ملاحظاتُ بوّابته النهائية) — الدرس ٢٩٦.

    «كل مرة تظهر نفس المشاكل»: الإصلاحُ كان يُثبَت على عيّناتٍ اصطناعية فقط؛
    هذا يعيد بناءَ **مستند العميل الفعليّ** لكلّ دراسةٍ مخزَّنة بالشيفرة الحالية
    ويشغّل عليه البوّابةَ النهائية وسجلَّ بلاغات المالك — قبل أن يراه المالك.
    """
    import tempfile
    import silk_quality_gate as Q
    from silk_reports import render_client_docx
    from docx import Document
    findings: list = []
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "client.docx")
        try:
            render_client_docx(view, out)
        except RuntimeError as e:          # البوّابةُ النهائية رفضت البناء
            findings.append({"check": "artifact_refused", "note": str(e)[:300]})
            return "", findings
        doc = Document(out)
        parts = [p.text for p in doc.paragraphs]
        parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
        text = "\n".join(x for x in parts if x.strip())
    lang = view.get("report_language") or "ar"
    findings += Q.run_client_artifact_text_gate(text, lang, view)
    return text, findings


def review(platform_db: str, engine_db: str, only: "int | None",
           show_headings: bool, client: bool = False,
           dump: bool = False, raw_compare: bool = True) -> list:
    """`raw_compare=False` داخل الخادم (مراجعة §58): مقارنةُ النصّ الخام تبدّل
    دالّةً عامّة مؤقّتاً — آمنةٌ في سطر أوامر، خطرةٌ مع طلباتٍ متزامنة."""
    import silk_ai_judge as judge
    import silk_render

    report = []
    for st in _studies(platform_db, only):
        row = {"study": st.get("id"), "product": st.get("product"),
               "market": st.get("market_pref"), "state": st.get("state")}
        aid = st.get("analysis_id")
        if not aid:
            row["status"] = "لا تقرير مخزَّن (لم تُشغَّل أو فشلت)"
            row["run_error"] = st.get("run_error")
            report.append(row)
            continue
        blob = _analysis_blob(engine_db, int(aid))
        if blob is None:
            row["status"] = "تعذّرت قراءة نتيجة التحليل"
            report.append(row)
            continue
        lang = st.get("report_language") or "ar"
        raw_text = (((blob.get("deep_research") or {}).get("report")
                     or {}).get("report") or "")
        row["chars"] = len(raw_text)
        row["sections_raw"] = _sections_seen(raw_text)
        row["sections_after"] = _sections_seen(
            judge.canonicalize_section_headings(raw_text, lang))
        if show_headings:
            row["heading_lines"] = _heading_lines(raw_text)
        try:
            # (أ) البوّابةُ على العرض **كما يراه المصنعُ** (الدرس ٢٩٧: نفسُ
            #     تهيئة سطح المصنع، والنثرُ البديل المخزَّن بلا نداء نموذج).
            fview = factory_view(blob, lang, int(aid))
            row["gate"] = _gate(fview)
            # (ب) وعلى النصّ **الخام** بتعطيل التطبيع مؤقّتاً — فيظهر بالضبط
            #     أيُّ حجبٍ كان سببُه شكلَ سطرِ العنوان وأيُّه عيبُ محتوىً.
            if raw_compare:
                _orig = judge.canonicalize_section_headings
                try:
                    judge.canonicalize_section_headings = lambda t, l="ar": t
                    row["gate_raw"] = _gate(silk_render.build_view(blob, lang))
                finally:
                    judge.canonicalize_section_headings = _orig
        except Exception as e:  # noqa: BLE001 — تدقيقٌ لا يسقط على دراسة
            row["gate_error"] = f"{type(e).__name__}: {e}"
        if client:
            try:
                import silk_quality_gate as _Q
                text, cf = client_text(fview if "fview" in locals()
                                       else factory_view(blob, lang, int(aid)))
                row["client_findings"] = [f"{f.get('check')}: {f.get('note')}"
                                          for f in cf]
                # ما يمنع التسليم فعلاً (مراجعة §58) — الملاحظاتُ الاستشارية لا تُعَدّ حجباً.
                row["client_blockers"] = [
                    f"{f.get('check')}: {f.get('note')}"
                    for f in (_Q.delivery_blockers(cf)
                              + [x for x in cf if x.get("check") == "artifact_refused"])]
                if dump:
                    row["client_text"] = text
            except Exception as e:  # noqa: BLE001
                row["client_error"] = f"{type(e).__name__}: {e}"
        row["status"] = "ok"
        report.append(row)
    return report


def _print(rows: list) -> None:
    if not rows:
        print("لا دراسات في هذه القاعدة.")
        return
    print(f"{'#':>4}  {'أقسام خام':>9}  {'بعد التطبيع':>11}  "
          f"{'محجوب':>6}  الحكم / السبب")
    print("-" * 78)
    healed = blocked = 0
    for r in rows:
        if r.get("status") != "ok":
            print(f"{r.get('study'):>4}  {'—':>9}  {'—':>11}  {'—':>6}  "
                  f"{r.get('status')}")
            continue
        g = r.get("gate") or {}
        if g.get("blocked"):
            blocked += 1
        if r.get("sections_after", 0) > r.get("sections_raw", 0):
            healed += 1
        drivers = "، ".join(g.get("fail_drivers") or []) or "—"
        raw = r.get("gate_raw") or {}
        if raw.get("blocked") and not g.get("blocked"):
            drivers += "  ← كان محجوباً بشكل العنوان وحدَه"
        print(f"{r.get('study'):>4}  {r.get('sections_raw', 0):>9}  "
              f"{r.get('sections_after', 0):>11}  "
              f"{('نعم' if g.get('blocked') else 'لا'):>6}  "
              f"{g.get('verdict')} / {drivers}")
        for ln in r.get("heading_lines") or []:
            print(f"        | {ln}")
        for ln in r.get("client_findings") or []:
            print(f"        ✗ {ln[:160]}")
        if r.get("client_error"):
            print(f"        ! {r['client_error']}")
        if r.get("client_text"):
            print(r["client_text"])
    print("-" * 78)
    print(f"دراسات: {len(rows)} · محجوبة التنزيل: {blocked} · "
          f"شفاها تطبيعُ العناوين: {healed}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--platform-db")
    ap.add_argument("--db", help="قاعدة المحرّك (analyses)")
    ap.add_argument("--study", type=int)
    ap.add_argument("--headings", action="store_true")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--client", action="store_true",
                    help="ابنِ مستندَ العميل الفعليّ وشغّل عليه البوّابة النهائية "
                         "وسجلَّ بلاغات المالك (الدرس ٢٩٦)")
    ap.add_argument("--dump", action="store_true",
                    help="مع --client: اطبع نصَّ مستند العميل كاملاً")
    ap.add_argument("--blob-file",
                    help="نتيجةُ تحليلٍ مخزَّنة بصيغة JSON (من GET /analyses/{id}) — "
                         "إعادةُ التشغيل بلا قواعد ولا railway ssh (الدرس ٢٩٧)")
    args = ap.parse_args()
    if args.blob_file:
        with open(args.blob_file, encoding="utf-8") as fh:
            blob = json.load(fh)
        lang = blob.get("report_language") or "ar"
        view = factory_view(blob, lang, blob.get("analysis_id"))
        text, cf = client_text(view)
        out = {"gate": _gate(view), "client_findings": [
            f"{f.get('check')}: {f.get('note')}" for f in cf]}
        if args.dump:
            out["client_text"] = text
        print(json.dumps(out, ensure_ascii=False, indent=2))
        return 0
    pdb, edb = _platform_db(args.platform_db), _engine_db(args.db)
    for p in (pdb, edb):
        if not os.path.exists(p):
            print(f"قاعدة غير موجودة: {p}", file=sys.stderr)
            return 2
    rows = review(pdb, edb, args.study, args.headings,
                  client=args.client or args.dump, dump=args.dump)
    if args.json:
        print(json.dumps(rows, ensure_ascii=False, indent=2))
    else:
        _print(rows)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
