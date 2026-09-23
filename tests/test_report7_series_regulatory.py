"""تقرير سِلك ٧ — السلسلة الزمنية والمتطلبات التنظيمية (§3.2، §4.3، الدرس ٢٧٩).

a year that could not be retrieved is a declared row in the client imports
table (never estimated or filled); every requirement row carries its type —
legally mandatory / mandatory by category or destination / optional — from
the curated reference, rendered in the client report; the regulatory-fit note
says what the score measures, not that the facility holds approvals.

كلُّ الحالات مصطنعةٌ لاختبارٍ هرمتيّ.
"""
import csv
import pathlib

from docx import Document

import silk_reports as SR
import silk_requirements_agent as RA

_ROOT = pathlib.Path(__file__).resolve().parents[1]


def test_a_missing_year_is_a_declared_row_in_the_imports_table():
    doc = Document()
    dr = {"imports": {"value_line": "واردات", "head": "الواردات",
                      "series": [{"year": 2021, "value": 51.0e6},
                                 {"year": 2024, "value": 89.4e6}],
                      "years_missing": [2022, 2023]}}
    SR._client_imports_section(doc, dr, "ar")
    t = doc.tables[-1]
    years = [r.cells[0].text for r in t.rows[1:]]
    assert years == ["2021", "2022", "2023", "2024"]
    assert "غير متاح" in t.rows[2].cells[1].text


def test_every_reference_row_declares_its_requirement_type():
    rows = list(csv.DictReader(open(_ROOT / "data/requirements_l1.csv",
                                    encoding="utf-8")))
    assert rows and all(r["requirement_status"] in RA.REQUIREMENT_STATUSES
                        for r in rows)
    # البندُ الذي يقول نصُّه «اختياري داعم لا إلزامي» مصنَّفٌ كذلك.
    vol = [r for r in rows if "اختياري داعم" in r["note"]]
    assert vol and all(r["requirement_status"] == "voluntary" for r in vol)


def test_requirement_rows_are_deterministic_and_carry_the_type():
    rows = RA.requirement_rows("NLD", "080410")
    assert rows and all(r["status"] for r in rows)
    assert {r["direction"] for r in rows} == {"entry", "exit"}


def test_the_client_report_splits_requirements_by_type():
    doc = Document()
    dr = {"requirements": [
        {"item": "بطاقة غذائية", "authority": "GSO", "direction": "entry",
         "source_url": "https://www.gso.org.sa", "status": "legal_mandatory"},
        {"item": "شهادة صحية", "authority": "NVWA", "direction": "entry",
         "source_url": "https://www.nvwa.nl", "status": "conditional_legal",
         "conditional_on_gate": True},
        {"item": "التسجيل في بوابة الصادرات", "authority": "هيئة الصادرات",
         "direction": "exit",
         "source_url": "https://saudiexports.sa", "status": "voluntary"}]}
    SR._client_requirements_table(doc, dr, "ar")
    t = doc.tables[-1]
    types = [r.cells[3].text for r in t.rows[1:]]
    assert types[0] == "إلزامي قانوناً" and types[2] == "اختياري داعم"
    assert types[1].endswith(SR._T("req_after_gate", "ar"))
    dirs = [r.cells[1].text for r in t.rows[1:]]
    assert dirs == [SR._T("req_dir_entry", "ar")] * 2 + [
        SR._T("req_dir_exit", "ar")]
    doc_en = Document()
    SR._client_requirements_table(doc_en, dr, "en")
    assert not doc_en.tables, "نصوصُ البنود عربية — لا جدولَ في تقريرٍ إنجليزيّ"
    body = "\n".join(p.text for p in doc_en.paragraphs)
    assert "3" in body and "1" in body, "لا صمتَ في الإنجليزي — عددٌ ونوع"
    assert not any("\u0600" <= ch <= "\u06ff" for ch in body)


def test_a_gap_row_is_rendered_as_a_declared_gap():
    doc = Document()
    SR._client_requirements_table(doc, {"requirements": RA.requirement_rows(
        "KEN", "080410")}, "ar")
    cells = [r.cells[3].text for r in doc.tables[-1].rows[1:]]
    assert SR._T("req_status_gap", "ar") in cells


def test_applies_to_drops_requirements_that_exclude_the_code():
    items = [r["item"] or "" for r in RA.requirement_rows("ARE", "080410")]
    assert items and not any("حلال" in i and "لحوم" in i for i in items)


def test_an_uncovered_market_gets_a_declared_entry_gap():
    rows = RA.requirement_rows("KEN", "080410")
    gaps = [r for r in rows if r.get("gap")]
    assert len(gaps) == 1 and gaps[0]["direction"] == "entry"
    assert any(r["direction"] == "exit" for r in rows)


def test_rows_after_the_eligibility_gate_are_conditional():
    rows = [r for r in RA.requirement_rows("DEU", "040900")
            if r["direction"] == "entry"]
    assert rows and not rows[0]["conditional_on_gate"]
    assert all(r["conditional_on_gate"] for r in rows[1:])
    assert not any(r["conditional_on_gate"] for r in
                   RA.requirement_rows("DEU", "080410"))


def test_the_latest_missing_year_and_a_single_observed_year_are_rows():
    doc = Document()
    dr = {"imports": {"value_line": "واردات", "head": "الواردات",
                      "series": [{"year": 2022, "value": 51.0e6}],
                      "years_missing": [2023, 2024]}}
    SR._client_imports_section(doc, dr, "ar")
    years = [r.cells[0].text for r in doc.tables[-1].rows[1:]]
    assert years == ["2022", "2023", "2024"]
    doc2 = Document()
    SR._client_imports_section(doc2, {"imports": {
        "value_line": "واردات", "head": "الواردات",
        "series": [{"year": 2024, "value": 1.0e6}]}}, "ar")
    assert not doc2.tables, "سنةٌ واحدة بلا غيابٍ معلن = لا جدول"


def test_regulatory_fit_is_explained_by_what_it_measures():
    import silk_i18n as I
    note = I.t("pillar_regulatory_note", "ar")
    assert "لا يعني حصولَ منشأتك على الموافقات" in note
