"""تقرير سِلك ٧ — موجة P0 + صفوف الأسعار (الدرسان ٢٧٠ و٢٧١).

report-7 P0 wave: one HHI fact (value + year + source) feeds the economics
view, the chart and the gate; the chart gate checks *metric identity*, not
mere presence of the number; MYR/ringgit prices are read from the one
currency vocabulary; and non-price notes never render as shelf prices.

كلُّ الحالات مصطنعةٌ لاختبارٍ هرمتيّ — لا تُقدَّم حالاتٍ ذهبيةً حقيقية.
"""
import silk_economics as E
import silk_render as R
from silk_data_layer import DataPoint


def _comp(*findings) -> dict:
    return {"missions": {"competitors": {"findings": list(findings),
                                         "failed": False, "summary": ""}}}


def _structured(hhi, year, src="UN Comtrade"):
    return DataPoint({"hhi": hhi, "year": year, "top_suppliers": []},
                     src, 0.9, "ملخّص المورّدين", "2026-09-17")


# ── (١) HHI: السنةُ لا تُقرأ قيمةً ─────────────────────────────────────────

def test_a_year_that_opens_the_sentence_is_not_read_as_the_hhi():
    for text in ("في عام 2024 بلغ مؤشر HHI نحو 1307",
                 "In 2024, the HHI stood at 1307",
                 "بلغ مؤشر HHI لسنة 2023 نحو 1307"):
        out = E.economics_view(_comp({"value": text, "note": "تركّز"}))
        assert out["hhi"] == 1307.0, (text, out["hhi"])


def test_a_real_hhi_of_2024_in_a_structured_field_is_accepted():
    """لا قائمةَ حظرٍ للسنوات: 2024 قيمةُ مؤشرٍ صالحة حين يحملها الحقلُ المهيكل."""
    out = E.economics_view(_comp(_structured(2024, 2023)))
    assert out["hhi"] == 2024.0
    assert out["hhi_year"] == 2023


def _two_sources(order: int) -> dict:
    older = DataPoint(1800.0, "WITS", 0.8, "HHI محسوب من حصص المورّدين",
                      "2026-09-17", data_year=2022)
    newer = _structured(2100, 2024)
    prose = {"value": "في عام 2023 كان مؤشر HHI نحو 900", "note": "تركّز"}
    rows = [older, newer, prose] if order == 0 else [prose, newer, older]
    return _comp(*rows)


def test_hhi_value_year_and_source_come_from_one_record():
    """قيمةٌ من قارئ وسنةٌ ومصدرٌ من قارئٍ آخر = رسمٌ يَنسب رقماً لسجلٍّ لا
    يحمله. الثلاثة من السجلّ المختار نفسه، أيّاً كان ترتيبُ الاكتشافات."""
    for order in (0, 1):
        dr = _two_sources(order)
        out = E.economics_view(dr)
        assert (out["hhi"], out["hhi_year"], out["hhi_source"]) == \
            (2100.0, 2024, "UN Comtrade"), (order, out["hhi"])


def test_the_chart_carries_the_same_hhi_record(monkeypatch):
    for f in ("SILK_CLIENT_METRIC_PRIVACY", "SILK_IMPORTS_SPOTLIGHT",
              "SILK_CONFIDENCE_DISCIPLINE"):
        monkeypatch.delenv(f, raising=False)
    monkeypatch.setenv("SILK_REPORT_CHARTS", "1")
    for order in (0, 1):
        res = {"market": {"name_ar": "ماليزيا", "name_en": "Malaysia",
                          "iso3": "MYS"},
               "product": "قهوة محمصة", "hs_code": "090121",
               "deep_research": dict(_two_sources(order), verdict={
                   "verdict": "GO", "confidence": 0.8},
                   report={"text": "## 1. الخلاصة التنفيذية\nنص.\n"}),
               "markets": []}
        charts = R.build_view(res)["deep_research"].get("charts") or []
        ch = next(c for c in charts if c["id"] == "supplier_concentration")
        assert ch["metric"] == "hhi"
        assert (ch["value"], ch["year"], ch["source"]) == \
            (2100.0, "2024", "UN Comtrade"), order


# ── (٢) بوابةُ الرسم: هويةُ المقياس لا وجودُ الرقم ────────────────────────

def _gauge(value, year="2024"):
    return {"id": "supplier_concentration", "kind": "gauge", "unit": "index",
            "metric": "hhi", "value": value, "year": year,
            "series": [{"label": "م", "value": value}]}


def test_a_chart_number_that_belongs_to_another_metric_is_rejected():
    """84.05 موجودةٌ في الأدلة — لكنها حصةُ مورّد لا HHI. وجودُ الرقم لا يكفي."""
    import silk_quality_gate as Q
    base = {"missions": {"competitors": {"findings": [
        DataPoint({"hhi": 2100, "year": 2024, "top_suppliers": [
            {"partner": "Brazil", "share": 84.05}]},
            "UN Comtrade", 0.9, "ملخّص المورّدين", "2026-09-17")]}}}
    ok = dict(base, charts=[_gauge(2100.0)])
    assert Q._check_chart_metric_identity(ok) == []
    wrong = dict(base, charts=[_gauge(84.05)])
    assert Q._check_chart_backing(wrong) == [], \
        "الحارسُ القديم يمرّره — وهذا بالضبط ما يسدّه فحصُ الهوية"
    hits = Q._check_chart_metric_identity(wrong)
    assert len(hits) == 1
    assert hits[0]["check"] == "chart_metric_identity_mismatch"
    assert "84.05" in hits[0]["note"]
    stale = dict(base, charts=[_gauge(2100.0, year="2022")])
    assert Q._check_chart_metric_identity(stale), "سنةٌ من سجلٍّ آخر تُلتقَط"
    # رسمٌ بلا مقياسٍ مُعلَن لا يُحكَم عليه هنا (يحرسه `_check_chart_backing`).
    anon = dict(base, charts=[dict(_gauge(84.05), metric=None)])
    assert Q._check_chart_metric_identity(anon) == []


# ── (٣) MYR والرينجيت من القاموس الواحد ──────────────────────────────────

def test_myr_prices_are_read_in_arabic_and_english():
    for text in ("سعر الرف 25.80 رينجيت ماليزي", "سعر الرف 25.80 رينجيت",
                 "shelf price 25.80 MYR", "RM25.80 per 250g pack",
                 "RM 25.80", "السعر 25.80 بالرينجيت"):
        assert E.price_numbers_in_text(text) == [25.8], text
    # المدى يُعيد طرفيه (والطرفُ الملتصق قد يتكرّر — سلوكُ المدى القائم).
    assert sorted(set(E.price_numbers_in_text("بين 20 و25 رينجيت"))) == \
        [20.0, 25.0]
    # «RM» بادئةٌ قبل رقم فقط — لا تُطابِق داخل كلمة ولا بلا رقم.
    for text in ("FORM 12", "ARM 12 units", "RM without a number"):
        assert E.price_numbers_in_text(text) == [], text


def test_myr_is_in_the_one_currency_vocabulary(monkeypatch):
    import silk_narrative as N
    # السعرُ الذي قُرئ بعملته لا يفقدها في المرساة — بالرايةِ وبدونها.
    for flag in ("", "1"):
        monkeypatch.setenv(E.RECOGNITION_VOCABULARY_FLAG, flag)
        assert E.currency_in_note("سعر رف 25.80 MYR") == "MYR", flag
        assert N.iso_currency(E.currency_in_note("25.80 رينجيت ماليزي")) \
            == "MYR", flag
    assert N.currency_in("سعر الرف 12 MYR") == "MYR"
    assert N.currency_in("سعر الرف 12 رينجيت ماليزي") == "رينجيت ماليزي"
    assert N.iso_currency("رينجيت ماليزي") == "MYR"
    assert N.iso_currency("رينجيت") == "MYR"


# ── (٤) صفوفُ الأسعار: سعرٌ فعليّ فقط، بلا تكرار ─────────────────────────

_VALID = DataPoint(25.8, "Lazada", 0.8, "سعر رف لعبوة 250 غ، MYR",
                   "2026-09-17")


def _price_view() -> dict:
    findings = [
        DataPoint(None, "Shopee", 0.0,
                  "استُبعد عرض لمنتج مختلف عن المنتج المدروس؛ لا يُستخدم "
                  "للمقارنة السعرية.", "2026-09-17"),
        DataPoint("المنتج يحمل شهادة حلال من JAKIM", "JAKIM", 0.7,
                  "حلال", "2026-09-17"),
        _VALID, _VALID,
        DataPoint("12.5 للعبوة", "Tesco", 0.6, "سعر رف", "2026-09-17"),
    ]
    res = {"market": {"name_ar": "ماليزيا", "name_en": "Malaysia",
                      "iso3": "MYS"},
           "product": "قهوة محمصة", "hs_code": "090121",
           "deep_research": {"missions": {"pricing_scout": {
               "findings": findings, "failed": False, "summary": ""}},
               "verdict": {"verdict": "GO", "confidence": 0.8},
               "report": {"text": "## 1. الخلاصة التنفيذية\nنص.\n"}},
           "markets": []}
    return R.build_view(res)


def test_non_price_notes_never_become_shelf_price_rows():
    rows = _price_view()["deep_research"]["price_rows"]
    blob = " ".join(str(r.get("note")) + str(r.get("value")) for r in rows)
    assert "استُبعد" not in blob, "صفُّ استبعادٍ عُرض سعرَ رف"
    assert "حلال" not in blob, "معلومةُ حلالٍ عُرضت سعرَ رف"
    assert any(r.get("value") == 25.8 for r in rows)
    # سعرٌ مرصودٌ تنقصه العملة يبقى — بحالةٍ معلنة لا بحذف.
    assert any(r.get("value") == "12.5 للعبوة" for r in rows)


def test_the_client_price_section_prints_each_observation_once():
    import silk_reports as SR
    from docx import Document
    doc = Document()
    SR._client_price_observations(doc, _price_view()["deep_research"], "ar")
    bullets = [p.text for p in doc.paragraphs
               if p.style.name == "List Bullet"]
    assert len(bullets) == len(set(bullets)), bullets
    assert not any("استُبعد" in b or "حلال" in b for b in bullets), bullets


# ── (٥) أقفالُ المراجعة الذاتية (§58) ─────────────────────────────────────

def test_a_supplier_share_near_the_keyword_never_becomes_the_hhi():
    """الاحتياطُ النثريّ بأرضية المحرّك نفسِها (100): 84.05 حصةٌ لا HHI."""
    out = E.economics_view(_comp({
        "value": "حصة أكبر مورد 84.05% ومؤشر التركّز مرتفع", "note": "تركّز"}))
    assert out["hhi"] is None
    # والكسرُ ≤1 صيغةٌ مقبولة تُعاد إلى المقياس القانونيّ.
    out = E.economics_view(_comp({"value": "HHI = 0.23", "note": "تركّز"}))
    assert out["hhi"] == 2300.0


def test_general_al_aam_is_not_a_year_marker_for_a_non_year_number():
    """«العام» بمعنى «الإجمالي»: 1850 ليست سنةً معقولة فتبقى قيمة."""
    out = E.economics_view(_comp({"value": "مؤشر HHI العام 1850 وفق كومتريد",
                                  "note": "تركّز"}))
    assert out["hhi"] == 1850.0


def test_the_gate_catches_an_out_of_unit_value_without_any_ledger():
    """حارسٌ مستقلّ عن القارئ: قيمةٌ خارج مقياس HHI تُلتقَط حتى بلا سجلّ."""
    import silk_quality_gate as Q
    dr = {"missions": {}, "charts": [_gauge(84.05)]}
    hits = Q._check_chart_metric_identity(dr, {"entries": {}})
    assert hits and "خارج مقياس hhi" in hits[0]["note"]
    # والسجلُّ المُمرَّر من العرض هو المرجع — لا إعادةُ قراءة `dr` المُقلَّص.
    ok = {"missions": {}, "charts": [_gauge(7118.0)]}
    ledger = {"entries": {"hhi": {"value": 7118.0, "year": 2024}}}
    assert Q._check_chart_metric_identity(ok, ledger) == []
    other = {"entries": {"hhi": {"value": 2100.0, "year": 2024}}}
    assert Q._check_chart_metric_identity(ok, other)


def test_iso_prefix_prices_and_plurals_are_read_and_word_prefixes_are_not():
    assert E.price_numbers_in_text("Carrefour: AED 18.50 for 500g") == [18.5]
    assert E.price_numbers_in_text("SAR 15 للعبوة") == [15.0]
    assert E.price_numbers_in_text("ر.س 15 للعبوة") == [15.0]
    assert E.price_numbers_in_text("السعر 5 دولارات") == [5.0]
    assert E.price_numbers_in_text("10 راندوم") == []
    assert E.currency_in_note("RM 25.80 per 1kg pack retail") == "MYR"
    from silk_render import _is_price_row
    assert _is_price_row("Carrefour: AED 18.50 for 500g", "price observed")


def test_long_observations_differing_after_the_cut_are_both_printed():
    import silk_reports as SR
    from docx import Document
    stem = "سعر رف لعبوة 250 غ بقيمة 25.80 MYR " + "تفاصيل " * 40
    dr = {"price_rows": [{"value": 25.8, "note": stem + "متجر أ", "reason": ""},
                         {"value": 25.8, "note": stem + "متجر ب", "reason": ""}]}
    doc = Document()
    SR._client_price_observations(doc, dr, "ar")
    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 2
