"""الموجة د-١ — البنية: خمسُ درجاتٍ في السجلّ، وصفوفُ مفاتيح بأعلام، وفئةُ
المنتج من ملفٍّ، وقائمةُ ممنوعاتٍ من ملفٍّ يقرؤها حارسُ التصدير.

بلاغُ المالك (2026-09-19، الموجة د): «كل نتيجة تدخل سجل الحقائق بحالتها
الداخلية (مرصود / تقدير / استنتاج / ضعيف / ناقص) ومصدرها… هذه الحالات داخلية
فقط، ولا تظهر أسماؤها للعميل»؛ «جداول الفئات والشهادات ملفات بيانات موثقة
المصدر، قابلة للتحديث دون تعديل الكود»؛ «حارس التصدير يرفضها كما يرفض {{،
بقائمة ممنوعات قابلة للتحديث، مع اختبار على كل المدوّنات».

test-first: كُتب قبل التنفيذ. هرمتي: صفر شبكة، صفر مفتاح، صفر إنفاق.
Run: python3 -m pytest tests/test_wave_d1_structure.py -q
"""
import importlib
import os
import re
import sys
import tempfile

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for _p in (_ROOT, os.path.join(_ROOT, "tools")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import silk_ai_judge as J                                     # noqa: E402
import silk_fact_ledger as L                                  # noqa: E402
import silk_render as R                                       # noqa: E402
import silk_reports as SR                                     # noqa: E402
import silk_style_contract as SC                              # noqa: E402
from gen_verdict_baseline import CANONICAL_BLOBS              # noqa: E402


def _keys():
    return sorted(CANONICAL_BLOBS)


def _view(key):
    mod, fn = CANONICAL_BLOBS[key]
    return R.build_view(getattr(importlib.import_module(mod), fn)())


# ── ١) خمسُ درجاتٍ داخلية — التقديرُ لا يُملأ رقماً عارياً أبداً ────────────

def test_the_ledger_knows_five_grades():
    assert L.ESTIMATE == "estimate" and L.INFERENCE == "inference"
    assert set(L.STATUSES) == {L.OBSERVED, L.WEAK, L.MISSING, L.ESTIMATE,
                               L.INFERENCE}


def _estimate():
    return L.insight_entry(
        "border_price_usd_kg", 4.2, grade=L.ESTIMATE, range=(3.6, 4.8),
        assumption="هامش تجزئة بين 15% و40% (سيناريوهات معلنة)",
        basis=("market_imports_usd",), flip_if="سعر رفّ مرصود أدنى من 3 دولارات",
        source="حساب من السجلّ")


def test_an_estimate_carries_its_range_and_assumption_in_every_rendering():
    """قيدُ المالك: «لا تقدير بلا افتراض معلن ونطاق» — حتى داخل خليّة جدول
    (`bind` يستدعي `render_value` مباشرةً هناك)."""
    e = _estimate()
    assert e["status"] == L.ESTIMATE and e["range"] == (3.6, 4.8)
    val = L.render_value(e, "ar")
    assert "بين" in val and "3.6" in val and "4.8" in val
    assert "إذا افترضنا" in val and "هامش تجزئة" in val
    val_en = L.render_value(e, "en")
    assert "between" in val_en and "assuming" in val_en
    sent = L.render_sentence(e, "ar")
    assert sent.startswith("نقدّر") and sent.endswith(".")


def test_an_estimate_without_assumption_or_range_is_refused():
    with pytest.raises(ValueError):
        L.insight_entry("border_price_usd_kg", 4.2, grade=L.ESTIMATE)
    with pytest.raises(ValueError):
        L.insight_entry("border_price_usd_kg", 4.2, grade=L.ESTIMATE,
                        range=(3.0, 5.0))


def test_an_inference_names_what_it_rests_on():
    e = L.insight_entry("saudi_share_pct", 2.5, grade=L.INFERENCE,
                        basis=("market_imports_usd", "top_supplier_share_pct"),
                        assumption="صافي التجارة مؤشّرٌ للإنتاج لا قياسٌ له")
    sent = L.render_sentence(e, "ar")
    assert sent.startswith("يشير") or sent.startswith("تشير")
    assert "واردات السوق" in sent and "حصة المورّد الأكبر" in sent


def test_grade_words_are_reader_language_not_status_names():
    """`{{status:key}}` يصل نصَّ العميل — فلا اسمَ حالةٍ عارياً."""
    obs = L._entry("market_imports_usd", 1e6, source="UN Comtrade",
                   confidence=0.9)
    weak = L._entry("market_imports_usd", 1e6, source="مدوّنة", confidence=0.3)
    miss = L._entry("market_imports_usd", None)
    for e in (obs, weak, miss, _estimate()):
        for lang in ("ar", "en"):
            s = L.render_status(e, lang)
            assert s not in {"مرصود", "ناقص", "تقدير", "استنتاج", "ضعيف",
                             "observed", "weak", "missing", "estimate",
                             "inference"}, s
    assert "مرصود بتوثيق ضعيف" not in L.render_value(weak, "ar")
    assert "غير رسمية" in L.render_value(weak, "ar")


def test_a_missing_key_that_becomes_an_estimate_is_not_flagged_stale():
    """الانتقالُ غيابٌ → تقديرٌ تحسينٌ لا انزياح؛ وقيمةٌ مرصودةٌ تتغيّر تبقى
    قديمةً كما كانت."""
    snap = {"border_price_usd_kg": {"value": None, "status": L.MISSING,
                                    "source": "", "year": None}}
    ledger = {"entries": {"border_price_usd_kg": _estimate()}}
    assert L.stale_keys(snap, ledger) == []
    snap2 = {"border_price_usd_kg": {"value": 3.0, "status": L.OBSERVED,
                                     "source": "x", "year": 2023}}
    assert L.stale_keys(snap2, ledger)


# ── ٢) صفوفُ المفاتيح بأعلام — قوائمُ التخطّي تُشتقّ لا تُعدَّد ─────────────

def test_key_rows_carry_flags_and_stay_index_compatible():
    row = L.KEYS[0]
    assert row.key == "market_imports_usd" == row[0]
    assert row.words == row[4] and row.label_ar == row[1]
    assert set(L.MANDATORY_TOKEN_KEYS) == {
        r.key for r in L.KEYS if r.mandatory}
    assert L._SELF_DESCRIBING_KEYS == frozenset(
        r.key for r in L.KEYS if r.self_describing)
    assert L._ENGINE_COMPUTED == frozenset(
        r.key for r in L.KEYS if r.engine_computed)


def test_derived_skip_sets_match_the_behaviour_that_was_hand_listed():
    no_check = {r.key for r in L.KEYS if not r.numeric_check}
    assert {"open_conditions", "requirements_count", "competitor_prices",
            "blocking_condition"} <= no_check
    assert "market_imports_usd" not in no_check
    hidden = {r.key for r in L.KEYS if r.writer_hidden}
    assert hidden == {"competitor_prices"}
    unlisted = {r.key for r in L.KEYS if not r.gap_listed}
    assert unlisted == {"blocking_condition", "imports_latest_year"}


def test_an_insight_key_ships_with_no_recognition_words():
    """مفتاحُ استنتاجٍ بكلماتِ تعرّفٍ كان سيُمسَح ضدّ نثر المدوّنات فيكسر
    قفلَ «صفر إنذار كاذب» (مراجعةُ الخطة)."""
    for r in L.KEYS:
        if r.kind == "insight":
            assert r.words == () and not r.numeric_check


@pytest.mark.parametrize("key", _keys())
def test_no_false_alarm_survives_the_restructure(key):
    view = _view(key)
    assert L.check(view, SR.render_markdown(view)) == []


# ── ٣) فئةُ المنتج من ملفٍّ — الثمانيةُ الأصليةُ حرفياً والفصولُ الناقصةُ مغطّاة ─

def test_the_category_table_is_read_from_the_data_file():
    assert J._HS_CATEGORY[3][1] == "معادن/مصنوعات معدنية"
    assert J._HS_CATEGORY[3][2] == "المعايير التقنية ومطابقة المواصفات القياسية للسوق"
    assert [c[1] for c in J._HS_CATEGORY[:8]] == [
        "منتج غذائي/زراعي", "منتج كيميائي/بلاستيكي", "منسوجات/ملابس/أحذية",
        "معادن/مصنوعات معدنية", "آلات/معدّات كهربائية", "مركبات/معدّات نقل",
        "أجهزة/أدوات دقيقة", "أثاث/ألعاب/سلع استهلاكية"]
    assert J._HS_CATEGORY[0][0] == range(1, 25)
    assert "والحلال حيث انطبق" in J._HS_CATEGORY[0][2]      # حرفيٌّ حتى د-٣


def test_previously_unmapped_chapters_now_have_a_category():
    assert J._product_category("270900")[0] == "خامات معدنية/وقود"
    assert J._product_category("410120")[0] == "جلود ومصنوعات جلدية"
    assert J._product_category("440710")[0] == "خشب وورق ومنتجاتهما"
    assert J._product_category("700510")[0] == "حجر وزجاج وخزف ومجوهرات"
    assert J._product_category("930100") is None       # أسلحة — عمداً بلا فئة
    assert J._product_category("990000") is None


@pytest.mark.parametrize("hs,expected", [
    ("020110", "affects"),    # لحوم
    ("040900", "affects"),    # عسل — إضافات/إنزيمات
    ("080410", "inherent"),   # تمور
    ("190219", "affects"),    # معكرونة — إضافات
    ("330499", "affects"),    # تجميل
    ("300490", "affects"),    # أدوية
    ("390210", "none"),       # بوليمرات
    ("410120", "affects"),    # جلود
    ("847989", "none"),       # آلات
    ("720610", "none"),       # معادن
])
def test_religion_relevance_is_a_general_rule_by_chapter(hs, expected):
    assert J.religion_relevance(hs) == expected


def test_processing_level_and_unmapped_are_declared():
    assert J.processing_level("080410") == "raw"
    assert J.processing_level("200899") == "processed"
    assert J.religion_relevance("990000") is None
    assert J.product_profile(None) is None


def test_every_category_has_terms_in_arabic_and_english():
    """فصولٌ تكتسب فئةً تصير خاضعةً لمِصفاة صلة الموزّعين — بلا مفرداتٍ
    تُحذَف روابطُ صحيحة."""
    names = {c[1] for c in J._HS_CATEGORY}
    for name in names:
        assert SC.product_terms(name, "ar"), name
        assert SC.product_terms(name, "en"), name


# ── ٤) قائمةُ الممنوعات من ملفٍّ — الحارسُ يقرؤها وينقّي أو يرفض ─────────────

def test_the_forbidden_terms_file_loads_with_compiled_patterns():
    rows = SC.client_forbidden_terms()
    labels = {r["label"] for r in rows}
    assert {"ledger_status_missing_label", "ledger_key_name",
            "mission_key", "ledger_token"} <= labels
    assert all(hasattr(r["regex"], "search") for r in rows)


@pytest.mark.parametrize("text,label", [
    ("| بنود الاشتراطات | (ناقص) |", "ledger_status_missing_label"),
    ("الحالة: مرصود بتوثيق ضعيف", "ledger_status_weak_label"),
    ("اعتمدنا market_imports_usd في الحساب", "ledger_key_name"),
    ("بعثة pricing_scout لم تجد", "mission_key"),
    ("أطلقها TrendsAgent قبل الكاتب", "agent_class"),
    ("الفحص ledger_value_mismatch أطلق", "ledger_check_name"),
])
def test_file_terms_are_caught_by_the_client_guard(text, label):
    hits = SR._client_forbidden_hits(text, "ar")
    assert any(h.startswith(label) for h in hits), hits


def test_a_status_word_inside_a_normal_sentence_is_not_a_label():
    """«سعر رفّ مرصود في متجر» نثرٌ مشروع؛ «(مرصود)» تسمية."""
    assert not [h for h in SR._client_forbidden_hits(
        "سعر رفّ مرصود في متجر محلي بتاريخ 2026-03-01", "ar")
        if h.startswith("ledger_status")]
    assert [h for h in SR._client_forbidden_hits("التعرفة (مرصود)", "ar")
            if h.startswith("ledger_status_observed")]


def test_redaction_uses_the_files_natural_replacements():
    out = SR._client_redact_text("| التعرفة | (ناقص) | — |", "ar")
    assert "ناقص" not in out and "لم يُعرَف بعد" in out
    out = SR._client_redact_text("لا نعرفه بعد — الناقص: حجم العبوة", "ar")
    assert "الناقص" not in out and "يلزم: حجم العبوة" in out
    out_en = SR._client_redact_text("| tariff | (missing) |", "en")
    assert "missing" not in out_en and "not yet known" in out_en


def test_a_refuse_row_is_never_redacted_only_refused():
    """`{{` لا يُنقّى إلى نصٍّ يبدو سليماً — يُرفض كما كان."""
    txt = "الرسوم {{tariff_applied_pct}} مرتفعة"
    assert "{{" in SR._client_redact_text(txt, "ar")
    assert any(h.startswith("ledger_token")
               for h in SR._client_forbidden_hits(txt, "ar"))


def _client_text(view):
    from docx import Document
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "c.docx")
        SR.render_client_docx(view, p)
        d = Document(p)
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.extend(c.text for c in r.cells)
    return "\n".join(parts)


@pytest.mark.parametrize("key", _keys())
def test_every_codex_client_report_is_free_of_file_terms(key):
    """الاختبارُ الذي طلبه المالك: على كلّ المدوّنات، لا مصطلحَ داخلياً في
    نصّ العميل — ولا يُرفض تصديرٌ كان يمرّ (نقِّ لا ترفض)."""
    txt = _client_text(_view(key))
    assert SR._client_forbidden_hits(txt, "ar") == []
    assert "{{" not in txt


def test_a_bad_regex_row_is_skipped_not_fatal(tmp_path, monkeypatch):
    bad = tmp_path / "terms.csv"
    bad.write_text("label,lang,pattern,replacement_ar,replacement_en,note\n"
                   "broken,both,(unclosed,,,\n"
                   "ok,ar,ناقص\\s*:,يلزم:,,\n", encoding="utf-8")
    monkeypatch.setattr(SC, "_FORBIDDEN_TERMS_PATH", str(bad))
    SC.client_forbidden_terms.cache_clear()
    try:
        rows = SC.client_forbidden_terms()
        assert [r["label"] for r in rows] == ["ok"]
    finally:
        SC.client_forbidden_terms.cache_clear()
