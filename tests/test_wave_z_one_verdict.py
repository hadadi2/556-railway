"""الموجة Z — تدقيقٌ خصوميّ مستقلّ: **حكمٌ واحد يصل المصنوع النهائي**.

> **ما الذي كُسِر فعلاً.** الموجة D أوصلت `silk_decision` إلى المسار العميق
> وسجّلت البند مغلقاً. والقياسُ المستقلّ نفى الإغلاق من ثلاث جهات:
>
> 1. **حكمان في مصنوعٍ واحد** (`Z-01`): المحرّكُ يقول `NO-GO` («الدرجة 0.31
>    دون عتبة الرفض») وغلافُ تقرير العميل المُسلَّم — docx وPDF — يقول «توصية
>    أولية بالدخول»، وحكمُ المحرّك صفرُ ظهورٍ في المستند. السببُ ترتيبيّ:
>    القرارُ كان يُحسَب **بعد** الكاتب وبعد التخزين.
> 2. **اشتقاقٌ ثانٍ للحقل الواحد** (`Z-02`): الموجة B أدخلت
>    `_coverage_basis_tone` في موضعٍ واحد، فصار العرضُ الواحد يحمل شارةً صادقة
>    على الويب وغلافَ docx يقول عكسَها — لأنّ `silk_reports` تُعيد الاشتقاق من
>    الرمز الخام فتتخطّاها.
> 3. **تصديرٌ ينهار على حالةٍ عادية** (`Z-03`): حكمُ الجورية
>    «PRELIMINARY / INCONCLUSIVE» — الصادرُ كلّما أعلنت بعثةٌ واحدة فجوة —
>    يُصنَّف `inconclusive`، ولا لونَ له في `silk_reports`، فيرفع
>    **KeyError** ويسقط تصديرُ تقرير العميل كلَّه قبل أن يُكتَب حرف.
>
> ومعها ثلاثةٌ أصغر: معرّفاتٌ داخلية في نصٍّ يقرؤه إنسان (`Z-04`)، وانقلابُ
> قطبِ عمود المنافسة في الحجّة المضادّة (`Z-05`)، وأساسُ الحكم الذي لا يصل
> المصنعَ إطلاقاً (`Z-06`).

كلُّ قفلٍ هنا مربوطٌ بإعادةِ إنتاجٍ مُقاسة في `docs/WAVE_Z_AUDIT.md`.
"""
from __future__ import annotations

import copy
import io
import os
import re
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tools"))

import silk_decision as D                                # noqa: E402
import silk_deep_pillars as DP                           # noqa: E402
import silk_render                                       # noqa: E402
import silk_reports                                      # noqa: E402
from canonical_netherlands import netherlands_research_blob  # noqa: E402
from silk_requirements_agent import regulatory_state     # noqa: E402


# ── أدواتٌ مشتركة ────────────────────────────────────────────────────────

def _jury(state: str) -> dict:
    """حكمُ الجورية **بشكله الإنتاجيّ** — لا قاموسٌ مصنوع باليد.

    الفرقُ حاسم: القاموسُ المصنوع بلا `basis="data_coverage"` يتخطّى تصحيحَ
    الموجة B فيُنتِج إعادةَ إنتاجٍ زائفة. هذا يبني التقريرَ فتُصدِره
    `JuryCommittee` نفسُها.
    """
    import silk_agents
    from silk_data_layer import DataPoint

    def dp(v):
        return DataPoint(value=v, source="UN Comtrade", confidence=0.8,
                         note="—")
    ok = silk_agents.AgentReport(agent_name="a", findings=[dp(1.0)],
                                 failed=False)
    if state == "full":
        reports = [ok, silk_agents.AgentReport(agent_name="b",
                                               findings=[dp(2.0)], failed=False)]
    elif state == "partial":
        reports = [ok, silk_agents.AgentReport(agent_name="b",
                                               findings=[dp(None)], failed=True)]
    else:
        reports = [silk_agents.AgentReport(agent_name="b", findings=[dp(None)],
                                           failed=True)]
    v = silk_agents.JuryCommittee.evaluate(reports)
    v["synthesis_stage"] = 1
    return v


def _view(*, jury_state: str = "full", engine: "str | None" = "NO-GO",
          lang: str = "ar", capped_confidence: "float | None" = None) -> dict:
    """عرضٌ كاملٌ على شكل المسار العميق بعد الموجة Z."""
    reg = regulatory_state("NLD", "080410")
    r = netherlands_research_blob()
    dr = r["deep_research"]
    # **قرارٌ حقيقيّ من `decide` لا قاموسٌ مُعدَّلٌ بعده** (مراجعةٌ ذاتية §٥٨):
    # كان الحكمُ يُكتَب فوق ناتج المحرّك، فيبقى `counter_case` مبنيّاً على
    # الحكم الأصليّ — والفرعُ الشرطيّ وحدَه بلا أرقام. فقفلُ «لا كسرَ آليّ
    # خام» كان يفحص نصّاً **لا يمكن** أن يحمل كسراً: حارسٌ لا يستطيع الفشل.
    # الآن يُنتَج NO-GO بإفقارِ المدخلات فعلاً، فيمرّ النصُّ كاملاً بالقفل.
    if engine == "NO-GO":
        pi = DP.build_pillar_inputs(dr, regulatory=reg)
        pi["market_attractiveness"] = {"tam_usd": 2.0e5, "import_cagr_pct": -9.0,
                                       "gdp_per_capita_usd": 800.0,
                                       "saudi_share_pct": 19.0}
        pi["competition_intensity"] = {"hhi": 6000.0,
                                       "top_supplier_share_pct": 80.0,
                                       "named_company_count": 1}
        pi["regulatory_fit"] = {"tariff_applied_pct": 28.0,
                                "entry_requirements_count": 1,
                                "eligibility_gate": False}
        pi["profitability"], pi["risk"] = {}, {}
        decision = D.decide({"pillar_inputs": pi, "coverage": 1.0})
        assert decision["verdict"] == "NO-GO", decision["verdict"]
    elif engine == "GO":
        # وفرعُ الإيجاب بنفس الانضباط: `GO` **حقيقيّ** من `decide` بمدخلاتٍ
        # قويّة، لا حكمٌ يُكتَب فوق ناتجٍ آخر (وإلّا صار السيناريو صوريّاً).
        pi = DP.build_pillar_inputs(dr, regulatory=reg)
        pi["market_attractiveness"] = {"tam_usd": 8.0e9, "import_cagr_pct": 18.0,
                                       "gdp_per_capita_usd": 62_000.0,
                                       "saudi_share_pct": 1.0}
        pi["competition_intensity"] = {"hhi": 300.0,
                                       "top_supplier_share_pct": 6.0,
                                       "named_company_count": 9}
        pi["regulatory_fit"] = {"tariff_applied_pct": 0.0,
                                "entry_requirements_count": 8,
                                "eligibility_gate": False}
        pi["profitability"] = {"border_unit_value_usd_kg": 9.0,
                               "saudi_border_unit_value_usd_kg": 3.0,
                               "margin_at_border_pct": 45.0}
        pi["risk"] = {"political_stability_wgi": 1.4, "fx_volatility_pct": 0.6,
                      "supplier_concentration_hhi": 300.0,
                      "critical_risk": False}
        decision = D.decide({"pillar_inputs": pi, "coverage": 1.0})
        assert decision["verdict"] == "GO", decision["verdict"]
    elif engine:
        decision = copy.deepcopy(DP.decide_for_deep(dr, regulatory=reg))
        decision["verdict"] = engine
    else:
        decision = None
    dr["verdict"] = DP.promote_engine_verdict(_jury(jury_state), decision)
    if capped_confidence is not None:
        # ترتيبُ الإنتاج: `cap_confidence_for_flagged_hs` يسبق التخزين
        # ويسبق `build_view` — فالتسقيفُ يُحقَن **قبل** بناء العرض لا بعده.
        dr["verdict"]["confidence"] = capped_confidence
    r["markets"] = [{"iso3": "NLD", "name_ar": "هولندا", "name_en": "Netherlands",
                     "rank": 1, "deep": True, "regulatory": reg,
                     "components": DP.build_components(dr),
                     "decision": decision}]
    return silk_render.build_view(r, lang)


def _docx_lines(view: dict, tmp_path, client: bool = True) -> list[str]:
    p = str(tmp_path / "z.docx")
    (silk_reports.render_client_docx if client
     else silk_reports.render_docx)(view, p)
    xml = zipfile.ZipFile(p).read("word/document.xml").decode("utf-8")
    txt = re.sub(r"<[^>]+>", "",
                 xml.replace("</w:p>", "\n").replace("</w:tc>", " | "))
    return [l.strip() for l in txt.split("\n") if l.strip()]


# ── Z-01 · حكمُ المحرّك يسبق الكاتبَ ويحكم كلَّ سطح ───────────────────────

def test_the_engine_verdict_is_promoted_into_the_one_verdict_field():
    """المُقدَّمُ حكمُ المحرّك؛ وقراءةُ التغطية تبقى **مُعلَنةً** بجواره."""
    jury = _jury("full")
    assert jury["basis"] == "data_coverage"          # شكلُ الإنتاج مُثبَت
    out = DP.promote_engine_verdict(
        jury, {"verdict": "NO-GO", "confidence": 0.62, "score": 0.31,
               "why": "الدرجة دون العتبة"})
    assert out["verdict"] == "NO-GO"
    assert out["confidence"] == 0.62
    assert out["basis"] == "silk.decision/v1", (
        "بقاءُ الأساس «تغطيةَ أدلة» يُعيد تسميةَ حالةِ الأدلة فوق حكمِ محرّك")
    assert out["coverage_verdict"] == jury["verdict"], "قراءةُ التغطية فُقدت"
    assert out["coverage_state"] == "coverage_full"


def test_no_engine_result_leaves_the_verdict_byte_for_byte_unchanged():
    """تكافؤٌ رجعيّ تامّ: بلا محرّك، السلوكُ كما كان قبل الموجة حرفياً."""
    jury = _jury("partial")
    for absent in (None, {}, {"error": "boom"}, {"verdict": ""}):
        assert DP.promote_engine_verdict(jury, absent) == jury


def test_the_engine_verdict_reaches_the_delivered_client_document(tmp_path):
    """**الحادثةُ الأصلية:** محرّكٌ يقول NO-GO وغلافٌ يقول «توصية أولية»."""
    v = _view(jury_state="full", engine="NO-GO")
    assert v["decision"]["verdict"] == "NO-GO"
    assert v["deep_research"]["verdict_label"] == "عدم الدخول حالياً"
    lines = _docx_lines(v, tmp_path)
    assert "عدم الدخول حالياً" in lines[:8], (
        f"غلافُ المستند لا يحمل حكمَ المحرّك — {lines[:6]}")
    assert not any("توصية أولية بالدخول" in l for l in lines), (
        "عدّادُ تغطيةِ البعثات ما يزال يُقدَّم توصيةً تجارية في المُسلَّم")


def test_every_surface_of_one_view_carries_the_same_verdict(tmp_path):
    """الشارةُ والمختصرُ والغلافُ لا يتباعدون — أياً كانت الحالة."""
    for jury_state, engine in (("full", "NO-GO"), ("full", None),
                               ("partial", None), ("none", None)):
        v = _view(jury_state=jury_state, engine=engine)
        label = v["deep_research"]["verdict_label"]
        assert label in v["brief"][0], (
            f"المختصرُ يخالف التسمية ({jury_state}/{engine}): {v['brief'][0]}")
        lines = _docx_lines(v, tmp_path)
        assert label in lines[:8], (
            f"الغلافُ يخالف التسمية ({jury_state}/{engine}): {lines[:6]}")


# ── Z-02 · حقلٌ واحد، لا اشتقاقٌ ثانٍ ────────────────────────────────────

def test_the_reconciled_tone_round_trips_through_the_shared_classifier():
    """تمريرُ مفتاحِ التصنيف يعيده كما هو — فتصل القيمةُ المُصالَحة كلَّ عارض."""
    for key in silk_render._VERDICT_LABELS_AR:
        assert silk_render._verdict_tone(key) == key
    # ولا يفسد تصنيفُ الرموز الخام:
    assert silk_render._verdict_tone("PRELIMINARY GO — مبدئي") == "preliminary"
    assert silk_render._verdict_tone("NO-GO (insufficient data)") == "nogo"
    assert silk_render._verdict_tone("CONDITIONAL-GO") == "conditional"


def test_the_report_layer_reads_the_reconciled_field_first():
    """`silk_reports` لا تُعيد الاشتقاقَ من الرمز الخام فوق الحقل المُصالَح."""
    dr = {"verdict_tone": "data_partial",
          "verdict": {"verdict": "PRELIMINARY / INCONCLUSIVE — ناقص"}}
    assert silk_reports._resolve_vtxt(dr) == "data_partial"
    # وبلا الحقل (مدوّنةٌ مخزَّنة قديمة) يبقى المسارُ القديم كما هو:
    assert silk_reports._resolve_vtxt(
        {"verdict": {"verdict": "CONDITIONAL-GO"}}) == "CONDITIONAL-GO"


def test_an_evidence_state_is_not_announced_as_a_recommendation():
    """C9: «التوصية: بحثٌ ناقص» جملةٌ تناقض نفسَها — البادئةُ تتبع الحقيقة."""
    v = _view(jury_state="partial", engine=None)
    assert v["deep_research"]["verdict_tone"].startswith("data_")
    assert v["brief"][0].startswith("الحالة:"), v["brief"][0]
    v2 = _view(jury_state="full", engine="NO-GO")
    assert v2["brief"][0].startswith("التوصية:"), v2["brief"][0]


def test_a_second_render_path_for_the_decision_never_returns():
    """قفلُ الموجة D يبقى: لا مسارَ عرضٍ ثانٍ للحكم المحسوب."""
    src = io.open(os.path.join(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))), "silk_render.py"), encoding="utf-8").read()
    assert "_decision_from_engine" not in src


# ── Z-03 · التسليمُ لا ينهار على تصنيفٍ بلا لون ──────────────────────────

def test_every_verdict_tone_has_a_colour_and_an_academic_sentence():
    """خريطتان تتباعدان = تصديرٌ ينهار. القفلُ يمنع التباعد لا يُصلِح أثرَه."""
    missing_colour = [t for t in silk_render._VERDICT_LABELS_AR
                      if t not in silk_reports._VERDICT_TEXT_COLORS]
    missing_rec = [t for t in silk_render._VERDICT_LABELS_AR
                   if t not in silk_reports._ACADEMIC_MAIN_REC]
    assert not missing_colour, f"تصانيفُ بلا لون تُسقِط التصدير: {missing_colour}"
    assert not missing_rec, f"تصانيفُ بلا نصٍّ أكاديمي: {missing_rec}"


def test_a_study_with_one_declared_gap_still_exports(tmp_path):
    """**الحادثة:** بعثةٌ واحدة أعلنت فجوة ⇒ KeyError ⇒ لا تقرير إطلاقاً."""
    v = _view(jury_state="partial", engine=None)
    assert v["deep_research"]["verdict_tone"] == "data_partial"
    lines = _docx_lines(v, tmp_path)          # كان يرفع KeyError('inconclusive')
    assert len(lines) > 40


def test_an_unknown_future_tone_never_blocks_delivery(tmp_path):
    """اللونُ زينةٌ لا حكم: نقصُه لا يجوز أن يمنع تسليمَ تقرير."""
    from docx import Document
    doc = Document()
    silk_reports._add_verdict_badge(doc, "tone_that_does_not_exist", "ar")
    assert doc.paragraphs


# ── Z-04 · لا معرّفٍ داخليٍّ في نصٍّ يقرؤه إنسان ─────────────────────────

_INTERNAL_KEYS = ("political_stability", "regulatory_quality", "fx_stability",
                  "price_position", "saudi_momentum", "named_density",
                  "top_share", "tam_log", "legibility")


def test_conditions_name_the_gap_in_words_not_in_identifiers():
    d = D.decide({"coverage": 0.5, "pillar_inputs": {
        "market_attractiveness": {"tam_usd": 2.0e8},
        "competition_intensity": {"hhi": 1200},
        "regulatory_fit": {"entry_requirements_count": 5},
        "profitability": {}, "risk": {}}})
    blob = " ".join(d["conditions"])
    leaked = [k for k in _INTERNAL_KEYS if k in blob]
    assert not leaked, f"معرّفاتٌ داخلية في نصٍّ للقارئ: {leaked}"
    assert "الاستقرار السياسي" in blob and "أداء اللوجستيات" in blob


# ── Z-05 · قطبُ عمود المنافسة في الحجّة المضادّة ─────────────────────────

def _bundle(hhi: float, top: float, named: int) -> dict:
    return {"coverage": 1.0, "pillar_inputs": {
        "market_attractiveness": {"tam_usd": 3e8, "import_cagr_pct": 6.0,
                                  "gdp_per_capita_usd": 45_000,
                                  "saudi_share_pct": 2.0},
        "competition_intensity": {"hhi": hhi, "top_supplier_share_pct": top,
                                  "named_company_count": named},
        "regulatory_fit": {"tariff_applied_pct": 4.0,
                           "entry_requirements_count": 6,
                           "eligibility_gate": False},
        "profitability": {"border_unit_value_usd_kg": 5.0,
                          "saudi_border_unit_value_usd_kg": 4.2,
                          "margin_at_border_pct": 16.0},
        "risk": {"political_stability_wgi": 0.9, "fx_volatility_pct": 1.2,
                 "supplier_concentration_hhi": 900, "critical_risk": False}}}


def test_a_fragmented_market_is_never_argued_as_a_reason_against_entry():
    """**الحادثة:** HHI=200 (أقوى ما يدعم الدخول) قُدِّم «يدعم رفضاً»."""
    case = D.decide(_bundle(200.0, 4.0, 9))["counter_case"]["case"]
    assert "شدة المنافسة" not in case, (
        f"سوقٌ مُفتَّتة ما تزال تُقدَّم حجّةً ضدّ الدخول: {case}")


def test_a_concentrated_market_is_the_argument_against_entry():
    """والاتجاهُ الآخر يبقى صحيحاً — لا قلبٌ في الجهة المقابلة."""
    case = D.decide(_bundle(4800.0, 68.0, 2))["counter_case"]["case"]
    assert "شدة المنافسة" in case


def test_the_counter_case_prints_human_percentages_not_raw_fractions():
    """الدرس ٩٢: لا كسرٌ آليٌّ خام على وجه التقرير."""
    cc = D.decide(_bundle(4800.0, 68.0, 2))["counter_case"]
    blob = f"{cc['case']} {cc['rebuttal']}"
    assert "%" in blob
    assert not re.search(r"\(0\.\d+\)", blob), blob


def test_the_why_line_carries_no_raw_machine_fraction():
    """التعليقُ فوق السطر يقول ذلك منذ المرحلة ٥ — والسطرُ كان يطبع «0.31»."""
    for bundle in (_bundle(200.0, 4.0, 9), _bundle(4800.0, 68.0, 2)):
        for cov in (1.0, 0.4):
            b = copy.deepcopy(bundle)
            b["coverage"] = cov
            why = D.decide(b)["why"]
            assert not re.search(r"\b0\.\d+\b", why), why


def test_pillar_strength_mirrors_the_score_transform_exactly():
    assert D.pillar_strength("competition", 0.19) == 0.81
    assert D.pillar_strength("market", 0.85) == 0.85
    assert D.pillar_strength("profit", None) is None


# ── Z-06 · أساسُ الحكم يصل المصنع فعلاً ──────────────────────────────────

def test_the_client_report_shows_what_the_decision_rests_on(tmp_path):
    """**الحادثة:** صفرُ ذكرٍ للأعمدة أو القاعدة أو الحجّة المضادّة."""
    lines = _docx_lines(_view(jury_state="full", engine="NO-GO"), tmp_path)
    blob = "\n".join(lines)
    assert "على أيّ أساس صدر هذا الحكم" in blob
    assert "قاعدة الحكم مُعلنة قبل النظر في الأرقام" in blob
    assert "جاذبية السوق" in blob and "شدة المنافسة" in blob
    assert "أقوى ما يُقال ضدّ هذا الحكم" in blob


def test_an_uncomputed_pillar_is_declared_with_what_it_needs(tmp_path):
    """الغيابُ يُعلَن ويقول ما ينقصه — لا خانةٌ فارغة ولا صفرٌ مختلَق."""
    blob = "\n".join(_docx_lines(_view(jury_state="full", engine="NO-GO"),
                                 tmp_path))
    # هدف الدراسة الاحترافية (البند ٣): صياغة العميل «لا نعرفه بعد /
    # ينقصنا:» بدل مصطلح القياس — النية نفسها: الغياب معلن ويقول ما ينقصه.
    assert "لم يُحسب بعد" in blob
    assert "ينقصنا:" in blob
    leaked = [k for k in _INTERNAL_KEYS if k in blob]
    assert not leaked, f"معرّفاتٌ داخلية في مُسلَّم العميل: {leaked}"


def test_no_raw_machine_score_reaches_the_client_face(tmp_path):
    """الدرس ٩٢ على المُسلَّم النهائي: النِّسَبُ بشرية لا كسورٌ خام."""
    blob = "\n".join(_docx_lines(_view(jury_state="full", engine="NO-GO"),
                                 tmp_path))
    assert "قوة هذه الفرصة في تقييمنا" in blob   # البند ٣: جملة معنى
    assert not re.search(r"\b0\.\d{2,3}\b", blob), (
        "كسرٌ عشريٌّ خام على وجه تقرير العميل")


def test_the_basis_section_removes_itself_when_no_decision_was_computed(
        tmp_path):
    """قسمٌ بلا مادّة يُسقِط نفسَه — لا هيكلَ فارغ يوحي بحسابٍ لم يجرِ."""
    blob = "\n".join(_docx_lines(_view(jury_state="full", engine=None),
                                 tmp_path))
    assert "على أيّ أساس صدر هذا الحكم" not in blob


def test_the_basis_section_is_authored_in_english_too():
    """C9: نصٌّ مكتوبٌ أصالةً في اللغتين — لا ترجمةَ آليّة ولا تسرّبَ لغة."""
    import silk_i18n
    for key in ("decision_basis_head", "decision_rule_lead", "pillar_col",
                "pillar_not_computed", "pillar_missing_lead",
                "counter_case_head", "counter_case_computed"):
        en = silk_i18n.t(key, "en", go=65, nogo=45, conf=60, parts="x",
                         strong="A", weak="B", strong_pct=1, weak_pct=2,
                         score=1)
        assert en and not silk_i18n.foreign_prose_spans(en, "en"), key


# ── السيناريوهاتُ الثلاثة الحرجة (أمرُ المالك بعد حكم Z) ────────────────
#
# الثابتُ المطلوب: مدخلاتٌ خام ← تحليل ← `silk_decision` ← نموذجُ التقرير ←
# DOCX/PDF. ولا شيءَ يكتب فوق حكم المحرّك في المصنوع النهائيّ.

@pytest.mark.parametrize("engine,label", [
    ("NO-GO", "عدم الدخول حالياً"),
    ("GO", "التوصية بالدخول"),
])
def test_the_engine_verdict_is_exactly_what_the_document_says(
        engine, label, tmp_path):
    """الاتجاهان معاً: محرّكٌ سالب ⇒ مستندٌ سالب، وموجبٌ ⇒ موجب."""
    v = _view(engine=engine)
    assert v["decision"]["verdict"] == engine
    assert v["deep_research"]["verdict_label"] == label
    lines = _docx_lines(v, tmp_path)
    assert label in lines[:8], f"غلافُ المستند يخالف المحرّك: {lines[:6]}"
    # ولا تسميةَ حكمٍ **أخرى** في المستند كلِّه:
    others = {l for l in silk_render._VERDICT_LABELS_AR.values()} - {label}
    leaked = sorted({o for o in others for l in lines if l.strip() == o})
    assert not leaked, f"تسميةُ حكمٍ ثانية في المصنوع: {leaked}"


def test_no_default_recommendation_can_overwrite_the_engine(tmp_path):
    """قراءةُ التغطية وقراءةُ كلود كلتاهما لا تعلوان حكمَ المحرّك."""
    reg = regulatory_state("NLD", "080410")
    r = netherlands_research_blob()
    dr = r["deep_research"]
    dec = DP.decide_for_deep(dr, regulatory=reg)
    jury = _jury("full")                       # «PRELIMINARY GO — مبدئي إيجابي»
    jury["ai"] = {"verdict": "NO-GO", "confidence": 0.9,
                  "reasoning": "قراءةٌ استشارية مخالِفة"}
    dr["verdict"] = DP.promote_engine_verdict(jury, dec)
    r["markets"] = [{"iso3": "NLD", "name_ar": "هولندا", "rank": 1,
                     "deep": True, "regulatory": reg,
                     "components": DP.build_components(dr), "decision": dec}]
    v = silk_render.build_view(r, "ar")
    want = silk_render._verdict_label(
        silk_render._verdict_tone(dec["verdict"]), "ar")
    assert v["deep_research"]["verdict_label"] == want
    assert want in _docx_lines(v, tmp_path)[:8]
    # والقراءتان محفوظتان **مُعلَنتين** لا ممحوّتين:
    assert v["deep_research"]["verdict"]["coverage_verdict"]
    assert (v["deep_research"]["verdict"].get("ai") or {}).get("verdict")


def test_incomplete_evidence_yields_a_valid_inconclusive_report(tmp_path):
    """الحالةُ الثالثة **مواطنٌ من الدرجة الأولى**: تقريرٌ صالح، حالةٌ مُعلَنة،
    ولا انزلاقَ صامتٌ إلى إيجابٍ ولا إلى رفض."""
    v = _view(jury_state="partial", engine=None)
    tone = v["deep_research"]["verdict_tone"]
    assert tone == "data_partial", tone
    label = v["deep_research"]["verdict_label"]
    for commercial in ("التوصية بالدخول", "عدم الدخول حالياً", "دخول مشروط",
                       "توصية أولية بالدخول"):
        assert commercial != label, (
            f"حالةُ أدلةٍ ناقصة انزلقت إلى حكمٍ تجاريّ: {label}")
    lines = _docx_lines(v, tmp_path)                 # لا KeyError، وتقريرٌ فعليّ
    assert len(lines) > 40
    assert label in lines[:8]
    assert v["brief"][0].startswith("الحالة:")


# ── Z-06 · سطحُ المصنع يعرض أساسَ الحكم فعلاً ────────────────────────────

def test_the_decision_basis_is_built_once_and_consumed_everywhere():
    """بانٍ واحد: `silk_render.decision_basis`؛ والمُصدِّرُ يعرض ولا يبني."""
    v = _view(engine="NO-GO")
    basis = (v.get("decision") or {}).get("basis")
    assert basis and basis.get("pillars"), "الأساسُ لا يصل نموذجَ العرض"
    assert basis.get("rule_line") and basis.get("counter_case_line")
    assert any(r["strength_pct"] is None for r in basis["pillars"])
    src = io.open(os.path.join(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))), "silk_reports.py"), encoding="utf-8").read()
    body = src[src.index("def _client_decision_basis"):
               src.index("def render_client_docx")]
    for banned in ("pillar_strength", "part_labels", "pillar_label",
                   "_D._GO", "counter_case_computed"):
        assert banned not in body, (
            f"المُصدِّرُ يُعيد بناءَ الأساس بدل عرضه: {banned}")


def test_the_factory_surface_renders_the_decision_basis():
    """`web/platform.html` يقرأ البنيةَ الجاهزة — لا حسابَ ولا تعريبَ في JS."""
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    html = io.open(os.path.join(root, "web", "platform.html"),
                   encoding="utf-8").read()
    assert "v.decision && v.decision.basis" in html, (
        "سطحُ المصنع ما يزال لا يعرض أساسَ الحكم (Z-06)")
    for key in ("col_pillar", "col_strength", "col_note", "not_computed",
                "conditions_head", "counter_case_line", "rule_line",
                "score_line"):
        assert key in html, f"مفتاحٌ من البنية لا يُعرَض: {key}"
    # ولا تسمياتِ أعمدةٍ مكتوبةً في JS (خريطةٌ ثانية تتباعد):
    for banned in ("جاذبية السوق", "شدة المنافسة", "الملاءمة التنظيمية"):
        assert banned not in html, f"تسميةُ عمودٍ منسوخةٌ في JS: {banned}"
    # وشارةُ الحكم: لكلّ تصنيفٍ صنفُ CSS — نفسُ درسِ `Z-03` على الويب.
    for tone in silk_render._VERDICT_LABELS_AR:
        assert f".rverdict.{tone}" in html, f"تصنيفٌ بلا لونٍ على الشارة: {tone}"


def test_an_absent_denominator_is_declared_not_invented():
    """Z-10: «0/4» على صفٍّ بلا مكوّنات — بسطٌ صادقٌ ومقامٌ مُثبَّت."""
    r = netherlands_research_blob()
    r["markets"] = [{"iso3": "NLD", "name_ar": "هولندا", "rank": 1,
                     "deep": True}]
    # تحديث قفل معلَن (موجة سدّ الفجوات F4 — البند 16): مفردتا الغياب —
    # «غير متاح» بدل «غير مرصود».
    assert silk_render.build_view(r, "ar")["markets"][0][
        "components_present"] == "غير متاح"
    assert silk_render.build_view(r, "en")["markets"][0][
        "components_present"] == "Not available"
    # ومع مكوّناتٍ فعلية: المقامُ **مقيسٌ** — يساوي عددَها لا رقماً مُثبَّتاً.
    v = _view(engine="NO-GO")
    n = len(v["markets"][0]["components_detail"])
    assert v["markets"][0]["components_present"].endswith(f"/{n}")


# ── ما التقطته المراجعةُ الذاتية (§٥٨) على هذه الموجة نفسِها ────────────

def _stored_pre_z_view():
    """مدوّنةٌ **مخزَّنة قبل الترقية**: صفٌّ بقرارِ محرّك وحكمُ جوريةٍ خام."""
    reg = regulatory_state("NLD", "080410")
    r = netherlands_research_blob()
    dr = r["deep_research"]
    dr["verdict"] = _jury("partial")          # بلا `promote_engine_verdict`
    r["markets"] = [{"iso3": "NLD", "name_ar": "هولندا", "rank": 1,
                     "deep": True, "regulatory": reg,
                     "components": DP.build_components(dr),
                     "decision": DP.decide_for_deep(dr, regulatory=reg)}]
    return silk_render.build_view(r, "ar")


def test_a_stored_run_from_before_the_promotion_still_shows_one_verdict(
        tmp_path):
    """الترقيةُ تخدم الجديد؛ والمخزَّنُ يُصالَح عند العرض لا يُترك منشطراً."""
    v = _stored_pre_z_view()
    engine = (v.get("decision") or {}).get("verdict")
    assert engine, "المدوّنةُ المخزَّنة بلا قرارٍ — الحالةُ غير مُختبَرة"
    label = v["deep_research"]["verdict_label"]
    assert label == silk_render._verdict_label(
        silk_render._verdict_tone(engine), "ar"), (
        f"الشارةُ «{label}» تخالف حكمَ المحرّك «{engine}» على مدوّنةٍ مخزَّنة")
    assert label in _docx_lines(v, tmp_path)[:8]


def test_the_verdict_structure_valve_still_turns_the_layer_off(tmp_path):
    """الدرس ٩١: طبقةٌ مرئية بلا صمّامٍ عامل ليست طبقةً محكومة."""
    prev = os.environ.get("SILK_VERDICT_STRUCTURE_ENABLED")
    os.environ["SILK_VERDICT_STRUCTURE_ENABLED"] = "0"
    try:
        blob = "\n".join(_docx_lines(_view(engine="NO-GO"), tmp_path))
    finally:
        if prev is None:
            os.environ.pop("SILK_VERDICT_STRUCTURE_ENABLED", None)
        else:
            os.environ["SILK_VERDICT_STRUCTURE_ENABLED"] = prev
    assert "قاعدة الحكم مُعلنة قبل النظر في الأرقام" not in blob
    assert "أقوى ما يُقال ضدّ هذا الحكم" not in blob


def test_one_confidence_number_in_one_document(tmp_path):
    """PR A §A1: رمزٌ مُعلَّم يسقف الثقة — والقسمُ الجديد يعرض المسقوفة."""
    capped = 0.5
    v = _view(engine="NO-GO", capped_confidence=capped)
    blob = "\n".join(_docx_lines(v, tmp_path))
    assert f"مستوى الثقة في هذا التقييم نحو {round(capped * 100)}%" in blob, (
        "القسمُ يعرض ثقةَ المحرّك الخام بدل المسقوفة — غلافٌ برقمٍ وقسمٌ بآخر")


def test_the_brief_line_stays_in_one_language(tmp_path):
    """سطرُ المختصر أحاديُّ اللغة **بلغة التقرير** — كان عربياً دائماً حتى
    على lang=en (فيُفتتح تقرير مصنع إنجليزي بواجهة عربية)؛ صيد الفجوات ٣
    (الدرس 165) جعله يتبع اللغة، والعقد الأصلي باقٍ: لا سطر نصف مترجَم."""
    import silk_i18n
    line_en = _view(engine="NO-GO", lang="en")["brief"][0]
    assert not silk_i18n.foreign_prose_spans(line_en, "en"), line_en
    line_ar = _view(engine="NO-GO", lang="ar")["brief"][0]
    assert not silk_i18n.foreign_prose_spans(line_ar, "ar"), line_ar


def test_the_counter_case_heading_never_stands_alone(tmp_path):
    """قسمٌ بلا متن هو الهيكلُ الفارغ الذي تمنعه الدالةُ بنصّ توثيقها."""
    lines = _docx_lines(_view(engine="NO-GO", lang="en"), tmp_path) \
        if False else _docx_lines(_view(engine="NO-GO"), tmp_path)
    for head in ("أقوى ما يُقال ضدّ هذا الحكم", "ما يجب إغلاقه قبل الالتزام"):
        if head in lines:
            i = lines.index(head)
            assert i + 1 < len(lines) and lines[i + 1] != "", f"عنوانٌ يتيم: {head}"


def test_the_named_conditions_promised_by_the_rule_are_actually_printed(
        tmp_path):
    """قاعدةُ الحكم تَعِد بشروطٍ «مسمّاة أدناه» — فلتكن أدناه فعلاً."""
    blob = "\n".join(_docx_lines(_view(engine="NO-GO"), tmp_path))
    assert "ما يجب إغلاقه قبل الالتزام" in blob
    assert "لم يُحسب بعد" in blob   # البند ٣: صياغة العميل


def test_the_coverage_disclaimer_is_not_dropped_by_the_promotion():
    """تحفّظُ عدم الاختلاق في `note` كان يُمحى بسببِ المحرّك."""
    jury = _jury("full")
    assert jury.get("note"), "الجوريةُ لم تعد تحمل تحفّظاً — القفلُ عمي"
    out = DP.promote_engine_verdict(jury, {"verdict": "NO-GO", "why": "س",
                                           "confidence": 0.6, "score": 0.3})
    assert out["note"] == jury["note"]
    assert out["decision_why"] == "س"


def test_the_confidence_basis_line_cannot_print_an_impossible_percentage():
    d = D.decide({**_bundle(200.0, 4.0, 9), "coverage": 3.0})
    assert "300%" not in d["confidence_basis"], d["confidence_basis"]
    assert "100%" in d["confidence_basis"]


def test_the_rebuttal_carries_no_raw_machine_fraction():
    """كلُّ فروع الحجّة المضادّة، لا الفرعُ الشرطيُّ وحده."""
    seen = set()
    for b in (_bundle(200.0, 4.0, 9), _bundle(4800.0, 68.0, 2)):
        for cov in (1.0, 0.4):
            d = D.decide({**b, "coverage": cov})
            seen.add(d["verdict"])
            cc = d["counter_case"]
            blob = f"{cc['case']} {cc['rebuttal']}"
            assert not re.search(r"\b0\.\d+\b", blob), blob
    # وفرعُ الرفض صراحةً (لا يبلغه المسحُ أعلاه):
    poor = {"coverage": 1.0, "pillar_inputs": {
        "market_attractiveness": {"tam_usd": 2e5, "import_cagr_pct": -9.0,
                                  "gdp_per_capita_usd": 800.0,
                                  "saudi_share_pct": 19.0},
        "competition_intensity": {"hhi": 6000.0, "top_supplier_share_pct": 80.0,
                                  "named_company_count": 1},
        "regulatory_fit": {"tariff_applied_pct": 28.0,
                           "entry_requirements_count": 1,
                           "eligibility_gate": False},
        "profitability": {}, "risk": {}}}
    d = D.decide(poor)
    assert d["verdict"] == "NO-GO"
    assert not re.search(r"\b0\.\d+\b", d["counter_case"]["rebuttal"])


def test_a_pillar_has_exactly_one_arabic_name_across_surfaces():
    """خريطتا أسماءٍ متباعدتان = اسمان للعمود نفسه في تقريرَي دراسةٍ واحدة."""
    src = io.open(os.path.join(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))), "silk_reports.py"), encoding="utf-8").read()
    assert "_PILLAR_AR = {" not in src, "عادت خريطةُ أسماءِ الأعمدة الثانية"
    assert silk_reports._pillar_ar("competition") == D.pillar_label(
        "competition", "ar")


def test_internal_keys_are_gone_from_the_operator_surfaces_too():
    """LESSONS ١٠١ على الموجة نفسِها: عدَّ المواضع قبل الإصلاح.

    جدولُ الأعمدة على سطح المشغّل (`/analyze`) كان يطبع مفاتيحَ المكوّنات
    الخام في عمود «مكوّنات غائبة» — نفسُ التسرّب الذي عولج في `conditions`
    وحدَها. الفحصُ على المُصيِّر مباشرةً: مسارُ `/research` لا يمرّ بهذا
    الجدول، فاختبارُه من هناك يكون **خاوياً** (وكان كذلك أوّلَ مرّة).
    """
    from docx import Document
    ed = D.decide({"coverage": 0.6, "pillar_inputs": {
        "market_attractiveness": {"tam_usd": 2e8},
        "competition_intensity": {"hhi": 1200},
        "regulatory_fit": {"entry_requirements_count": 5},
        "profitability": {}, "risk": {}}})
    ed = {**ed, "schema": "silk.decision/v1"}
    doc = Document()
    silk_reports._docx_entry_decision(doc, {"entry_decision": ed})
    blob = "\n".join([p.text for p in doc.paragraphs]
                      + [c.text for t in doc.tables for r in t.rows
                         for c in r.cells])
    assert "جاذبية السوق" in blob, "الجدولُ لم يُبنَ — الفحصُ خاوٍ"
    leaked = [k for k in _INTERNAL_KEYS if k in blob]
    assert not leaked, f"معرّفاتٌ داخلية في docx المشغّل: {leaked}"
    assert "دخل الفرد" in blob, "المكوّناتُ الغائبة لم تُعرَّب"


# ── الهدف ٤ · جذرٌ عارٍ لا يُنشِط فرعاً ──────────────────────────────────

@pytest.mark.parametrize("row", [
    {}, {"components": {}}, {"regulatory": {}}, {"decision": {}},
    {"components": {}, "regulatory": {}, "decision": {}},
])
def test_a_bare_market_root_activates_no_dependent_branch(row):
    """وجودُ كائن السوق ليس دليلَ اكتمالِه — لكلّ فرعٍ شرطُه المستقلّ."""
    r = netherlands_research_blob()
    r["markets"] = [{"iso3": "NLD", "name_ar": "هولندا", "rank": 1,
                     "deep": True, **row}]
    v = silk_render.build_view(r, "ar")
    top = v["markets"][0]
    assert not top.get("components_detail")
    assert top.get("has_competitive_position") is False
    assert (v.get("completeness") or {}).get("pct") == 0.0
    assert not (v["decision"].get("verdict") or "")
    assert "غيابُ قياسٍ لا نتيجةَ قياس" in (v["decision"].get("why") or ""), (
        "صفرُ قياسٍ يُعرَض «تغطية 0/0 وفجوات: لا شيء» — طمأنةٌ بلا قياس")


# ── الهدف ٦ · S-01 قفلاً دائماً ──────────────────────────────────────────

def _vintages(year: "int | None") -> list:
    reg = regulatory_state("NLD", "080410")
    r = netherlands_research_blob()
    comps = DP.build_components(r["deep_research"])
    if year is not None:
        for c in comps.values():
            if isinstance(c, dict) and c.get("value") is not None:
                c["data_year"] = year
    r["markets"] = [{"iso3": "NLD", "name_ar": "هولندا", "rank": 1,
                     "deep": True, "regulatory": reg, "components": comps}]
    det = silk_render.build_view(r, "ar")["markets"][0]["components_detail"]
    return [d.get("vintage") or "" for d in det]


def test_a_missing_data_year_yields_no_invented_vintage_warning():
    """غيابُ السنة ≠ بيانٌ حديث — ولا تُختلَق سنةٌ لإصدار تحذير."""
    assert not any(_vintages(None))


def test_a_recent_data_year_is_not_flagged_stale():
    assert not any(_vintages(2025))


def test_a_stale_data_year_is_flagged_with_its_year():
    out = [v for v in _vintages(2019) if v]
    assert out, "تحذيرُ القِدَم خامدٌ على بيانات 2019"
    assert all("2019" in v and "أقدم من 3 سنوات" in v for v in out)


def test_a_very_stale_data_year_escalates_to_the_second_tier():
    out = [v for v in _vintages(2016) if v]
    assert out and all("أقدم من 7 سنوات" in v and "سياقٌ فقط" in v
                       for v in out), out


# ── حالاتٌ خصومية على المحرّك ────────────────────────────────────────────

def test_an_open_eligibility_gate_can_never_produce_an_unconditional_go():
    """أفضلُ مدخلاتٍ ممكنة + بوّابةٌ مفتوحة ⇒ لا `GO` بحال."""
    best = {"coverage": 1.0, "pillar_inputs": {
        "market_attractiveness": {"tam_usd": 5e10, "import_cagr_pct": 40.0,
                                  "gdp_per_capita_usd": 90_000,
                                  "saudi_share_pct": 0.1},
        "competition_intensity": {"hhi": 100, "top_supplier_share_pct": 2.0,
                                  "named_company_count": 12},
        "regulatory_fit": {"tariff_applied_pct": 0.0,
                           "entry_requirements_count": 8,
                           "eligibility_gate": True},
        "profitability": {"border_unit_value_usd_kg": 20.0,
                          "saudi_border_unit_value_usd_kg": 2.0,
                          "margin_at_border_pct": 80.0},
        "risk": {"political_stability_wgi": 2.5, "fx_volatility_pct": 0.1,
                 "supplier_concentration_hhi": 100, "critical_risk": False}}}
    d = D.decide(best)
    assert d["verdict"] != "GO"
    assert "أهلية" in d["conditions"][0]
    assert d["pillars"]["regulatory"]["value"] <= 0.3
    clear = copy.deepcopy(best)
    clear["pillar_inputs"]["regulatory_fit"]["eligibility_gate"] = False
    assert D.decide(clear)["verdict"] == "GO", "القيدُ صار عامّاً لا خاصّاً"


def test_confidence_never_leaves_the_range_a_reader_can_read():
    """حارسٌ لا تغييرُ سلوك: المدى الحقيقيّ يبقى كما هو حرفياً."""
    b = _bundle(200.0, 4.0, 9)
    inside = D.decide({**b, "coverage": 0.8})["confidence"]
    assert inside == 0.8
    assert D.decide({**b, "coverage": 3.0})["confidence"] == 1.0
    assert D.decide({**b, "coverage": -1.0})["confidence"] == 0.0


def test_zero_evidence_never_produces_a_score_that_looks_computed():
    d = D.decide({"coverage": 0.0, "pillar_inputs": {
        k: {} for k in ("market_attractiveness", "competition_intensity",
                        "regulatory_fit", "profitability", "risk")}})
    # البند 2 (أمر إصلاح المحرّك): صفرُ أعمدةٍ لم يعد يُصدر حكماً آلياً —
    # NO-GO من لا-بيانات كان اختلاقَ حكم؛ الآن حكمٌ فارغ (لا يُرقّى ولا
    # يُعرَض توصيةً) وثقةٌ None لا 0.0 «تبدو محسوبة».
    assert d["score"] is None and d["confidence"] is None
    assert d["verdict"] == "" and "غير كافية" in d["why"]
    assert d["insufficient_pillars"] is True


def test_a_malformed_bundle_degrades_declaredly_not_silently():
    """مدخلٌ مشوّه لا يُنتِج قراراً يبدو محسوباً — يسقط مُعلَناً."""
    for bad in ({}, {"coverage": 1.0}, {"coverage": 1.0, "pillar_inputs": {}}):
        d = D.decide(bad)
        # البند 2: لا حكمَ آلياً من مدخلٍ مشوّه — سقوطٌ معلَن بلا توصية.
        assert d["score"] is None and d["verdict"] == ""
        assert d["insufficient_pillars"] is True
    # وبعثةٌ مشوّهة ⇒ لا قرارٌ **مختلَق**: إمّا `None`، وإمّا قرارٌ كلُّ
    # أعمدته مُعلَنةُ الغياب بدرجةٍ لا تبدو محسوبة.
    out = DP.decide_for_deep({"missions": {"trade_flow": "ليس قاموساً"}})
    if out is not None:
        assert out["score"] is None or out["confidence"] == 0.0, out
