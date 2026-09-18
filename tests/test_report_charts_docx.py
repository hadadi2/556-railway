"""الموجة الخامسة — الرسمُ يصل مُسلَّم العميل (Word ⇒ PDF) من العقد نفسِه.

Wave 5: the chart contract the dashboard draws also reaches the client
deliverable as images (`silk_chart_image.chart_png` → `doc.add_picture`), so
Word and PDF can never show a different picture than the screen. Rules: no
data ⇒ no image (the gap stays in the report's limits section only), a failed
raster drops the image **and** its caption (never an empty frame, never a
failed export), and every image carries a readable attribution line that the
client-vocabulary guard reads like any other paragraph.

المالك: «الويب + PDF/Word معاً»، و«يختفي الرسم كلياً» عند غياب البيانات.
"""
import inspect
import os
import re
import sys

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_TOOLS = os.path.join(_ROOT, "tools")
if _TOOLS not in sys.path:
    sys.path.insert(0, _TOOLS)

import silk_chart_image as CI              # noqa: E402
import silk_render as R                    # noqa: E402
import silk_reports as SR                  # noqa: E402

pytest.importorskip("docx")
pytest.importorskip("pymupdf")

_FLAGS = ("SILK_CLIENT_METRIC_PRIVACY", "SILK_IMPORTS_SPOTLIGHT",
          "SILK_REPORT_CHARTS", "SILK_CONFIDENCE_DISCIPLINE",
          "SILK_REPORT_CHARTS_DOCX_ENABLED")

_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def _blob(module: str = "canonical_libya_tahini"):
    """مدوّنةٌ قانونيةٌ محفوظة — الافتراضيةُ تحمل التركّزَ والاقتصادَ الكامل.

    ومسارُ الواردات يحتاج سنتين مرصودتين (لا مسارَ بسنةٍ واحدة)، وأكثرُ
    المدوّنات المحفوظة بسنةٍ واحدة — فمدوّنةُ اليمن للألبان هي مقياسُ رسوم
    قسم السوق: ثلاثُ سنواتٍ حقيقيةٍ كما تجلبها أداةُ الإنتاج.
    """
    import importlib
    M = importlib.import_module(module)
    fn = [o for n, o in vars(M).items()
          if callable(o) and not n.startswith("_")
          and getattr(o, "__module__", "") == M.__name__
          and not inspect.signature(o).parameters][0]
    return fn()


_MARKET_BLOB = "canonical_nadec_yemen_dairy"


def _view(monkeypatch, *, charts: bool, lang: str = "ar",
          module: str = "canonical_libya_tahini", **extra):
    for f in _FLAGS:
        monkeypatch.delenv(f, raising=False)
    if charts:
        monkeypatch.setenv("SILK_REPORT_CHARTS", "1")
        monkeypatch.setenv("SILK_IMPORTS_SPOTLIGHT", "1")
    for k, v in extra.items():
        monkeypatch.setenv(k, v)
    return R.build_view(_blob(module), lang=lang)


def _docx(view, tmp_path, name="client.docx"):
    from docx import Document
    path = SR.render_client_docx(view, str(tmp_path / name))
    return Document(path), path


# ── (١) الصورةُ في المستند مقابلَ كلّ رسمٍ في العرض ────────────────────────

def test_every_chart_in_the_view_becomes_one_picture_in_the_docx(monkeypatch,
                                                                 tmp_path):
    view = _view(monkeypatch, charts=True)
    charts = view["deep_research"]["charts"]
    assert len(charts) >= 3, [c["id"] for c in charts]
    doc, _ = _docx(view, tmp_path)
    logo = 1 if os.path.exists((SR._load_branding() or {}).get("logo_path")
                               or "") else 0
    assert len(doc.inline_shapes) == len(charts) + logo
    text = "\n".join(p.text for p in doc.paragraphs)
    for ch in charts:
        # سطرُ الإسناد بعنوان الرسم ومصدره — رقمٌ بلا مصدرٍ مستحيلٌ هنا أيضاً.
        head = ch["title"].split("(")[0].strip()
        assert head in text, head


def test_no_flag_means_no_picture_and_an_identical_document(monkeypatch,
                                                            tmp_path):
    off = _view(monkeypatch, charts=False)
    doc_off, _ = _docx(off, tmp_path, "off.docx")
    assert "charts" not in off["deep_research"]
    logo = 1 if os.path.exists((SR._load_branding() or {}).get("logo_path")
                               or "") else 0
    assert len(doc_off.inline_shapes) == logo


def test_the_docx_valve_removes_images_while_the_web_keeps_them(monkeypatch,
                                                                tmp_path):
    """صمّامُ الصور وحدَه: العرضُ يبقى حاملاً رسومَه (الويب يرسم) والمستندُ
    يخرج بلا صورة — تراجعٌ جزئيّ بلا إطفاء الراية كلها."""
    view = _view(monkeypatch, charts=True,
                 SILK_REPORT_CHARTS_DOCX_ENABLED="0")
    assert view["deep_research"]["charts"], "الويب يرسم كما هو"
    doc, _ = _docx(view, tmp_path, "valve.docx")
    logo = 1 if os.path.exists((SR._load_branding() or {}).get("logo_path")
                               or "") else 0
    assert len(doc.inline_shapes) == logo


def test_a_failed_raster_drops_the_image_and_its_caption_not_the_export(
        monkeypatch, tmp_path):
    view = _view(monkeypatch, charts=True)
    titles = [c["title"] for c in view["deep_research"]["charts"]]
    monkeypatch.setattr(CI, "chart_png", lambda ch, lang="ar": None)
    doc, _ = _docx(view, tmp_path, "noimg.docx")     # لا استثناء
    logo = 1 if os.path.exists((SR._load_branding() or {}).get("logo_path")
                               or "") else 0
    assert len(doc.inline_shapes) == logo
    body = "\n".join(p.text for p in doc.paragraphs)
    for t in titles:
        assert t not in body, "تعليقٌ بلا صورةٍ = هيكلٌ فارغ"


# ── (٢) موضعُ الصورة: قسمُها لا آخرُ المستند ───────────────────────────────

def _doc_sequence(doc) -> list:
    """تسلسلُ المستند: ("img","") لكلّ صورة و("txt", النص) لكلّ فقرةٍ غيرِ خالية."""
    from docx.text.paragraph import Paragraph
    blip = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    seq = []
    for child in doc.element.body.iterchildren():
        if not child.tag.endswith("}p"):
            continue
        if child.findall(".//" + blip):
            seq.append(("img", ""))
        else:
            txt = Paragraph(child, doc).text.strip()
            if txt:
                seq.append(("txt", txt))
    return seq


def test_pictures_sit_inside_their_own_section(monkeypatch, tmp_path):
    """موضعُ الصورة قسمُها لا آخرُ المستند: رسمُ السوق بعد عنوان «السوق
    بالأرقام»، ورسومُ الاقتصاد بعد عنوان قسم الاقتصاد. القسمانِ في مدوّنتين
    مختلفتين (لا مدوّنةَ واحدةَ تحمل الاثنين: المسارُ يحتاج سنتين)."""
    import silk_i18n
    eco_head = silk_i18n.t("eco_heading", "ar")
    seq = _doc_sequence(_docx(_view(monkeypatch, charts=True),
                              tmp_path, "placed.docx")[0])
    imgs = [i for i, (k, _t) in enumerate(seq) if k == "img"]
    assert len(imgs) >= 3, seq[:5]
    eco = next(i for i, (k, t) in enumerate(seq)
               if k == "txt" and t == eco_head)
    assert max(imgs) > eco, "رسومُ الاقتصاد داخل قسم الاقتصاد"
    comp = next(i for i, (k, t) in enumerate(seq)
                if k == "txt" and t == "المنافسة والتسعير والهامش")
    assert min(imgs) > comp, "رسمُ التركّز داخل قسم المنافسة لا قبله"
    assert min(imgs) < eco, "ولا يتأخّر إلى قسم الاقتصاد"
    for i in imgs:                # تعليقُ كلّ صورةٍ يليها مباشرةً
        assert seq[i + 1][0] == "txt" and seq[i + 1][1], seq[i:i + 2]

    seq2 = _doc_sequence(_docx(_view(monkeypatch, charts=True,
                                     module=_MARKET_BLOB),
                               tmp_path, "placed_market.docx")[0])
    imgs2 = [i for i, (k, _t) in enumerate(seq2) if k == "img"]
    market = next(i for i, (k, t) in enumerate(seq2)
                  if k == "txt" and t == "السوق بالأرقام")
    assert imgs2 and min(imgs2) > market, "رسمُ السوق داخل قسمه لا قبله"
    for i in imgs2:
        assert seq2[i + 1][0] == "txt" and seq2[i + 1][1], seq2[i:i + 2]


# ── (٣) الصورةُ نفسُها: PNG لكلّ نوع، ولا صورةَ لسلسلةٍ فارغة ───────────────

def _chart(kind="bars", **over):
    base = {"id": "t", "kind": kind, "unit": "USD", "section": "economics",
            "title": "سلّم التكلفة من المصنع إلى الرف (USD)",
            "series": [{"label": "سعر المصنع", "value": 3.1},
                       {"label": "الشحن", "value": 0.37, "muted": True}],
            "source": "محسوب", "year": "", "note": "ملاحظة"}
    base.update(over)
    return base


def test_chart_png_renders_each_kind_and_refuses_empty_series():
    bars = CI.chart_png(_chart())
    rng = CI.chart_png(_chart(
        "range", unit="SAR",
        series=[{"label": "كلفة الدخول", "low": 80000, "high": 140000,
                 "value": 110000},
                {"label": "أقصى خسارة", "low": 120000, "high": 120000,
                 "value": 120000}]))
    gauge = CI.chart_png(_chart(
        "gauge", unit="index", value=2100.0, band="moderate",
        band_label="متوسطة التركّز",
        bands=[{"key": "open", "label": "مفتوحة", "from": 0, "to": 1500},
               {"key": "moderate", "label": "متوسطة", "from": 1500,
                "to": 2500},
               {"key": "high", "label": "مركّزة", "from": 2500, "to": 10000}],
        series=[{"label": "متوسطة التركّز", "value": 2100.0}]))
    for png in (bars, rng, gauge):
        assert png and png.startswith(_PNG_MAGIC)
    assert CI.chart_png(_chart(series=[])) is None
    assert CI.chart_png({"id": "x", "kind": "bars"}) is None
    assert CI.chart_png(None) is None


def test_chart_png_marker_lands_on_the_measured_value():
    """هندسةٌ مقيسة لا مظهرٌ مُدّعى: مؤشّرُ التركّز عند موضع القيمة نفسِها."""
    import pymupdf
    value = 940.0
    png = CI.chart_png(_chart(
        "gauge", unit="index", value=value, band="open", band_label="مفتوحة",
        bands=[{"key": "open", "label": "مفتوحة", "from": 0, "to": 1500},
               {"key": "moderate", "label": "متوسطة", "from": 1500,
                "to": 2500},
               {"key": "high", "label": "مركّزة", "from": 2500, "to": 10000}],
        series=[{"label": "مفتوحة", "value": value}]))
    pix = pymupdf.open(stream=png, filetype="png")[0].get_pixmap()
    gold = [x for y in range(pix.height) for x in range(pix.width)
            if (lambda i: abs(pix.samples[i] - 0xC9) < 16
                and abs(pix.samples[i + 1] - 0xA2) < 16
                and abs(pix.samples[i + 2] - 0x27) < 16)(
                    (y * pix.width + x) * pix.n)]
    assert gold, "لا مؤشّر مرسوم"
    scale = pix.width / CI._W
    bar_w = CI._W - CI._GUT - CI._VAL - 16.0
    expected = (CI._VAL + 8.0 + bar_w * (1 - value / 10000.0)) * scale
    assert abs(sum(gold) / len(gold) - expected) <= 6, (gold[:3], expected)


def test_chart_png_prints_a_dash_for_a_missing_value():
    """قيمةٌ غائبة «—» بلا عمود — لا صفرَ مرسوم."""
    import pymupdf
    png = CI.chart_png(_chart(series=[{"label": "أ", "value": None},
                                      {"label": "ب", "value": 5.0}]))
    txt = pymupdf.open(stream=png, filetype="png")   # صورةٌ لا نصّ: نقيس الحبر
    assert png and txt.page_count == 1
    assert CI._num(_chart(), None) == "—"
    assert CI._num(_chart(unit="index"), 2100.4) == "2100"
    assert CI._num(_chart(unit="%"), 41.24) == "41.2%"
    assert CI._num(_chart(unit="USD"), 3.1) == "$3.1"
    assert CI._num(_chart(unit="USD"), 41300000) == "$41.3M"


def test_caption_carries_title_source_year_and_note():
    cap = CI.caption(_chart(source="UN Comtrade", year="2024",
                            note="حصة السعودية مميَّزة بالذهبي."))
    assert cap.split(" · ")[0].startswith("سلّم التكلفة")
    assert "UN Comtrade" in cap and "2024" in cap and "الذهبي" in cap
    assert CI.caption({"title": "", "source": "", "year": "", "note": ""}) == ""


# ── (٤) لغةُ التقرير: تسميةٌ عربيةٌ لا تُرسَم على مستندٍ إنجليزيّ ───────────

def test_an_english_report_draws_only_english_labelled_charts(monkeypatch,
                                                              tmp_path):
    """الفصلُ الصلب لا الترجمة: تسميةٌ عربيةٌ يكتبها المحرّك (أسماءُ أرقام
    القرار، درجاتُ السلّم) لا تُرسَم على تقريرٍ إنجليزيّ — تُسقَط كما يُسقِط
    `_client_decision_numbers_table` جدولَه.

    حدُّ الدليل: المدوّنةُ المخزَّنة تحمل نثرَ حكمٍ عربياً فبوّابةُ اللغة ترفض
    مستندَها الإنجليزيَّ كلَّه (عيبٌ سابقٌ لهذه الموجة ولا صلةَ له بالرسوم)،
    فيُقاس الإدراجُ هنا على الدالّة نفسِها بمستندٍ نظيف.
    """
    from docx import Document
    view = _view(monkeypatch, charts=True, lang="en", module=_MARKET_BLOB)
    charts = view["deep_research"]["charts"]
    ar = re.compile(r"[\u0600-\u06FF]")
    for ch in charts:
        for row in ch["series"]:
            assert not ar.search(str(row["label"])), (ch["id"], row["label"])
    # وعلى الأقلّ رسمٌ إنجليزيُّ التسميات يصل المستند فعلاً.
    en_market = [c for c in charts if c["section"] == "market"]
    assert en_market, [c["id"] for c in charts]
    doc = Document()
    SR._client_charts(doc, view["deep_research"], "market", "en")
    assert len(doc.inline_shapes) == len(en_market)
    body = "\n".join(p.text for p in doc.paragraphs)
    # ولا حرفَ عربيٍّ في تعليقٍ إنجليزيّ — حتى الفاصلةُ العربية «،» في قائمة
    # سنوات: حرفٌ عربيٌّ تُسقِط به بوّابةُ اللغة المستندَ كلَّه (§58).
    assert not ar.search(body), body
    assert "، " not in body


# ── (٥) الحارسُ يقرأ التعليقات كسائر الفقرات (لا مفردةَ قياسٍ داخليّ) ───────

def test_captions_pass_the_client_vocabulary_guard(monkeypatch, tmp_path):
    import silk_quality_gate as Q
    view = _view(monkeypatch, charts=True, SILK_CLIENT_METRIC_PRIVACY="1")
    doc, _ = _docx(view, tmp_path, "guarded.docx")   # `_client_assert_clean`
    blob = "\n".join([p.text for p in doc.paragraphs]
                     + [c.text for t in doc.tables for r in t.rows
                        for c in r.cells])
    assert SR._client_forbidden_hits(blob, "ar") == []
    assert Q.run_client_artifact_text_gate(blob, lang="ar", view=view) == []


def test_a_chart_whose_caption_is_emptied_is_not_drawn_at_all(monkeypatch,
                                                              tmp_path):
    """صورةٌ تحمل أرقاماً لا يقرؤها حارسُ النصّ — فإسنادُها شرطُ رسمها.
    تعليقٌ يُفرَّغ بالتطهير أو بالفصل اللغويّ ⇒ لا صورة (لا رقمَ بلا مصدر)."""
    from docx import Document
    view = _view(monkeypatch, charts=True)
    monkeypatch.setattr(CI, "caption", lambda ch: "")
    doc = Document()
    SR._client_charts(doc, view["deep_research"], "market", "ar")
    assert len(doc.inline_shapes) == 0
    assert not [p.text for p in doc.paragraphs if p.text.strip()]


# ── (٦) حصادُ المراجعة الذاتية §58: النصُّ المرسومُ والحبرُ المقيس ──────────

def test_text_drawn_inside_the_image_passes_the_client_guard(monkeypatch):
    """ثغرةٌ بنيوية: حارسُ مفردات العميل يقرأ الفقرات ولا يقرأ الصورة. فتسميةٌ
    من نصّ بعثةٍ («[الطلب] مبني على: mean interest 0-100…» — شكلُ الإنتاج
    حرفياً) تصل مُسلَّمَ العميل مرسومةً وهي عينُ ما ترفضه مكتوبة. فالتطهيرُ
    قبل الرسم، وما بقي ممنوعاً ⇒ لا رسمَ (§58، M1)."""
    from docx import Document
    scaffold = "[الطلب] مبني على: mean interest 0-100 for 'tahini' geo=LY"
    assert SR._client_forbidden_hits(scaffold, "ar"), "المرجعُ: البوّابةُ ترفضه"
    ch = {"id": "demand_interest", "kind": "bars", "unit": "index",
          "section": "market", "title": "اهتمام البحث النسبي بالمنتج",
          "series": [{"label": scaffold, "value": 74.0},
                     {"label": "طحينة", "value": 100.0}],
          "source": "Google Trends", "year": "2026", "note": "مؤشّر نسبي"}
    assert SR._client_chart_text_safe(ch, "ar") is None
    doc = Document()
    SR._client_charts(doc, {"charts": [ch]}, "market", "ar")
    assert len(doc.inline_shapes) == 0, "لا صورةَ لنصٍّ لا يعبر الحارس"
    assert not [p.text for p in doc.paragraphs if p.text.strip()]
    # ورسمٌ نظيفُ النصّ يُرسَم، ونصُّه المرسومُ هو النصُّ المطهَّر لا الخام.
    ok = dict(ch, series=[{"label": "طحينة", "value": 74.0},
                          {"label": "زيت", "value": 100.0}])
    safe = SR._client_chart_text_safe(ok, "ar")
    assert safe and [r["label"] for r in safe["series"]] == ["طحينة", "زيت"]
    doc2 = Document()
    SR._client_charts(doc2, {"charts": [ok]}, "market", "ar")
    assert len(doc2.inline_shapes) == 1


def test_a_label_emptied_by_the_sanitizer_drops_the_whole_chart(monkeypatch):
    """عمودٌ بلا تسميةٍ لا يُقرأ — فتسميةٌ يُفرِّغها التطهيرُ تُسقِط الرسمَ."""
    monkeypatch.setattr(SR, "_client_sanitize",
                        lambda t, lang="ar": "" if t == "طحينة" else str(t))
    ch = {"id": "x", "kind": "bars", "unit": "%", "section": "market",
          "title": "عنوان", "series": [{"label": "طحينة", "value": 5.0},
                                       {"label": "زيت", "value": 6.0}],
          "source": "م", "year": "2024", "note": "ن"}
    assert SR._client_chart_text_safe(ch, "ar") is None


def test_a_missing_value_draws_no_bar_measured_in_ink():
    """«قيمةٌ غائبة «—» بلا عمود» تُقاس حبراً لا تُدّعى: صفُّ القيمة الغائبة
    بلا بكسلٍ أزرقَ واحد، وصفُّ القيمة الحاضرة مليءٌ به (§58، M11)."""
    import pymupdf
    png = CI.chart_png(_chart(series=[{"label": "أ", "value": None},
                                      {"label": "ب", "value": 5.0}]))
    pix = pymupdf.open(stream=png, filetype="png")[0].get_pixmap()
    k = pix.height / (26.0 + 2 * CI._ROW + 6.0)          # بكسل لكلّ نقطة

    def blue(row_index: int) -> int:
        y0 = int((26.0 + row_index * CI._ROW + 2) * k)
        y1 = int((26.0 + row_index * CI._ROW + CI._BAR_H - 2) * k)
        n = 0
        for y in range(y0, y1):
            for x in range(pix.width):
                i = (y * pix.width + x) * pix.n
                if (abs(pix.samples[i] - 0x25) < 24
                        and abs(pix.samples[i + 1] - 0x63) < 24
                        and abs(pix.samples[i + 2] - 0xEB) < 24):
                    n += 1
        return n

    assert blue(0) == 0, "لا عمودَ لقيمةٍ غائبة"
    assert blue(1) > 100, "والقيمةُ الحاضرة عمودٌ مرسوم"


def test_the_image_renderer_uses_only_branding_colours():
    """تدقيقٌ قبل الدمج: اختبارُ الألوان كان يفحص `web/platform.html` وحدَها،
    ومُصيِّرُ الصورة يطبع النصَّ بلونٍ سادسٍ (رقمان مُبدَّلان عن حبر الهوية).
    السطحان يقرآن العقدَ نفسَه فليقرآ ملفَّ الهوية نفسَه — خمسةٌ لا ستّة."""
    import re
    allowed = {"#2563EB", "#C9A227", "#EEF2F7", "#64748B", "#111827"}
    src = open(CI.__file__, encoding="utf-8").read()
    hexes = {h.upper() for h in re.findall(r"#[0-9A-Fa-f]{6}", src)}
    assert hexes <= allowed, hexes - allowed
    # والمُتَّجهاتُ (للأشكال) هي الأربعةُ نفسُها بصيغة 0–1.
    tuples = {(f"#{r}{g}{b}").upper() for r, g, b in re.findall(
        r"0x([0-9A-Fa-f]{2}) / 255, 0x([0-9A-Fa-f]{2}) / 255, "
        r"0x([0-9A-Fa-f]{2}) / 255", src)}
    assert tuples <= allowed, tuples - allowed
    assert len(tuples) >= 4, tuples
    # وحبرُ النصّ نصٌّ (CSS) لا مُتَّجه: النصُّ يُرسَم بـ`insert_htmlbox`.
    assert CI._INK == "#111827" and isinstance(CI._INK, str)
