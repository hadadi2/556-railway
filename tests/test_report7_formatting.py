"""تقرير سِلك ٧ — المراجع والإخراج وبيانات الإصدار (§7، الدرس ٢٧٨).

references carry a light per-source note of the facts each source supports;
every client table repeats its header on each page and keeps a row from
splitting; the contacts table gives name/address wider columns and renders
phone/email/website left-to-right inside the RTL table; the report stamps the
decision-rules version (visible) and an input fingerprint (file properties).

كلُّ الحالات مصطنعةٌ لاختبارٍ هرمتيّ.
"""
from docx import Document
from docx.oxml.ns import qn

import silk_reports as SR


def _table(doc):
    return doc.tables[-1]


def test_every_table_repeats_its_header_and_rows_do_not_split():
    doc = Document()
    SR._add_table(doc, ["أ", "ب"], [["1", "2"], ["3", "4"]])
    t = _table(doc)
    first = t.rows[0]._tr.trPr
    assert first.find(qn("w:tblHeader")) is not None
    assert all(r._tr.trPr.find(qn("w:cantSplit")) is not None for r in t.rows)


def test_contact_cells_are_left_to_right_and_columns_weighted():
    doc = Document()
    dr = {"importer_leads": {"leads": [
        {"name": "Kopi Hub", "category": "Coffee wholesaler",
         "phone": "+60 3 2222 2222", "email": "sales@kopihub.my",
         "website": "https://kopihub.my", "address": "Kuala Lumpur, Malaysia"}]},
        "market": {"iso3": "MYS", "name_en": "Malaysia", "name_ar": "ماليزيا"},
        "product": "قهوة محمصة", "hs_code": "090121"}
    SR._docx_leads(doc, dr, lang="ar")
    t = doc.tables[0]   # الجدولُ الرئيس — تحته جدولُ القناة والخطوة (الدرس ٢٨٢)
    row = t.rows[1]
    for ci in (2, 3, 4):
        p = row.cells[ci].paragraphs[0]
        assert p._p.pPr.find(qn("w:bidi")).get(qn("w:val")) == "0"
        assert all(r._r.rPr.find(qn("w:rtl")).get(qn("w:val")) == "0"
                   for r in p.runs)
    # الاسمُ أعرضُ من التقييم.
    assert t.columns[0].cells[0].width > t.columns[5].cells[0].width


def test_references_name_the_facts_each_source_supports():
    doc = Document()
    dr = {"missions": {"trade_flow": {"findings": [
        {"value": 89400000.0, "source": "UN Comtrade", "confidence": 0.9,
         "note": "HS090121 إجمالي استيراد 2024, USD",
         "retrieved_at": "2026-09-17"}]}}}
    ledger = {"entries": {
        "market_imports_usd": {
            "value": 89400000.0, "source": "UN Comtrade", "status": "observed",
            "label_ar": "واردات السوق", "label_en": "market imports"},
        # استنتاجٌ بمصدرٍ عامّ لا يُعرَض حقيقةً يسندها المصدر.
        "supplier_nature": {
            "value": "مصنّعون", "source": "UN Comtrade", "status": "inference",
            "label_ar": "طبيعة المورّدين", "label_en": "supplier nature"},
        # «World Bank WITS» لا يُنسَب إلى «World Bank».
        "tariff_applied_pct": {
            "value": 5.0, "source": "World Bank WITS", "status": "observed",
            "label_ar": "التعرفة المطبّقة", "label_en": "applied tariff"}}}
    SR._client_references_section(doc, dr, "ar", ledger=ledger)
    lines = [p.text for p in doc.paragraphs
             if p.text.startswith("UN Comtrade — ")]
    assert lines and "يسند: واردات السوق" in lines[0]
    assert "طبيعة المورّدين" not in lines[0]
    sup = SR._supported_facts(ledger, "ar")
    assert "world bank" not in sup and "world bank wits" in sup


def test_rules_version_is_shown_and_fingerprint_is_in_file_properties():
    doc = Document()
    view = {"report_meta": {"study_id": 7}, "hs_code": "090121",
            "header": {"target_market": "ماليزيا", "product": "قهوة"},
            "ledger": {"entries": {"hhi": {"value": 2100, "year": 2024,
                                            "source": "UN Comtrade"}}}}
    SR._stamp_report_metadata(doc, view, "ar")
    body = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "silk.decision/v1" in body
    comments = doc.core_properties.comments
    assert "input_fingerprint=" in comments and "rules_version=" in comments
    fp = SR._input_fingerprint(view)
    assert len(fp) == 12 and fp == SR._input_fingerprint(dict(view))
    changed = dict(view, ledger={"entries": {"hhi": {"value": 2200,
                                                      "year": 2024}}})
    assert SR._input_fingerprint(changed) != fp


def test_widths_reach_the_grid_and_long_rows_may_split():
    doc = Document()
    SR._add_table(doc, ["أ", "ب"], [["قصير", "x"], ["ن" * 900, "y"]],
                  widths=[3, 1])
    t = _table(doc)
    grid = t._tbl.tblGrid.findall(qn("w:gridCol"))
    assert int(grid[0].get(qn("w:w"))) > int(grid[1].get(qn("w:w")))
    assert t.rows[1]._tr.trPr.find(qn("w:cantSplit")) is not None
    assert t.rows[2]._tr.trPr.find(qn("w:cantSplit")) is None


def test_the_fingerprint_follows_the_card_and_never_breaks_the_export():
    base = {"hs_code": "090121", "header": {"target_market": "ماليزيا"},
            "ledger": {"entries": {}},
            "deep_research": {"economics": {"cost_currency": "SAR"}}}
    other = dict(base, deep_research={"economics": {"cost_currency": "USD"}})
    assert SR._input_fingerprint(base) != SR._input_fingerprint(other)
    doc = Document()
    SR._stamp_report_metadata(doc, {"ledger": {"entries": ["legacy"]},
                                    "hs_code": "090121"}, "ar")
