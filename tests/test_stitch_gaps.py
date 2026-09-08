"""سدّ الخياطات (أمر المالك «ابحث عن الفجوات وسدّها») — الأقسام الجديدة تصل
كل سطوح العرض فعلاً: تقرير العميل، لوحة الويب، وفحص الشكل الموصى به. هرمتي.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def test_client_docx_carries_reverse_solve_headline(tmp_path):
    import silk_render as R
    import silk_reports as REP
    from docx import Document
    from tools.canonical_netherlands import netherlands_research_blob
    path = str(tmp_path / "client.docx")
    REP.render_client_docx(R.build_view(netherlands_research_blob()), path)
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "أقصى سعر مصنع قابل للمنافسة" in text
    assert "افتراضات معلنة قابلة للتعديل" in text


def test_client_economics_absent_when_no_model_no_empty_skeleton(tmp_path):
    import silk_reports as REP

    class _Doc:
        def __init__(self):
            self.parts = []

        def add_heading(self, t, level=1):
            self.parts.append(t)

        def add_paragraph(self, t, style=None):
            self.parts.append(t)
    d = _Doc()
    REP._client_economics_section(d, {"economics": {"reverse_solve": None}})
    assert d.parts == []   # لا هيكل فارغ


def test_dashboard_renders_economics_and_gap_register():
    html = open(os.path.join(_ROOT, "web/index.html"), encoding="utf-8").read()
    assert "النموذج الاقتصادي — أقصى سعر مصنع قابل للمنافسة" in html
    assert "dr.economics" in html
    assert "dr.gap_register" in html
    assert "gap_class_label" in html


def test_anchor_form_mismatch_check_fires_only_with_declared_form():
    import silk_quality_gate as Q
    dr_conflict = {"economics": {
        "product_form": "حليب UHT طويل الأجل",
        "anchor_price": {"source": "سعر رف حليب طازج 1 لتر", "note": ""}}}
    f = Q._check_anchor_matches_recommended_form(dr_conflict)
    assert f and f[0]["repairable"] is True
    assert "anchor_form_mismatch" not in Q.FAIL_TRIGGER_CHECKS
    # بلا حقل الشكل = نائم (لا لغويات تخمينية).
    dr_dormant = {"economics": {
        "product_form": None,
        "anchor_price": {"source": "سعر رف حليب طازج"}}}
    assert Q._check_anchor_matches_recommended_form(dr_dormant) == []
    # شكل مطابق = صمت.
    dr_ok = {"economics": {
        "product_form": "UHT",
        "anchor_price": {"source": "سعر رف حليب UHT 1 لتر"}}}
    assert Q._check_anchor_matches_recommended_form(dr_ok) == []


def test_product_form_carried_from_product_card():
    import silk_economics as E
    eco = E.economics_view({"missions": {}},
                           product_card={"form": "UHT",
                                         "cost_per_unit": 2.0})
    assert eco["product_form"] == "UHT"
    assert E.economics_view({"missions": {}})["product_form"] is None


def test_directive_registry_includes_form_check():
    import silk_quality_gate as Q
    assert any(c == "anchor_form_mismatch"
               for c, _ in Q.DIRECTIVE_AUDIT_CHECKS.values())
