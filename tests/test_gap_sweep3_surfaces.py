"""موجة صيد الفجوات الثالثة — أقفال أسطح التسليم وطبقة API.

الصيادان ٣/٤ (2026-08-25): حقن عربي في docx الإنجليزي، حقول عرض ميتة
(price_rows/hs_derivation)، مختصر عربي على lang=en، بوابة نص المُنتَج تعود
501 على السطح الجذري، طلب استرجاع كلمة المرور بلا خانق، حرق سقف قبل الشرط.
هرمتي. Run: python3 -m pytest tests/test_gap_sweep3_surfaces.py -q
"""
import ast
import os
import re
import sys
from unittest.mock import patch

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "tools"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def _nl_view(lang="ar"):
    from canonical_netherlands import netherlands_research_blob
    import silk_render as R
    return R.build_view(netherlands_research_blob(), lang)


# ── C1: لا تعقيم عاري اللغة داخل مصيّر docx العميل ──────────────────────

def test_no_bare_client_sanitize_inside_client_docx_renderer():
    """`_client_sanitize` بلا `lang` داخل `render_client_docx` يطبّق
    الاستبدالات العربية على نثر إنجليزي («The overall التقييم is…») —
    المطهّر المربوط `_sanitize` هو المسموح."""
    src = open(os.path.join(_ROOT, "silk_reports.py"), encoding="utf-8").read()
    tree = ast.parse(src)
    fn = next(n for n in ast.walk(tree)
              if isinstance(n, ast.FunctionDef)
              and n.name == "render_client_docx")
    bad = []
    for node in ast.walk(fn):
        if (isinstance(node, ast.Call)
                and getattr(node.func, "id", "") == "_client_sanitize"):
            args = [ast.unparse(a) for a in node.args]
            kws = {k.arg for k in node.keywords}
            if len(args) < 2 and "lang" not in kws:
                bad.append((node.lineno, args))
    assert not bad, f"نداء تعقيم عاري اللغة داخل مصيّر العميل: {bad}"


# ── C4/C5/C6: لا عربية خام على مخرجات lang=en ───────────────────────────

def test_eco_terms_translate_for_english():
    import silk_reports as SR
    assert SR._eco_term("متوسط", "en") == "mid"
    assert SR._eco_term("كجم", "en") == "kg"
    assert SR._eco_term("متوسط", "ar") == "متوسط"


def test_view_brief_and_target_market_follow_report_language():
    v_ar = _nl_view("ar")
    v_en = _nl_view("en")
    assert v_ar["brief"][0].startswith("التوصية")
    assert v_en["brief"][0].startswith(("Recommendation", "Status"))
    assert "سوق" not in v_en["brief"][0]
    assert v_en["header"]["target_market"] == "Netherlands"
    assert v_ar["header"]["target_market"] == "هولندا"


def test_cover_wordmark_follows_language():
    import silk_reports as SR
    assert SR._SILK_WORDMARK["en"] == "Silk"
    src = open(os.path.join(_ROOT, "silk_reports.py"), encoding="utf-8").read()
    assert "_add_cover_wordmark(doc, branding, lang)" in src, \
        "غلاف docx العميل لا يمرر لغة التقرير للشعار"


# ── C7/C8/C9: الحقول المبنية تُعرض فعلاً (عائلة الدرس 104) ──────────────

def test_price_row_reason_consults_the_note_for_numeric_values():
    """قيمة رقمية («7.49») كانت تُوسم «الوزن غير متاح» رغم أن الوزن معلن في
    الملاحظة — التصنيف على القيمة ثم الملاحظة."""
    dr = _nl_view("ar")["deep_research"]
    rows = dr["price_rows"]
    assert rows, "لا صفوف أسعار في المدوّنة القانونية"
    r0 = rows[0]
    assert "note" in r0 and "store" not in r0   # المفتاح سُمّي بما هو
    assert "1 كجم" in r0["note"]
    assert r0["reason"] == "", r0


def test_client_docx_renders_price_rows_and_derivation(tmp_path):
    import silk_reports as SR
    from docx import Document
    v = _nl_view("ar")
    p = SR.render_client_docx(v, str(tmp_path / "c.docx"))
    txt = "\n".join(par.text for par in Document(p).paragraphs)
    assert "الأسعار المرصودة على الرف" in txt
    # قفل محدَّث معلن (#13 الدرس 170): سطر الفتح بلغة الزائر لا لغة المدخلات.
    assert "سعر المصنع وسعر المنافس بعملة ووحدة متطابقتين" in txt   # سطر الفتح الوحيد
    assert "كيف حُدّد الرمز الجمركي" in txt            # سجل الاشتقاق §3.1


def test_substitution_note_reaches_the_client_when_code_unconfirmed(tmp_path):
    """جملة إفصاح الاستبدال الإلزامية (§3.2) كانت تُبنى في العرض ولا يعرضها
    أي سطح — بوابة spec_evidence_coherence كانت تعدّ وجودها البنيوي عرضاً."""
    from canonical_yemen import yemen_research_blob
    import silk_render as R
    import silk_reports as SR
    from docx import Document
    v = R.build_view(yemen_research_blob(), "ar")
    deriv = (v["deep_research"].get("hs_derivation") or {})
    assert deriv.get("substitution_note"), "المدوّنة اليمنية بلا رمز معلَّم؟"
    p = SR.render_client_docx(v, str(tmp_path / "y.docx"))
    txt = "\n".join(par.text for par in Document(p).paragraphs)
    assert "لا تُنقل إلى فئة المنتج الفعلية" in txt
    # وشارة «التركّز سياق» (Wave 3.2) تصل القارئ أيضاً.
    assert v["deep_research"].get("concentration_context_only") is True
    assert "تُقرأ سياقاً عاماً للفئة" in txt


# ── C10/C14: قصّ معلن وتعقيم ملاحظة الحكم ───────────────────────────────

def test_dead_table_replacement_clips_at_word_boundary():
    import silk_render as R
    header = ("| مؤشر التغطية التفصيلية للسلسلة الزمنية رقم واحد "
              "| عمود المؤشرات التفصيلية رقم اثنين للسلسلة الزمنية "
              "| عمود ثالث | عمود رابع |")
    sep = "|---|---|---|---|"
    row = "| — | — | — | — |"
    out = R._collapse_dead_tables("\n".join([header, sep, row]))
    assert "غير متاح: جدول" in out
    assert "…" not in out
    line = next(l for l in out.splitlines() if "غير متاح: جدول" in l)
    m = re.search(r"«(.+?)»", line)
    assert m and ("(مختصر" in m.group(1)
                  or len(m.group(1)) <= 90), line


def test_verdict_note_is_stripped_of_internal_english_half():
    from canonical_netherlands import netherlands_research_blob
    import silk_render as R
    blob = netherlands_research_blob()
    blob["deep_research"]["verdict"]["note"] = (
        "أولية فقط؛ المصادر الناقصة معلنة لا مقدّرة. Preliminary only; "
        "missing sources flagged, not estimated.")
    v = R.build_view(blob, "ar")
    note = v["deep_research"]["verdict"].get("note") or ""
    assert "Preliminary only" not in note
    assert "أولية فقط" in note


# ── C2/C3: الواجهتان تقرآن حمولة الحجب ──────────────────────────────────

def test_platform_pdf_button_reads_the_gate_payload_not_raw_json():
    src = open(os.path.join(_ROOT, "web", "platform.html"),
               encoding="utf-8").read()
    assert 'window.open(API + "/studies/" + id + "/report.pdf"' not in src, \
        "زر PDF ما زال يفتح JSON الحجب خاماً في تبويب"
    assert "blocked_checks" in src
    assert "out.quality" in src, "ملخّص الجودة (G-02) ما زال بلا مستهلك"


def test_operator_download_toast_names_the_fail_drivers():
    src = open(os.path.join(_ROOT, "web", "index.html"),
               encoding="utf-8").read()
    assert "fail_drivers" in src, "اللوحة تسقط قائد الحجب من رسالة الفشل"
    # القصّ الخام في خلية ملخص البعثة استُبدل بالمعلن (clip140).
    assert 'esc((mm.summary||"").slice(0,140))' not in src
    assert "clip140(" in src


# ── D1: بوابة نص المُنتَج على السطح الجذري = 409 منظّم ──────────────────

def test_root_artifact_gate_refusal_is_409_not_501(monkeypatch, tmp_path):
    os.environ.pop("SILK_API_KEY", None)
    monkeypatch.setenv("SILK_DB", str(tmp_path / "t.db"))
    import importlib
    import api as api_mod
    importlib.reload(api_mod)
    from fastapi.testclient import TestClient
    import silk_storage
    import silk_reports
    from canonical_netherlands import netherlands_research_blob
    with patch("requests.get", side_effect=OSError("hermetic")), \
            patch("requests.post", side_effect=OSError("hermetic")):
        blob = netherlands_research_blob()
        aid = silk_storage.save_analysis(blob)
        app = api_mod.create_app()
        cl = TestClient(app)

        def _boom(view, path):
            raise silk_reports.ClientArtifactGateError(
                "رفض", findings=[{"check": "artifact_language_leak",
                                  "note": "تسرب"}])

        with patch.object(silk_reports, "render_client_docx",
                          side_effect=_boom), \
                patch.object(silk_reports, "render_academic_docx",
                             side_effect=_boom, create=True):
            r = cl.get(f"/analyses/{aid}/report.docx")
    assert r.status_code == 409, (r.status_code, r.text[:300])
    d = r.json()["detail"]
    assert d.get("error") == "quality_gate_fail"
    assert any(f.get("check") == "artifact_language_leak"
               for f in d.get("findings", [])), d


# ── D2: خانق استرجاع كلمة المرور ────────────────────────────────────────

def test_password_reset_request_is_throttled(monkeypatch):
    from platform_helpers import seed, client
    seed(monkeypatch)
    cl = client()
    codes = [cl.post("/platform/auth/password-reset/request",
                     json={"email": f"x{i}@example.com"}).status_code
             for i in range(7)]
    assert 429 in codes, codes
    assert codes[0] == 200


# ── D3: لا حرق سقف قبل شرط 409 ──────────────────────────────────────────

def test_regen_report_checks_checkpoints_before_reserving_cap():
    src = open(os.path.join(_ROOT, "api.py"), encoding="utf-8").read()
    seg = src[src.find("def regenerate_report"):]
    if "def regenerate_report" not in src:
        seg = src[src.find("no mission checkpoints stored") - 2000:]
    i_ck = seg.find("load_mission_checkpoints")
    i_cap = seg.find("_free_ai_extras_allowed")
    assert 0 < i_ck < i_cap, "حجز السقف ما زال يسبق شرط نقاط التفتيش"


# ── D4/D6: تنقيح الأسرار واسم صنف الاستثناء ─────────────────────────────

def test_redact_covers_the_serper_alias(monkeypatch):
    import importlib
    monkeypatch.setenv("SERPER_API_KEY", "sk-serper-SECRET-123")
    import silk_diagnostics as DG
    importlib.reload(DG)
    out = DG._redact("probe failed: key sk-serper-SECRET-123 rejected")
    assert "sk-serper-SECRET-123" not in out


def test_bridge_generic_failure_carries_no_exception_class_name():
    src = open(os.path.join(_ROOT, "silk_platform", "engine_bridge.py"),
               encoding="utf-8").read()
    assert 'f"{type(exc).__name__}: {str(exc)[:300]}"' not in src, \
        "run_error ما زال يحمل اسم صنف الاستثناء (رطانة على سطح مصنع)"
    assert "تعذّر إكمال الدراسة لعطل تقني" in src
