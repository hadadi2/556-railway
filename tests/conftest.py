"""أدوات اختبار مشتركة — shared test helpers (M0).

يوفّر `block_network` القانوني الواحد بدل النسخ المكرَّرة في ملفات الموجات
(الأثر التاريخي يُنظَّف في M9). الاختبارات الجديدة تستورد من هنا حصراً.
Canonical network guard for hermetic tests; new tests import from here only.
"""
import contextlib
import os
import socket
import sys

# كنسُ أيتام المنصّة مشغَّل افتراضياً في الإنتاج (إصلاح «قيد الإعداد» الأبدي)؛
# في العملية الهرمتية نطفئه صراحةً: خيطٌ يوفّق صفوفاً في قاعدةٍ مؤقّتة يتقاطع
# مع اختبارات التزامن ونوافذها الموسّعة. صريحٌ هنا لا مُكتشَفٌ في شيفرة الإنتاج.
os.environ.setdefault("SILK_PLATFORM_ORPHAN_SWEEP", "0")
# مُشرِف تشغيلات الدراسات (R2) — نفس المنطق: خيطُ نبضةٍ/كنسٍ/مطالبة دوريّ يتقاطع
# مع اختبارات التزامن؛ الاختبارات تستدعي `study_runtime.tick()` صراحةً بدله.
os.environ.setdefault("SILK_PLATFORM_RUN_SUPERVISOR", "0")
# R2b: حاصدُ `/research` الجذري الدوريّ — نفسُ المنطق: خيطٌ دوريّ يتقاطع مع اختبارات
# التزامن؛ الاختبارات تنادي `silk_research_runtime.tick()` صراحةً بدله.
os.environ.setdefault("SILK_ORPHAN_REAP_INTERVAL_S", "0")

# حدّ المعدل صفر أساسَ جلسةِ الاختبار صراحةً (دورة C4 لصيد الفجوات ٣):
# `api.app` المفردُ يخبز `_rl_max` عند أول استيرادٍ، وهويةُ `X-API-Key`
# المشتركة («secret») تتجاوز 120 طلباً في نافذة دقيقةٍ تحت حمل السلسلة
# الكاملة. كانت السلسلةُ كلُّها مستفيدةً **صامتة** من تسرُّب `SILK_RATE_LIMIT=0`
# من اختبارٍ مبكّر بلا استرجاع؛ سدُّ التسرب (عازل `_rate_limit_env_guard`)
# كشف الاعتماد فصار الأساسُ معلناً هنا — واختباراتُ الحدّ نفسِه تضبط قيمتها
# صراحةً عبر `_env` وتبني تطبيقاً جديداً فلا تتأثر.
os.environ.setdefault("SILK_RATE_LIMIT", "0")

# الإيقاف المبكر (هدف الدراسة الاحترافية، البند ٨) مطفأ أساسَ جلسة الاختبار
# صراحةً: الموكات القانونية للبعثات («بند مموّه» الواحد) تُحسب أعمدتها 2 من 3
# عمداً — دون الحد الأدنى — والاختبارات القائمة تقيس الذيل الكامل (محلل+كاتب
# مموّهان، صفر كلفة فعلية). في الإنتاج افتراضه مفعّل (`api._early_halt_enabled`)
# واختبارات الموجة ٨ تضبط `SILK_EARLY_HALT=1` صراحةً لقياس الإيقاف نفسه.
os.environ.setdefault("SILK_EARLY_HALT", "0")

# دفتر الاستخدام الدولاري (الموجة p6): الاختبارات الهرمتية التي لا تضبط
# SILK_USAGE_DB كانت تحجز 3$ لكل /research في **دفتر الجهاز الحقيقي**
# (data/usage.db) — تراكم 146$ «تجريبية» في يوم واحد. ومع حارس الإنفاق عند
# حدود المراحل (T9) صار هذا خطراً حقيقياً: تشغيلة مشروعة على جهاز يشغّل
# الاختبارات قد تُوقَف بسقف يومي ملأه pytest. لذا لكل جلسة اختبار دفترُها
# المؤقّت افتراضياً؛ اختبارٌ يضبط SILK_USAGE_DB بنفسه يبقى كما هو (setdefault
# لا يستبدل)، واختباراتُ اشتقاق المسارات تحذف المتغيّر صراحةً قبل القياس.
# مراجعة §58 #14: يُنشَأ المجلّد **فقط** إن لم يكن المتغيّر مضبوطاً (كان
# `mkdtemp()` يُقيَّم دوماً كوسيطٍ لـsetdefault فيُخلّف مجلّداً يتيماً لكل
# تشغيلة)، ويُنظَّف عند انتهاء الجلسة.
if "SILK_USAGE_DB" not in os.environ:
    import atexit as _atexit
    import shutil as _shutil
    import tempfile as _tempfile
    _USAGE_TMPDIR = _tempfile.mkdtemp(prefix="silk-test-usage-")
    os.environ["SILK_USAGE_DB"] = os.path.join(_USAGE_TMPDIR, "usage.db")
    _atexit.register(_shutil.rmtree, _USAGE_TMPDIR, ignore_errors=True)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


@contextlib.contextmanager
def block_network():
    """اقطع الشبكة مؤقتاً — make outbound sockets fail so 'no data' paths hold.

    بلاغ حي (تسريب تسلسل اختبارات CI): جلسة requests المشتركة الدائمة
    (silk_data_layer._session — تجميع اتصالات keep-alive للأداء الإنتاجي)
    قد تحمل اتصالاً TCP حياً فعلياً تركه نداء سابق غير محظور في نفس عملية
    pytest (تشغيل تسلسلي واحد لكل ملفات tests/). إعادة استعمال اتصال
    مجمَّع قائم لا يستدعي socket.socket() من جديد، فيتجاوز الحجب أدناه
    صامتاً ويُرجع بيانات حقيقية رغم دخول هذا السياق — ظهر هذا حين أضاف
    ملف اختبار جديد بضعة نداءات فأزاح ترتيب التنفيذ فكشف اتصالاً مجمَّعاً
    كان يبقى خاملاً غير مستغَل سابقاً. إغلاق تجمّعات الاتصال المعروفة عند
    كل دخول يمنع نجاة اتصال حيّ لاختبار يُفترض به حجب كامل — Session.close()
    يُغلق التجمّع الحالي فقط لا الكائن نفسه، فيُعاد فتح اتصال جديد طبيعياً
    خارج هذا السياق حين تُستأنف الشبكة.
    """
    real = socket.socket

    def _no_net(*a, **k):  # noqa: ANN002, ANN003
        # صياغة بلا كلمة hermetic عمداً: حارس تقارير الإنتاج يرفض أي أثر يحمل
        # الكلمة (إصلاح مراجعة Stage 5) — قطع الشبكة حالة تشغيل صادقة لا بديل
        # بيانات، فلا يجوز أن تسمّم ملاحظاتُه تقريراً مشتقاً في اختبار.
        raise OSError("network disabled for offline test")

    # حيادُ البروكسي (أوّل CI حيّ، 2026-08-27): رسالةُ الفشل المعروضة تُشتقّ من
    # **نوع** الاستثناء (`silk_narrative._EXC_FAMILY_AR`): بيئةٌ خلف بروكسي
    # تُنتِج `ProxyError` ⇒ «تعذّر تأمين الاتصال بالمصدر»، وحجبُ socket نقيّ
    # يُنتِج `ConnectionError` ⇒ «تعذّر الاتصال بالمصدر». فالنصّ المولَّد كان
    # يحمل بصمةَ بيئة التوليد، وعيّنةٌ ملتزَمة وُلدت خلف بروكسي تخالف مخرَجَ CI
    # بلا أيّ انحدارٍ حقيقيّ. تصفيرُ متغيّرات البروكسي يجعل النكهة واحدةً
    # أينما وُلِّدت. Neutralise ambient proxy config: it changes the exception
    # type, and the failure-message flavour is derived from it.
    for _var in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY",
                 "http_proxy", "https_proxy", "all_proxy"):
        os.environ.pop(_var, None)
    try:
        import silk_data_layer
        silk_data_layer._session.close()
    except Exception:  # noqa: BLE001 — أفضل جهد؛ الحجب الأساسي (socket) نافذ بدونه
        pass

    socket.socket = _no_net
    try:
        yield
    finally:
        socket.socket = real


def docx_all_text(path: str) -> str:
    """كل نص مستند Word — فقرات + خلايا جداول (مراجعة المشروع: بعض أقسام
    render_docx صارت جداولاً حقيقية بدل نقاط سردية؛ `doc.paragraphs` وحدها
    لا تصل خلايا الجداول، فتفوّت اختباراتٌ محتوًى انتقل إليها بلا انحدار فعلي).
    """
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


import tempfile

import pytest


# ── تسرُّبُ المجلّدات المؤقّتة · leaked temp dirs ─────────────────────────────
#
# تدقيق 2026-08-30: ٦٧ ملفَّ اختبارٍ ينادي `tempfile.mkdtemp()` بلا تنظيف، فتُخلَّف
# مجلّداتٌ إلى الأبد — بلغت التركةُ ٢٠٩٬٦٤٨ مجلّداً على آلة التطوير. وعند ذلك
# الحجم يتباطأ إنشاءُ الملفات على ويندوز حتى تتعثّر اختباراتٌ لا علاقة لها بما
# تقيسه: رُصِد تعليقٌ داخل `apply_migrations` على قاعدةٍ **جديدة تماماً**، بخيطٍ
# واحد وبلا مالكِ قفل، ثمّ نجحت الشيفرةُ نفسُها في الإعادة. عطلُ بيئةٍ يُقرأ
# انحدارَ شيفرة، ويُهدر وقتَ تشخيصٍ في كلّ موجة.
#
# **إصلاحُ عائلةٍ لا موقع**: تغليفُ `mkdtemp` مرّةً هنا يغطّي الملفّاتِ السبعةَ
# والستّين وكلَّ ما يُكتَب بعدها، بدل ٦٧ تعديلاً يسهو أحدُها. التنظيفُ عند
# **نهاية الجلسة** لا في تفكيك كل اختبار: أثناء الاختبار قد تبقى اتصالاتُ
# SQLite مفتوحةً فيرفض ويندوز الحذف، فيتحوّل التنظيفُ إلى فشلٍ كاذب.
# مجلّدات pytest نفسِها تُستثنى — يملكها إطارُ العمل ويُبقي آخرَ التشغيلات
# عمداً للتشخيص.
_LEAKED_TMPDIRS: list[str] = []
_REAL_MKDTEMP = tempfile.mkdtemp


def _tracked_mkdtemp(*a, **k):
    d = _REAL_MKDTEMP(*a, **k)
    if not os.path.basename(d).startswith("pytest"):
        _LEAKED_TMPDIRS.append(d)
    return d


def _install_tmpdir_tracking() -> None:
    """يُنصَب في `pytest_configure` لا في تجهيزةٍ (fixture).

    بعضُ الملفّات ينادي `mkdtemp()` **وقت الاستيراد** (جمعُ الاختبارات)، أي
    قبل أن تعمل أيُّ تجهيزةٍ ولو كانت `session`-scope — فكانت تلك المجلّداتُ
    تفلت من الكنس. `pytest_configure` يسبق الجمعَ فيغطّيها."""
    tempfile.mkdtemp = _tracked_mkdtemp


def _sweep_leaked_tmpdirs() -> None:
    """اكنس ما خلّفته الجلسةُ — يُستدعى من `pytest_unconfigure`."""
    import shutil
    tempfile.mkdtemp = _REAL_MKDTEMP
    for d in _LEAKED_TMPDIRS:
        shutil.rmtree(d, ignore_errors=True)
    _LEAKED_TMPDIRS.clear()


@pytest.fixture(scope="session")
def _isolated_analyses_db_path():
    """المسار المشترَك طوال الجلسة لِعزل قاعدة التحليلات (Issue #240) —
    يُحسَب مرّةً واحدة فقط؛ التطعيمُ الفعليّ في `_reapply_isolated_
    analyses_db` أدناه (لكل اختبار، لا هنا) كي ينجو من `importlib.reload
    (silk_storage)` الذي تنفّذه بعضُ الاختبارات (يُعيد `_DEFAULT_PATH` إلى
    حرفيّته الأصلية من المصدر، ماحياً أيّ `setattr` سابق)."""
    return os.path.join(tempfile.mkdtemp(), "silk.db")


@pytest.fixture(autouse=True)
def _reapply_isolated_analyses_db(monkeypatch, _isolated_analyses_db_path):
    """Issue #240: قاعدة التحليلات (`analyses`/`research_missions`) كانت
    الوحيدة غير المعزولة عن شجرة العمل الحقيقية — `api.py` يستدعي
    `silk_storage.reap_orphan_research_runs()` بلا مسارٍ صريح عند كل
    `create_app()` (المسار الافتراضي `data/silk.db` في شجرة العمل نفسها لا
    ملفٍّ مؤقّت). اختبارُ جسر منصّةٍ يُطلق دراسةً حقيقية عبر المحرّك بلا
    تنظيفٍ بعدها كان يترك صفّاً `running` دائماً هناك؛ تشغيلةُ الحزمة
    التالية تسمه `failed` عند إقلاعها الأول (المكنَس)، وتترك اختباراتٌ أخرى
    صفوفاً جديدة — تراكمٌ لا نهائيّ عبر كل تشغيلة (اكتُشف أثناء التحقّق من
    توفّر بيانات بحثٍ حقيقية مُلتقَطة محلياً لمحرك دراسة السوق، PR #238).

    **يُطعَّم `silk_storage._DEFAULT_PATH` مباشرةً — لا `SILK_DATA_DIR`/
    `SILK_DB` بيئيّاً.** جُرِّب ضبطُ متغيّرَي البيئة أولاً وكسر ٩ اختباراتٍ
    أخرى بالضبط: عشراتُ الاختبارات (`test_client_report_export.py`،
    `test_wave1_cleanups.py`، `test_wave5c_reports.py`، …) تعزل نفسَها
    بنمطٍ مختلفٍ تماماً — قراءة/كتابة `storage._DEFAULT_PATH` كسمةِ وحدةٍ
    مباشرةً (`saved = storage._DEFAULT_PATH; storage._DEFAULT_PATH = db`)
    لا عبر متغيّرات البيئة إطلاقاً. `_db_path()` يُعطي `SILK_DB` ثم
    `SILK_DATA_DIR` الأولوية فوق `_DEFAULT_PATH` — فضبطُ أيٍّ منهما هنا كان
    يتجاوز تقنية تلك الاختبارات تماماً (تجدُ الاستدعاءاتُ مسارَ هذا التثبيت
    بدل مسارها الخاص المزروع، فتُخفِق بـ404 بدل السلوك المتوقَّع). تطعيمُ
    `_DEFAULT_PATH` نفسِه فقط يصل التوجيهَ إلى **آخر** حلقة القرار — أسفل
    كلا المتغيّرين — فلا يتعارض مع أيّ تقنية عزلٍ قائمة أياً كانت.

    **مُعاد تطبيقه قبل كل اختبار (`autouse` بلا `scope="session"`) عمداً**:
    اختبارُ `test_engine_persist_writes_to_canonical_db_path_not_relative_
    literal` يُنفّذ `importlib.reload(silk_storage)` — يُعيد تنفيذ الوحدة
    كاملةً فيمحو أيّ `setattr` سابق على `_DEFAULT_PATH` (يعود لحرفيّته
    الأصلية `"data/silk.db"` من المصدر). تطعيمٌ بنطاق الجلسة وحده كان
    سيُمحى عند أوّل إعادة تحميلٍ كهذه ويترك كل اختبارٍ لاحقٍ بلا حماية —
    إعادةُ التطعيم بنفس المسار المخزَّن (`_isolated_analyses_db_path`) قبل
    كل اختبار تضمن الحمايةَ دوماً بصرف النظر عمّا فعله اختبارٌ سابق."""
    import silk_storage
    monkeypatch.setattr(silk_storage, "_DEFAULT_PATH", _isolated_analyses_db_path)


@pytest.fixture(autouse=True)
def _reset_http_breaker():
    """EXT-1: القاطعُ صار **يرفض** بلا نداء (`CircuitOpen`) بدل محاولةٍ واحدة — وحالتُه
    حالةُ عمليةٍ عامّة. بلا تصفيرٍ بين الاختبارات يُسقِط اختبارٌ سابقٌ خمسةَ فشلٍ على مضيفٍ
    فيرفض القاطعُ نداءَ اختبارٍ لاحقٍ لا علاقةَ له به (رُصد على الحزمة الكاملة: فشلٌ
    يعتمد على الترتيب في `test_p4_resilient_fetch` و`test_wave7_live_incident_fixes`)."""
    import silk_circuit
    silk_circuit.http_breaker.reset()
    yield
    silk_circuit.http_breaker.reset()


@pytest.fixture(autouse=True)
def _isolated_fact_store(monkeypatch):
    """عزل مخزن الحقائق لكل اختبار — كتابة M2 العابرة دفّأت المخزن الافتراضي
    فتسرّبت حقائق حقيقية بين الاختبارات (اكتُشف عبر test_engine_localprice_layer_offline
    بعد تشغيلات تدقيق Stage 1). Every test gets its own store unless it overrides."""
    monkeypatch.setenv("SILK_STORE_DB",
                       os.path.join(tempfile.mkdtemp(), "store.db"))
    # قاعدة المنصّة (المستأجرون/المصادقة/المحافظ) تُعزَل هنا أيضاً: api.py الجذر
    # يركّب /platform داخل create_app()، فأيّ اختبار قديم يمسّ مساراً تحتها كان
    # سيهيّئ `data/platform.db` في شجرة العمل ويسرّب حالة بين الاختبارات — نفس
    # عائلة الحادثة التي وُلد لها هذا التثبيت. Isolate the 5th store too.
    monkeypatch.setenv("SILK_PLATFORM_DB",
                       os.path.join(tempfile.mkdtemp(), "platform.db"))
    # 1b: عطّل مباعدة النداءات في الاختبارات — الشبكة مقطوعة أصلاً، والمباعدة
    # 250ms × مئات النداءات الفاشلة كانت ستبطئ الحزمة بلا فائدة.
    monkeypatch.setenv("SILK_HTTP_MIN_GAP_MS", "0")
    # نافذة كومتريد الخاصة (بلاغ 429، افتراضي 1100ms) تُصفَّر أيضاً — بلا
    # هذا نامت الحزمة الهيرمتية ~ساعتين فعلياً (1.1ث × مئات نداءات كومتريد
    # المقطوعة الشبكة) — اكتُشف حياً عند إضافة النافذة.
    monkeypatch.setenv("SILK_COMTRADE_MIN_GAP_MS", "0")
    # الموجة ٦ (V5): عزل ملفات التتبّع أيضاً — بلا هذا، اختبارات /research
    # الحقيقية (TestClient) تكتب data/traces/*.jsonl فعلياً على القرص.
    monkeypatch.setenv("SILK_TRACE_DIR", tempfile.mkdtemp())
    # **المخزن الخامس — ذاكرة الطلبات على القرص** (أول CI حيّ بعد عودة Actions،
    # 2026-08-27). كان هذا التثبيت يعزل أربعةً وينسى `silk_cache` وحده، وهو
    # المخزن **الوحيد الذي يُغني عن الشبكة تماماً**: `comtrade_trade`
    # (`silk_data_layer.py:706`) و`world_bank` (`:848`) يمرّان بـ
    # `silk_cache.cached_get` الذي يقرأ ملفَ JSON من `data/cache/` **بلا أيّ
    # socket** — فحجبُ `block_network()` للـsockets لا يلمسه إطلاقاً.
    #
    # الأثر الحقيقي (مُثبَت بإعادة إنتاج مباشرة): على عاملٍ **متّصل** بالإنترنت،
    # أوّلُ اختبارٍ يجلب حيّاً يملأ `data/cache/`، فكلّ اختبار «بلا شبكة» بعده
    # يقرأ قيماً حقيقية من القرص وتسقط تأكيداتُه («فجوة معلنة» تصير رقماً).
    # اثنا عشر اختباراً سقطت على أوّل تشغيلة CI حقيقية بهذا السبب وحده
    # (`test_wave6_trend`، `test_smoke::test_engine_localprice_layer_offline`،
    # `test_stage2a`، `test_stage5_review_fixes`) — وكلّها خضراء في صندوقٍ
    # شبكتُه محجوبة أصلاً، فبقي الثقب غير مرئيّ محلياً.
    #
    # نفس عائلة الحادثة المذكورة في هذا التثبيت أعلاه («تسرّبت حقائق حقيقية بين
    # الاختبارات») — طُبّق العلاج على أربعة مخازن ولم يصل الخامس.
    # The request cache is the one store that serves data with NO socket at all.
    monkeypatch.setenv("SILK_CACHE_DIR", tempfile.mkdtemp())
    # عزل عدّاد data_economics بين الاختبارات — contextvar بلا حدود عملية
    # مستقلة (pytest يُشغّل كل الاختبارات على نفس الخيط)، فاختبار سابق ترك
    # عدّاداً بأرقام عالية كان سيُفعِّل سقف silk_llm_runtime._run_loop
    # الكلي زوراً في اختبار لاحق لا علاقة له (انحدار اكتُشف فعلياً، الموجة
    # ٦). الإنتاج غير متأثر: كل طلب /research يستدعي begin_data_counter()
    # صراحة قبل أي استخدام.
    import silk_context
    silk_context._data_counter.set(None)
    # R2: سجلّ مقابض التشغيلات وعلمُ الإغلاق حالةُ وحدةٍ (لا حدود عملية بين
    # الاختبارات) — اختبارُ إغلاقٍ سابق كان سيمنع كل إطلاقٍ لاحق من المطالبة.
    try:
        from silk_platform import study_runtime as _srt
        _srt.reset_for_tests()
    except Exception:  # noqa: BLE001 — بيئة بلا المنصّة
        pass
    # R2b: سجلّ تشغيلات `/research` الجذري حالةُ وحدةٍ أيضاً (مقابض + سقف + حاصد).
    try:
        import silk_research_runtime as _rrt
        _rrt.reset_for_tests()
    except Exception:  # noqa: BLE001
        pass
    # لا حاجة لتصفير خنق الدخول يدوياً: حالته صارت في قاعدة المنصّة (المعزولة
    # لكل اختبار بالسطر أعلاه) لا في ذاكرة الوحدة — فالعزل يأتي مجّاناً.
    # Login-throttle state lives in the (per-test isolated) platform DB.
    #
    # عامل عمل كلمات المرور مُخفَّض **في الاختبارات فقط**: قياس فعلي — تجزئة
    # bcrypt بعامل ١٢ = ~٢٧٧ms وتحقّق = ~٢٧٣ms، فحزمة المنصّة قفزت ٢٥ث→١٣٣ث
    # واتجاهها يسوء مع كل اختبار جديد. الإنتاج غير متأثّر بثلاث حمايات
    # (الافتراضي ١٢، وقيمة تالفة ⇒ ١٢، وحارس الإقلاع يرفض أقلّ من ١٢ في الإنتاج)،
    # واختبار عامل ١٢ يمسح هذا المتغيّر ويدفع تكلفة تجزئة حقيقية واحدة ليُثبِته.
    # Test-only work-factor reduction; production is protected three ways.
    monkeypatch.setenv("SILK_PLATFORM_BCRYPT_ROUNDS", "4")
    monkeypatch.setenv("SILK_PLATFORM_SCRYPT_N", "1024")


@pytest.fixture(autouse=True)
def _hermetic_env_guard():
    """عزل علم `SILK_HERMETIC` لكل اختبار — اختبارات كانت تضبطه خاماً
    (`os.environ[...] = "1"`) بلا استرجاع فيتسرّب لكل اختبار لاحق: build_view
    يسم test_run زوراً فتظهر لافتة «نموذج توضيحي» في PDF عميلٍ إنتاجي المسار
    (اكتُشف بإعادة إنتاج مباشرة: القفل البصري
    test_visual_pdf_lock_production_entrypoint_bare_no_split_no_leaks يسقط على
    «⚠» فقط حين تسبقه اختبارات التصدير في نفس الحزمة وأدوات PDF مثبَّتة).
    كل اختبار يبدأ بلا العلم، وأي ضبطٍ خام داخله يُمسَح بعده حتماً."""
    old = os.environ.pop("SILK_HERMETIC", None)
    yield
    if old is None:
        os.environ.pop("SILK_HERMETIC", None)
    else:
        os.environ["SILK_HERMETIC"] = old


@pytest.fixture(autouse=True)
def _rate_limit_env_guard():
    """عزل `SILK_RATE_LIMIT`/`SILK_API_KEY` لكل اختبار (صيد الفجوات ٣ —
    عائلة الدرس السابق نفسها): ستةُ اختبارات تضبطهما خاماً وتزيلهما بسطرٍ
    آخرَ الجسدِ يُتخطّى عند فشل توكيدٍ قبله (لا try/finally) — فيتسرّبان
    لكل اختبار لاحق في نفس الحزمة. عزلٌ بنيوي واحد بدل مطاردة كل موضع."""
    saved = {k: os.environ.get(k) for k in ("SILK_RATE_LIMIT", "SILK_API_KEY")}
    yield
    for k, v in saved.items():
        if v is None:
            os.environ.pop(k, None)
        else:
            os.environ[k] = v


# ── طبقة الدخان الحية (opt-in) — gated live-integration lane ─────────────────
# اختبارات موسومة `live` تضرب الشبكة الحقيقية (مصادر مجانية بلا مفتاح فقط —
# لا حرق أرصدة). تُتخطّى دائماً في CI الافتراضي (`pytest tests/ -q`) وتعمل
# فقط حين SILK_RUN_LIVE=1 (مسار workflow_dispatch يدوي، راجع
# .github/workflows/live-smoke.yml). هذا يبدأ إغلاق أكبر فجوة اختبار: غياب
# أي اختبار تكامل حي (كل شيء آخر يقطع الشبكة).

def pytest_configure(config):
    _install_tmpdir_tracking()
    config.addinivalue_line(
        "markers",
        "live: real-network integration smoke test — skipped unless "
        "SILK_RUN_LIVE=1 (opt-in lane, never runs on default CI).")
    # رُتبتا ٢–٣ (خادم حقيقي + متصفّح حقيقي) — تُقلِعان uvicorn فعلياً (وrung 3
    # تشغّل chromium)، أبطأ من الحزمة الهرمتية وخارج ضمانتها «بلا شبكة/بلا
    # عملية خارجية». تُجمَع في وظيفة CI المخصّصة `e2e-live-shape` حصراً
    # (SILK_RUN_E2E=1)، فتبقى `pytest tests/ -q` الافتراضية هرمتية سريعة.
    config.addinivalue_line(
        "markers",
        "e2e: real-server / real-browser rung — boots uvicorn (rung 2) and "
        "may drive chromium (rung 3); skipped unless SILK_RUN_E2E=1 "
        "(the e2e-live-shape CI job — see .github/workflows/e2e-live-shape.yml).")


def pytest_collection_modifyitems(config, items):
    live_on = os.environ.get("SILK_RUN_LIVE") == "1"
    e2e_on = os.environ.get("SILK_RUN_E2E") == "1"
    skip_live = pytest.mark.skip(
        reason="live-network test; set SILK_RUN_LIVE=1 to run "
               "(opt-in lane — see .github/workflows/live-smoke.yml)")
    skip_e2e = pytest.mark.skip(
        reason="real-server/browser rung; set SILK_RUN_E2E=1 to run "
               "(the e2e-live-shape CI job).")
    for item in items:
        if "live" in item.keywords and not live_on:
            item.add_marker(skip_live)
        if "e2e" in item.keywords and not e2e_on:
            item.add_marker(skip_e2e)


def pytest_unconfigure(config):
    """نهايةُ الجلسة — اكنس المجلّدات المؤقّتة التي خلّفتها."""
    _sweep_leaked_tmpdirs()
