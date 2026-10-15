"""الدرس ٢٩٦ — مراجعة المالك لدراسة قهوة البحرين (7/10): أخطاءٌ واقعية تُقال
بثقة، وتناقضاتٌ بين الأقسام، وتسرّبُ قوالب ولغةِ نظام، وقائمةُ عملاءٍ لم تتغيّر.

كلُّ اختبارٍ هنا يعيد إنتاج عبارةٍ من البلاغ نفسِه (direct reproduction على
المُصيِّر أو البوّابة)، ثم يثبت الإصلاح عند منبعه. والبوّاباتُ الثلاث التي طلبها
المالك: (١) تسميةُ حقلٍ/اسمُ وكيلٍ بدل القيمة، (٢) قيمةٌ في السجلّ لا تُوصف
«غير متاحة»، (٣) ادّعاءٌ اقتصاديٌّ كليٌّ أو تنظيميٌّ حادّ بلا سندٍ رسميّ.
"""
import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

import silk_fact_ledger as L                                  # noqa: E402
import silk_quality_gate as Q                                 # noqa: E402


def _led(**entries):
    return {"entries": entries, "order": list(entries), "gaps": []}


def _view(text: str, missions: dict | None = None, iso3: str = "BHR") -> dict:
    return {"market": {"iso3": iso3},
            "deep_research": {"report": {"text": text},
                              "missions": missions or {}}}


# ── بوّابة المالك ١: تسميةُ حالةٍ أو اسمُ وكيلٍ بدل القيمة ───────────────────

def test_status_token_in_a_value_slot_is_filled_with_the_value():
    """«يسجل {{status:x}}» خرج «يسجل مستنتج من أرقام مرصودة» (ص9)."""
    e = L._entry("import_cagr_pct", 13.7, source="UN Comtrade",
                 confidence=0.9, year=2024)
    out, _ = L.bind("معدل النمو المركّب يسجل {{status:import_cagr_pct}} سنوياً.",
                    _led(import_cagr_pct=e))
    assert "13.7%" in out and "مستنتج" not in out and "متاح" not in out


def test_status_token_in_a_value_slot_drops_the_sentence_when_missing():
    e = L._entry("import_cagr_pct", None)
    out, info = L.bind("النمو يسجل {{status:import_cagr_pct}} سنوياً. والسوق مفتوح.",
                       _led(import_cagr_pct=e))
    assert out == "والسوق مفتوح." and len(info["dropped"]) == 1


def test_status_token_in_a_value_slot_is_sent_back_to_the_writer():
    e = L._entry("import_cagr_pct", 13.7, source="UN Comtrade", confidence=0.9)
    issues = L.draft_issues("النمو يسجل {{status:import_cagr_pct}} سنوياً.",
                            _led(import_cagr_pct=e))
    assert any("حيث تلزم القيمة" in i for i in issues)


def test_gate_blocks_a_status_label_printed_as_a_value():
    f = Q._check_field_label_as_value("وتبلغ مستنتَج من أرقام مرصودة خلال الفترة.")
    assert [x["check"] for x in f] == ["status_label_as_value"]
    assert "status_label_as_value" in Q.DELIVERY_BLOCKING_CHECKS
    assert Q.delivery_blockers(f)


def test_gate_blocks_an_agent_name_but_not_urls_or_emails():
    f = Q._check_field_label_as_value("المصدر reach_agent.")
    assert [x["check"] for x in f] == ["agent_name_leak"]
    assert not Q._check_field_label_as_value(
        "راجع https://example.com/a_b أو info_desk@example.com.")
    assert Q.delivery_blockers(f)


def test_client_sanitizer_rewrites_failed_missions_and_contributing_count():
    import silk_reports as R
    out = R._client_sanitize("فشل البعثات أثّر في التحليل، وعدد المؤشرات "
                             "المرصودة المساهمة 97. المصدر Web Search (Exa via Agent Reach).")
    assert "عمليات جمع البيانات" not in out and "97" not in out
    assert "Agent Reach" not in out and "بحث الويب" in out


def test_writer_context_no_longer_carries_the_contributing_count():
    src = open(os.path.join(_ROOT, "silk_ai_judge.py"), encoding="utf-8").read()
    assert 'f"عدد المؤشرات المرصودة المساهمة' not in src


def test_one_other_component_is_named_not_counted():
    """«و1 مكونات أخرى» — باقٍ واحدٌ يُسمّى، والعددُ بمطابقة العربية."""
    import silk_render as R
    conds = R._flip_conditions("conditional", False, {}, "البحرين",
                               missing_components=["c1", "c2", "c3", "c4"])
    text = " ".join(c["condition"] for c in conds)
    assert "و1 " not in text and "c4" in text
    assert R._others_ar(2) == "مكوّنان آخران"


# ── بوّابة المالك ٢: قيمةٌ في السجلّ لا تُوصف «غير متاحة» ─────────────────────

def test_value_in_ledger_declared_unavailable_blocks_even_without_diacritics():
    """ص5: «المعدل المركب غير متاح» بينما السلسلة تعطيه — بلا شدّة."""
    e = L._entry("import_cagr_pct", 13.7, source="UN Comtrade", confidence=0.9)
    f = L.check({"ledger": _led(import_cagr_pct=e)},
                "أما المعدل المركب فهو غير متاح لهذه الفترة.")
    hit = [x for x in f if x["check"] == "ledger_status_mismatch"]
    assert hit and hit[0]["repairable"] is False
    assert Q.delivery_blockers(hit)


def test_cannot_determine_an_observed_share_is_caught_and_repaired():
    """ص19: «لا يمكن تحديد حصة السعودية» والتقرير يحدّدها 3.53%."""
    e = L._entry("saudi_share_pct", 3.53, source="UN Comtrade", confidence=0.9)
    text, rep = L.repair("حصة السعودية لا يمكن تحديدها بدقة. والسوق مفتوح.",
                         _led(saudi_share_pct=e), kinds=("status",))
    assert rep and "3.53%" in text


def test_status_repair_runs_without_the_enforce_flag():
    import silk_export_gate as G
    e = L._entry("import_cagr_pct", 13.7, source="UN Comtrade", confidence=0.9)
    view = {"ledger": _led(import_cagr_pct=e),
            "deep_research": {"report": {"text": "معدل النمو السنوي المركّب غير متاح."}}}
    os.environ.pop(L.ENFORCE_FLAG, None)
    G._ledger_repair_before_gate(view)
    assert "13.7%" in view["deep_research"]["report"]["text"]


def test_cagr_comes_from_the_drawn_series_with_its_source():
    """«21%» عنواناً و«غير متاح» في قسمٍ آخر — والسلسلة المباشرة تعطي غيرَهما."""
    import silk_deep_pillars as P
    rows = [{"value": v, "data_year": y, "source": src, "confidence": c,
             "note": "إجمالي واردات السوق"}
            for y, v, src, c in ((2018, 2.93e6, "UN Comtrade (mirror)", 0.6),
                                 (2019, 4.81e6, "UN Comtrade", 0.9),
                                 (2024, 9.17e6, "UN Comtrade", 0.9))]
    missions = {"trade_flow": {"findings": rows}}
    s = P.import_series(missions)
    assert 13.0 < s["cagr_pct"] < 14.5                 # 2019–2024 المباشرة
    led = L.build_ledger({"deep_research": {"missions": missions}})
    e = led["entries"]["import_cagr_pct"]
    assert e["status"] == L.OBSERVED and abs(e["value"] - s["cagr_pct"]) < 1e-9
    assert "2019–2024" in e["note"]


def test_mirror_and_gcc_sources_are_not_called_unofficial():
    mirror = L._entry("market_imports_usd", 9.17e6, source="UN Comtrade",
                      confidence=0.6, mirrored=True)
    assert mirror["status"] == L.WEAK
    assert "غير رسمية" not in L.render_value(mirror)
    assert "غير رسمية" not in L.render_status(mirror)
    from silk_fact_records import classify_tier
    assert classify_tier("GCC Secretariat") == "A"
    assert classify_tier("هيئة التقييس الخليجية") == "A"


# ── بوّابة المالك ٣: الادّعاءُ الحادّ بلا سندٍ رسميّ ─────────────────────────

def test_sharp_macro_contraction_without_official_source_blocks():
    f = Q._check_sharp_claim_unverified(_view(
        "يواجه الاقتصاد البحريني انكماشاً بنسبة 11.3% في عام 2026."))
    assert [x["check"] for x in f] == ["sharp_claim_unverified"]
    assert Q.delivery_blockers(f)


def test_sharp_macro_backed_by_imf_passes():
    m = {"demographics_economy": {"findings": [
        {"value": -11.3, "source": "IMF WEO", "note": "نمو الناتج المحلي"}]}}
    assert not Q._check_sharp_claim_unverified(_view(
        "يواجه الاقتصاد انكماشاً بنسبة 11.3% في عام 2026.", m))


def test_import_growth_is_not_a_macro_claim():
    assert not Q._check_sharp_claim_unverified(_view(
        "نمت الواردات بنسبة 213% خلال الفترة، وسجل الاقتصاد نمواً بنسبة 2.1%."))


def test_mandatory_decree_without_official_source_blocks():
    """«إلزامية شهادة الحلال منذ ديسمبر 2024 (القرار 61/2024)» بلا رابط."""
    text = "وأصبحت شهادة الحلال إلزامية منذ ديسمبر 2024 بموجب القرار 61/2024."
    assert Q._check_sharp_claim_unverified(_view(text))
    m = {"customs_requirements": {"findings": [
        {"value": "حلال", "note": "القرار 61/2024", "url": "https://www.gac.org.sa/x",
         "source": "web"}]}}
    assert not Q._check_sharp_claim_unverified(_view(text, m))


def test_study_gdp_ignores_bare_growth_words_and_unofficial_sharp_values():
    """«انكماش 11.3%» التقطته كلمةُ «النمو» العارية؛ رقمُ IMF بلا data_year."""
    import silk_study_case as C
    dr = {"missions": {"trade_flow": {"findings": [
        {"value": -11.3, "data_year": 2026, "source": "web", "note": "النمو المركب للواردات"}]}}}
    assert C._macro_gdp(dr) is None
    dr["missions"]["risk_news"] = {"findings": [
        {"value": -11.3, "data_year": 2026, "source": "web", "note": "GDP growth"}]}
    assert C._macro_gdp(dr) is None                     # حادٌّ بلا مصدرٍ رسميّ
    dr["missions"]["demographics_economy"] = {"findings": [
        {"value": 2.8, "data_year": 2026, "source": "IMF WEO",
         "note": "نمو الناتج المحلي الإجمالي الحقيقي — تقدير/تنبؤ"}]}
    g = C._macro_gdp(dr)
    assert g["value"] == 2.8 and g["forecast"] is True


def test_imf_datapoint_carries_a_structural_year():
    src = open(os.path.join(_ROOT, "silk_imf_agent.py"), encoding="utf-8").read()
    assert "data_year=int(yr)" in src


# ── الوقائع: المسار البرّي، المؤشرات، الأسعار، العدّ، القائمة ─────────────────

def test_bahrain_logistics_mission_gets_the_causeway_route():
    import silk_missions as M
    ctx = M.transport_route_context("BHR")
    assert "جسر الملك فهد" in ctx and "نقطة الدخول الوحيدة" in ctx
    assert M.transport_route_context("NLD") == ""


def test_causeway_dash_is_not_a_reversed_route_for_a_land_market():
    assert not Q._check_reversed_route(_view(
        "المنفذ: البحرين – الدمام عبر جسر الملك فهد."))


def test_writer_prompt_carries_the_new_factual_rules():
    src = open(os.path.join(_ROOT, "silk_ai_judge.py"), encoding="utf-8").read()
    assert "المؤشران القابلان للقياس" not in src
    for needle in ("جسر الملك فهد", "التدفّقُ العكسي", "بدولتها الأم",
                   "لا تصف قناةً بأنها «موثّقة»", "الادّعاءات الحادّة"):
        assert needle in src


def test_border_unit_values_are_not_shelf_price_rows():
    import silk_render as R
    assert R._border_price_dp({"source": "UN Comtrade", "note": "16.48 USD/kg"})
    assert R._border_price_dp({"source": "web", "note": "قيمة الاستيراد عند الحدود"})
    assert not R._border_price_dp({"source": "web", "note": "سعر الرف 3.2 BHD"})


def test_supplier_nature_summary_agrees_in_number():
    from silk_narrative import count_sentence
    assert count_sentence(2, "منتج", "منتجان", "منتجين", "منتجاً") == "منتجان"
    src = open(os.path.join(_ROOT, "silk_commercial_analysis.py"), encoding="utf-8").read()
    assert 'f"{n_prod} منتج و{n_re} معيد تصدير مرجَّح"' not in src


def test_cafes_roasters_and_equipment_are_not_importers():
    from silk_style_contract import lead_local_competitor, lead_product_fit
    for lead in ({"name": "Moments Café", "category": "Cafe"},
                 {"name": "Bean Machines", "category": "Coffee machine supplier"}):
        assert lead_product_fit(lead, "090121", "قهوة محمصة", "BHR")[0] is False
    # المحمصةُ منافسٌ: تُسقَط من قائمة المستوردين، وتبقى «منافساً» في الدراسة.
    roaster = {"name": "Hub Coffee Roasters", "category": ""}
    assert lead_local_competitor(roaster, "090121")
    assert not lead_local_competitor(roaster, "090111")
    import silk_reports as R
    dr = {"market": {"iso3": "BHR"}, "hs_code": "090121", "product": "قهوة محمصة",
          "report": {"text": ""}}
    assert R._clean_leads([dict(roaster, phone="+973 1700 0000")], dr) == []
    assert lead_product_fit({"name": "Ashraf FMCG", "category": "Distributor"},
                            "090121", "قهوة محمصة", "BHR")[0] is True


def test_report_named_entry_door_stays_on_the_list_without_contact():
    import silk_reports as R
    dr = {"market": {"iso3": "BHR", "name_en": "Bahrain"}, "hs_code": "090121",
          "product": "قهوة محمصة",
          "report": {"text": "أول باب دخول هو Ashraf FMCG لتوزيع السلع الاستهلاكية."}}
    out = R._clean_leads([{"name": "Ashraf FMCG"}, {"name": "Nobody Trading"}], dr)
    assert [x["name"] for x in out] == ["Ashraf FMCG"]


def test_web_candidates_are_not_cut_by_maps_leads():
    import silk_gmaps as G
    maps = [{"name": f"Cafe {i}"} for i in range(G._TOP_N)]
    merged = G._merge_web_candidates(maps, ["Ashraf FMCG"])
    assert len(merged) == G._TOP_N and merged[-1]["name"] == "Ashraf FMCG"


def test_study_template_has_a_land_route_and_a_forecast_variant():
    import silk_study_case as C
    assert C._land_route("BHR") == {"land": True,
                                    "land_route": "عبر جسر الملك فهد من المنطقة الشرقية"}
    tpl = open(os.path.join(_ROOT, "data", "study_templates_ar.yaml"),
               encoding="utf-8").read()
    assert "when: logistics_land" in tpl and "when: macro_forecast" in tpl


# ── «fix it all»: المحرّك والعدّ والتكرار ─────────────────────────────────────

def test_decision_engine_reads_the_series_cagr():
    """قرار المالك: عمودُ السوق يقرأ مركّبَ السلسلة نفسَه الذي يقرؤه النصّ."""
    import silk_deep_pillars as P
    rows = [{"value": v, "data_year": y, "source": "UN Comtrade", "confidence": 0.9,
             "note": "إجمالي واردات السوق"}
            for y, v in ((2019, 4.81e6), (2024, 9.17e6))]
    dr = {"missions": {"trade_flow": {"findings": rows}}}
    pi = P.build_pillar_inputs(dr)
    flat = {k: v for g in pi.values() if isinstance(g, dict) for k, v in g.items()}
    assert flat["import_cagr_pct"] == P.import_series(dr["missions"])["cagr_pct"]


def test_requirements_count_is_the_printed_table_rows():
    """«ثمانية بنود» بجوار جدولٍ من سبعة — العددُ من صفوف الجدول نفسِه."""
    import silk_render as R
    result = {"market": {"iso3": "BHR"}, "hs_code": "090121",
              "deep_research": {"missions": {"customs_requirements": {"findings": [
                  {"value": f"بند {i}", "note": "x"} for i in range(8)]}}}}
    rows = [r for r in R._requirements_safe(result) if not r.get("gap")]
    led = L.build_ledger(result)
    assert rows and led["entries"]["requirements_count"]["value"] == len(rows)


def test_the_same_printed_price_line_appears_once():
    import silk_reports as R
    from docx import Document
    base = "سعر رف قهوة محمصة 250 غ " + "و" * 200
    dr = {"price_rows": [{"value": "3.2 BHD", "note": base + " متجر أ"},
                         {"value": "3.2 BHD", "note": base + " متجر ب"}]}
    doc = Document()
    R._client_price_observations(doc, dr, "ar")
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == len(set(bullets))


# ── «كل مرة تظهر نفس المشاكل»: سجلُّ بلاغات المالك ───────────────────────────

def test_every_owner_flag_matches_its_own_example():
    rows = Q.owner_flags()
    assert len(rows) >= 10
    for r in rows:
        iso = sorted(r["markets"])[0] if r["markets"] else ""
        assert Q.owner_flag_hits(r["example"], iso), r["id"]


def test_owner_flags_block_delivery_and_return_drafts_to_the_writer():
    f = Q._check_owner_flags("ميناء خليفة نقطة الدخول الوحيدة.", "BHR")
    assert f and Q.delivery_blockers(f)
    assert not Q._check_owner_flags("ميناء خليفة نقطة الدخول الوحيدة.", "SGP")
    led = {"entries": {}, "market_iso3": "BHR"}
    assert any("رفضها المالك" in i for i in L.draft_issues(
        "قدّرنا العبور قياساً على شحنات من إثيوبيا.", led))


def test_no_committed_sample_carries_an_owner_flagged_phrase():
    from docx import Document
    sdir = os.path.join(_ROOT, "samples")
    for name in os.listdir(sdir):
        if name.startswith("kuwait_peanut_butter"):
            continue          # طبعةٌ تاريخيةٌ مثبَّتة قبل الإصلاح (samples/README.md)
        path = os.path.join(sdir, name)
        if name.endswith(".md"):
            text = open(path, encoding="utf-8").read()
        elif name.endswith(".docx"):
            d = Document(path)
            text = "\n".join([p.text for p in d.paragraphs] + [
                c.text for t in d.tables for r in t.rows for c in r.cells])
        else:
            continue
        assert not Q.owner_flag_hits(text, "BHR"), name


def test_replay_tool_rebuilds_the_real_client_document():
    """الأداةُ تبني مستندَ العميل الفعليّ من نتيجةٍ مخزَّنة وتفحصه — الرتبةُ
    التي كانت غائبة: الإصلاحُ يُثبَت على الدراسة نفسِها لا على عيّنة."""
    sys.path.insert(0, os.path.join(_ROOT, "tools"))
    import audit_study_reports as A
    from gen_client_report_sample import build_sample_view
    text, findings = A.client_text(build_sample_view("ar"))
    assert len(text) > 500 and isinstance(findings, list)


# ── المراجعة الذاتية §58 قبل الدمج — كلُّ ملاحظةٍ بقفل ─────────────────────────

def test_review_import_growth_near_the_word_economy_is_not_macro():
    assert not Q._check_sharp_claim_unverified(_view(
        "رغم تباطؤ الاقتصاد، سجّل نمو الواردات 21% خلال الفترة."))


def test_review_status_window_stops_at_the_sentence_and_at_a_stated_value():
    e = L._entry("saudi_share_pct", 3.5, source="UN Comtrade", confidence=0.9)
    led = {"ledger": _led(saudi_share_pct=e)}
    assert not [f for f in L.check(led, "الحصة السعودية 3.5%. تعذّر تحديد سعر الرف.")
                if f["check"] == "ledger_status_mismatch"]
    text, rep = L.repair("تستحوذ السعودية على 3% من السوق ولم يُرصد سعر رف مثبت.",
                         _led(saudi_share_pct=e), kinds=("status",))
    assert not rep and "سعر رف" in text


def test_review_brand_named_cafe_stays_in_study_entities():
    from silk_style_contract import lead_local_competitor, lead_product_fit
    lead = {"name": "Café Najjar", "category": ""}
    assert lead_product_fit(lead, "090121", "قهوة محمصة", "BHR")[0] is True
    assert lead_local_competitor(lead, "090121")          # خارج قائمة المستوردين فقط


def test_review_short_maps_list_keeps_every_web_candidate():
    import silk_gmaps as G
    names = [f"Distributor {i}" for i in range(10)]
    merged = G._merge_web_candidates([{"name": "Alpha Co"}, {"name": "Beta Co"}], names)
    assert len(merged) == 12


def test_review_growth_window_ends_at_the_base_year():
    import silk_deep_pillars as P
    rows = [{"value": v, "data_year": y, "source": src, "confidence": 0.9,
             "note": "إجمالي واردات السوق"}
            for y, v, src in ((2019, 4e6, "UN Comtrade"), (2022, 6e6, "UN Comtrade"),
                              (2024, 9e6, "UN Comtrade (mirror)"))]
    s = P.import_series({"trade_flow": {"findings": rows}})
    assert s["growth_pct"] == 125.0                       # 2019→2024 لا 2019→2022


def test_review_study_inflation_follows_the_gdp_year():
    import silk_study_case as C
    dr = {"missions": {"demographics_economy": {"findings": [
        {"value": 2.8, "data_year": 2024, "source": "IMF WEO", "note": "نمو الناتج المحلي"},
        {"value": 1.1, "data_year": 2024, "source": "IMF WEO", "note": "التضخم"},
        {"value": 3.9, "data_year": 2027, "source": "IMF WEO", "note": "التضخم — تنبؤ"}]}}}
    assert C._indicator(dr, ("التضخم",), -10, 200, year=2024)["value"] == 1.1
