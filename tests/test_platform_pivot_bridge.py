"""أقفال جسر المنصّة⇄المحرّك — hermetic locks for the market-study pivot.

قرارات المالك الثلاثة (2026-08-17) مقفولة هنا:
١) الإطلاق يشغّل المحرّك فعلاً (لا بريد، لا SMTP)؛
٢) «الباقة فقط» — الرصيد الصفري/السالب لا يمنع الإطلاق؛
٣) لا تكلفة داخلية تعبر لمصنع (تجريد المفاتيح بنيوي).
+ عقود الصدق: فشل التشغيل يعيد مسودّةً بسببٍ معلن ويُرجع الحصّة؛ ETA لا
يُقدَّم «مقاساً» بلا قياس؛ كنس الأيتام لا يحكم بلا دليل.

لا شبكة: المحرّك يُحاكى بـmonkeypatch على `engine_bridge._run_engine` (أو
مقعد `SILK_PLATFORM_FAKE_ENGINE` الذي يعبر مسار `silk_storage` الحقيقي).
"""
from __future__ import annotations

import base64
import os
import sqlite3

import pytest

from tests.platform_helpers import hdr, login, make_factory, seed, client


ADMIN = ("admin@silk.local", None)


@pytest.fixture()
def env(monkeypatch):
    """بذر + عميل + مصنع فضي بلا رصيد — the standard pivot fixture."""
    monkeypatch.setenv("SILK_SEED_ADMIN_PASSWORD", "AdminPass1")
    info = seed(monkeypatch)
    cl = client()
    fac = make_factory("silver", "pivot@f.local")
    tok = login(cl, fac["email"], fac["password"])
    return {"cl": cl, "fac": fac, "tok": tok, "info": info}


def _mk_study(cl, tok, **over):
    body = {"product": "تمور سكري", **over}
    r = cl.post("/platform/studies", json=body, headers=hdr(tok))
    assert r.status_code == 200, r.text
    return r.json()


def _study_row(sid):
    from silk_platform import db as pdb
    conn = pdb.connect()
    try:
        return dict(conn.execute("SELECT * FROM studies WHERE id = ?",
                                 (sid,)).fetchone())
    finally:
        conn.close()


def _month_count(aid):
    from silk_platform import db as pdb
    conn = pdb.connect()
    try:
        return int(conn.execute(
            "SELECT current_month_study_count FROM accounts WHERE id = ?",
            (aid,)).fetchone()[0])
    finally:
        conn.close()


def _canned_result(analysis_id=777, classified=True):
    # سوق جوهري واحد (واقعية المحاكاة — §58): حارس «المكتملة الفارغة» في الجسر
    # يعيد نتيجة بلا أي نقطة بيانات إلى مسودّة؛ أسواق فارغة هنا كانت تحاكي
    # بالضبط الحالة التي شكاها المالك، لا التشغيلة الناجحة التي تدّعيها.
    return {"product": "تمور سكري", "hs_code": "080410", "classified": classified,
            "markets": [{"country": "الإمارات", "iso3": "ARE",
                         "total_score": 0.5, "confidence": 0.4,
                         "components": {}}],
            "analysis_id": analysis_id,
            "hs_note": "اختبار", "note": "canned"}


# ── ١) الإطلاق يشغّل المحرّك ويكتمل — launch actually runs the engine ────────
def test_launch_runs_engine_and_completes(env, monkeypatch):
    import silk_platform.engine_bridge as eb
    calls = {}

    def fake_run(product, hs_code, market_pref, *_a, **_k):
        calls["args"] = (product, hs_code, market_pref)
        return _canned_result()

    monkeypatch.setattr(eb, "_run_engine", fake_run)
    s = _mk_study(env["cl"], env["tok"], hs_code="080410", market_pref="ARE")
    r = env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["state"] == "in_progress"
    assert "eta_s" in out and out["eta_basis"] in ("measured_median",
                                                   "declared_default")
    assert eb.wait_idle(15)
    row = _study_row(s["id"])
    assert row["state"] == "completed"
    assert row["analysis_id"] == 777
    assert row["run_finished_at"]
    assert calls["args"] == ("تمور سكري", "080410", "ARE")


# ── ٢) قرار المالك: الرصيد لا يمنع الإطلاق — wallet is NOT a launch gate ─────
def test_launch_ignores_wallet_balance(env, monkeypatch):
    """عكسُ الاختبار القديم `insufficient_balance_blocks_launch` عمداً —
    قرار المالك «الباقة فقط»: مصنع برصيد صفري (بل ومدين) يُطلق دراسته."""
    import silk_platform.engine_bridge as eb
    monkeypatch.setattr(eb, "_run_engine",
                        lambda *a, **_k: _canned_result(analysis_id=778))
    # اجعل الحساب مديناً صراحةً — أقسى من الصفر.
    from silk_platform import db as pdb, wallet
    from silk_platform.models import Operation
    conn = pdb.connect()
    try:
        wallet.post_entry(conn, account_id=env["fac"]["account_id"],
                          actor_user_id=env["fac"]["user_id"],
                          operation=Operation.EMAIL_SENT, amount=-500,
                          description="debt fixture", allow_negative=True)
    finally:
        conn.close()
    s = _mk_study(env["cl"], env["tok"])
    r = env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    assert r.status_code == 200, r.text
    assert eb.wait_idle(15)
    assert _study_row(s["id"])["state"] == "completed"


# ── ٣) لا إطلاق بلا منتج — a study without a product cannot launch ───────────
def test_launch_requires_product(env):
    r = env["cl"].post("/platform/studies", json={"title_ar": "قديمة"},
                       headers=hdr(env["tok"]))
    assert r.status_code == 200
    sid = r.json()["id"]
    r = env["cl"].post(f"/platform/studies/{sid}/launch", headers=hdr(env["tok"]))
    assert r.status_code == 422
    assert r.json()["detail"]["error"] == "study_no_product"
    # لم تُستهلك حصّة — the refused launch consumed nothing.
    assert _month_count(env["fac"]["account_id"]) == 0


# ── ٤) الفشل = مسودّة بسبب معلن + إرجاع الحصّة — failure reverts honestly ────
def test_engine_failure_reverts_to_draft_and_releases_quota(env, monkeypatch):
    import silk_platform.engine_bridge as eb

    def boom(*a, **_k):
        raise RuntimeError("engine exploded")

    monkeypatch.setattr(eb, "_run_engine", boom)
    s = _mk_study(env["cl"], env["tok"])
    r = env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    assert r.status_code == 200
    # لا نفحص العدّاد الوسيط هنا — خيط الفشل قد يسبق القراءة (سباق اختبار لا
    # إنتاج): الحجز ثم الإرجاع مثبتان معاً بالحالة النهائية أدناه وبردّ 200.
    assert eb.wait_idle(15)
    row = _study_row(s["id"])
    assert row["state"] == "draft"
    # صيد الفجوات ٣ (الدرس 166): run_error سطح مصنع — لا اسم صنف استثناء
    # (رطانة)؛ رسالة عميل + التفصيل المنقّح، والصنف الكامل في سجل التشغيل.
    assert "RuntimeError" not in (row["run_error"] or "")
    assert "تعذّر إكمال الدراسة" in (row["run_error"] or "")
    assert "engine exploded" in (row["run_error"] or "")
    assert row["launched_by_user_id"] is None            # حصّة المستخدم تحرّرت
    assert _month_count(env["fac"]["account_id"]) == 0   # وحصّة الحساب أُرجعت


def test_unclassified_result_reverts_with_guidance(env, monkeypatch):
    """`classified=False` = تقرير بلا أسواق — لا يُحرَق على المصنع كنجاح."""
    import silk_platform.engine_bridge as eb
    monkeypatch.setattr(
        eb, "_run_engine",
        lambda *a, **_k: _canned_result(analysis_id=779, classified=False))
    s = _mk_study(env["cl"], env["tok"])
    env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    assert eb.wait_idle(15)
    row = _study_row(s["id"])
    assert row["state"] == "draft"
    assert "تصنيف" in (row["run_error"] or "")
    assert _month_count(env["fac"]["account_id"]) == 0


def test_missing_analysis_id_is_failure_not_fake_success(env, monkeypatch):
    """نتيجة بلا `analysis_id` (فشل حفظ مبتلَع في المحرّك) ليست «مكتملة»."""
    import silk_platform.engine_bridge as eb
    monkeypatch.setattr(eb, "_run_engine",
                        lambda *a, **_k: {"classified": True, "markets": []})
    s = _mk_study(env["cl"], env["tok"])
    env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    assert eb.wait_idle(15)
    row = _study_row(s["id"])
    assert row["state"] == "draft"
    assert "حفظ" in (row["run_error"] or "")


# ── ٥) ETA صادق — measured only when actually measured ───────────────────────
def test_eta_declared_default_without_history(env):
    from silk_platform import db as pdb
    import silk_platform.engine_bridge as eb
    conn = pdb.connect()
    try:
        eta, basis = eb.eta_seconds(conn)
    finally:
        conn.close()
    assert basis == "declared_default"
    assert eta == eb._DEFAULT_ETA_S


def test_eta_measured_median_with_history(env):
    from silk_platform import db as pdb
    import silk_platform.engine_bridge as eb
    conn = pdb.connect()
    try:
        for i, dur in enumerate((10, 20, 30)):
            conn.execute(
                "INSERT INTO studies (owner_id, state, analysis_id, "
                "run_started_at, run_finished_at, created_at, updated_at) "
                "VALUES (?, 'completed', ?, ?, ?, '', '')",
                (env["fac"]["account_id"], 900 + i,
                 "2026-08-17T10:00:00+00:00",
                 f"2026-08-17T10:00:{dur:02d}+00:00"))
        conn.commit()
        eta, basis = eb.eta_seconds(conn)
    finally:
        conn.close()
    assert basis == "measured_median"
    assert eta == 20


def test_get_study_attaches_eta_only_in_progress(env, monkeypatch):
    import silk_platform.engine_bridge as eb
    s = _mk_study(env["cl"], env["tok"])
    r = env["cl"].get(f"/platform/studies/{s['id']}", headers=hdr(env["tok"]))
    assert "eta_s" not in r.json()
    ev = __import__("threading").Event()
    monkeypatch.setattr(eb, "_run_engine",
                        lambda *a, **_k: (ev.wait(10), _canned_result(781))[1])
    env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    r = env["cl"].get(f"/platform/studies/{s['id']}", headers=hdr(env["tok"]))
    body = r.json()
    assert body["state"] == "in_progress" and "eta_s" in body
    ev.set()
    assert eb.wait_idle(15)


# ── ٦) كنس الأيتام — the restart-orphan sweep, both branches ─────────────────
def test_orphan_sweep_completes_and_reverts(env, monkeypatch):
    import silk_platform.engine_bridge as eb
    from silk_platform import db as pdb
    aid = env["fac"]["account_id"]
    # R1 (تدقيق 2026-09-01): «مقروء» لم يعد يكفي — صفّ محرّك **مكتمل** حقيقي
    # في القاعدة المعزولة، أحدثُ من محاولة الدراسة، بدل محاكاة `get_analysis`.
    import silk_storage
    done_aid = silk_storage.create_research_run("تمور سكري", "ARE", "080410", {})
    silk_storage.update_research_status(done_aid, "completed")
    conn = pdb.connect()
    try:
        # يتيمة أنجزت (تشغيلة محرّك مكتملة) ويتيمة لم تنجز.
        conn.execute(
            "INSERT INTO studies (owner_id, state, analysis_id, launched_at, "
            "launched_by_user_id, run_started_at, created_at, updated_at) "
            "VALUES (?, 'in_progress', ?, '2026-08-17T09:00:00+00:00', ?, "
            "'2026-08-17T09:00:00+00:00', '', '')",
            (aid, done_aid, env["fac"]["user_id"]))
        done_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        conn.execute(
            "INSERT INTO studies (owner_id, state, launched_at, "
            "launched_by_user_id, run_started_at, created_at, updated_at) "
            "VALUES (?, 'in_progress', '2026-08-17T09:00:00+00:00', ?, "
            "'2026-08-17T09:00:00+00:00', '', '')", (aid, env["fac"]["user_id"]))
        lost_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        # كأنّ حجزها وقع فعلاً — العدّاد 2 قبل الكنس.
        conn.execute("UPDATE accounts SET current_month_study_count = 2, "
                     "quota_period = ? WHERE id = ?",
                     (__import__("silk_platform.quota", fromlist=["q"])
                      .current_period(), aid))
        conn.commit()
        swept = eb.sweep_orphans(conn)
    finally:
        conn.close()
    assert swept == {"completed": 1, "reverted": 1, "skipped": 0}
    done = _study_row(done_id)
    assert done["state"] == "completed"
    assert done["run_finished_at"] is None      # لا اختلاق وقت انتهاء
    lost = _study_row(lost_id)
    assert lost["state"] == "draft"
    assert "إعادة نشر" in lost["run_error"]
    assert _month_count(aid) == 1               # أُرجعت حصّة اليتيمة الضائعة فقط


def test_orphan_sweep_skips_active_runs(env):
    import silk_platform.engine_bridge as eb
    from silk_platform import db as pdb
    aid = env["fac"]["account_id"]
    conn = pdb.connect()
    try:
        conn.execute(
            "INSERT INTO studies (owner_id, state, created_at, updated_at) "
            "VALUES (?, 'in_progress', '', '')", (aid,))
        sid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        conn.commit()
        # السجلّ صار {معرّف: لحظة البدء} — «حيّ» له حدٌّ أعلى الآن (تشغيلةٌ
        # تجاوزت مهلتها تُكنَس ولو بقي خيطها معلَّقاً).
        import time as _t
        with eb._LOCK:
            # R2 (تدقيق 2026-09-01): القيمة زوجٌ (رمز، لحظة) منذ إصلاح الرمز —
            # الطفوُ الخام كان يرفع TypeError يبتلعه الكنس فيمرّ الاختبار فارغاً.
            eb._ACTIVE[sid] = ("", _t.monotonic())
        try:
            swept = eb.sweep_orphans(conn)
        finally:
            with eb._LOCK:
                eb._ACTIVE.pop(sid, None)
    finally:
        conn.close()
    assert swept == {"completed": 0, "reverted": 0, "skipped": 0}
    assert _study_row(sid)["state"] == "in_progress"


# ── ٧) release_launch محروس — guarded, never below zero, period-aware ────────
def test_release_launch_guarded(env):
    from silk_platform import db as pdb, quota
    aid = env["fac"]["account_id"]
    conn = pdb.connect()
    try:
        assert quota.release_launch(conn, aid) is False   # لا شيء محجوزاً
        conn.execute("UPDATE accounts SET current_month_study_count = 1, "
                     "quota_period = ? WHERE id = ?",
                     (quota.current_period(), aid))
        conn.commit()
        assert quota.release_launch(conn, aid) is True
        assert _month_count(aid) == 0
        assert quota.release_launch(conn, aid) is False   # لا نزول تحت الصفر
        # فترة قديمة ⇒ لا مساس بعدّاد شهرٍ سيُصفَّر أصلاً.
        conn.execute("UPDATE accounts SET current_month_study_count = 3, "
                     "quota_period = '2020-01' WHERE id = ?", (aid,))
        conn.commit()
        assert quota.release_launch(conn, aid) is False
    finally:
        conn.close()


# ── ٨) تقرير العميل — client report, cost keys structurally stripped ─────────
def _complete_study(env, sid, analysis_id=901):
    from silk_platform import db as pdb
    conn = pdb.connect()
    try:
        conn.execute("UPDATE studies SET state='completed', analysis_id=? "
                     "WHERE id = ?", (analysis_id, sid))
        conn.commit()
    finally:
        conn.close()


def test_report_409_before_completion(env):
    s = _mk_study(env["cl"], env["tok"])
    r = env["cl"].get(f"/platform/studies/{s['id']}/report", headers=hdr(env["tok"]))
    assert r.status_code == 409
    assert r.json()["detail"]["error"] == "no_report_yet"


def test_report_strips_cost_keys_recursively(env, monkeypatch):
    s = _mk_study(env["cl"], env["tok"])
    _complete_study(env, s["id"])
    import silk_render
    import silk_storage
    monkeypatch.setattr(silk_storage, "get_analysis",
                        lambda i, path=None: {"id": i, "product": "تمور"})
    # محاكاة بشكل build_view **الحقيقي** (brief قائمة أسطر، صفوف الأسواق
    # country/score) — المحاكاة بمفاتيح مخترعة أعمت الحارس عن خلل عرضٍ فعلي
    # (ملاحظة §58: الواجهة كانت تقرأ name/total_score فتعرض «—» لكل صف).
    monkeypatch.setattr(
        silk_render, "build_view",
        lambda found, lang="ar": {"brief": ["القرار: صالح",
                                             "الموقع: متوسط"],
                       "cost_usd": 9.9,
                       "data_economics": {"llm_calls": 3},
                       "markets": [{"country": "الإمارات", "score": 0.5,
                                    "confidence": 0.6,
                                    "cost_usd_estimate": 1.1}],
                       "nested": {"research_costs": {}, "keep": 1}})
    r = env["cl"].get(f"/platform/studies/{s['id']}/report", headers=hdr(env["tok"]))
    assert r.status_code == 200, r.text
    view = r.json()["view"]
    import json as _json
    flat = _json.dumps(view, ensure_ascii=False)
    from silk_platform.engine_bridge import _COST_KEYS
    for k in _COST_KEYS:
        assert k not in flat, f"cost key {k} leaked to a factory surface"
    assert view["brief"] == ["القرار: صالح", "الموقع: متوسط"]
    assert view["nested"]["keep"] == 1
    assert view["markets"][0]["country"] == "الإمارات"
    # والصفحة تقرأ المفاتيح الحقيقية لا المخترعة.
    import pathlib
    page = (pathlib.Path(__file__).resolve().parent.parent
            / "web" / "platform.html").read_text(encoding="utf-8")
    assert "m.country || m.name" in page
    assert "m.score != null ? m.score : m.total_score" in page
    assert 'briefLines.map(esc).join("\\n")' in page


def test_report_declared_gap_when_analysis_missing(env, monkeypatch):
    s = _mk_study(env["cl"], env["tok"])
    _complete_study(env, s["id"], analysis_id=999999)
    import silk_storage
    monkeypatch.setattr(silk_storage, "get_analysis", lambda i, path=None: None)
    r = env["cl"].get(f"/platform/studies/{s['id']}/report", headers=hdr(env["tok"]))
    assert r.status_code == 404
    assert r.json()["detail"]["error"] == "analysis_missing"


def test_report_brief_is_sanitized_of_internal_jargon(env, monkeypatch):
    """سطر الخلاصة كان يسرّب معرّف شيفرة للعميل: «أضف بطاقة منتجك (product_card)»
    (رصد جولة 2026-08-17) — التنقية خادمية فتشمل العرض وdocx وPDF معاً."""
    s = _mk_study(env["cl"], env["tok"])
    _complete_study(env, s["id"])
    import silk_render
    import silk_storage
    monkeypatch.setattr(silk_storage, "get_analysis",
                        lambda i, path=None: {"id": i, "product": "تمور"})
    monkeypatch.setattr(
        silk_render, "build_view",
        lambda found, lang="ar": {"brief": [
            "أضف بطاقة منتجك (product_card) للحصول على موقعك التنافسي",
            "التوصية: صالح للمضي"], "markets": []})
    r = env["cl"].get(f"/platform/studies/{s['id']}/report", headers=hdr(env["tok"]))
    assert r.status_code == 200, r.text
    brief = r.json()["view"]["brief"]
    assert "product_card" not in " ".join(brief)
    # الجملة العربية تبقى مفهومة، والسطر السليم لا يُمَسّ.
    assert brief[0] == "أضف بطاقة منتجك للحصول على موقعك التنافسي"
    assert brief[1] == "التوصية: صالح للمضي"


def test_report_docx_carries_no_internal_jargon(env):
    """§58 M1: «(product_card)» كانت تُطبع حرفياً في قسم «موقعك التنافسي» من
    كل Word/PDF ينزّله المصنع (منبعها silk_render وكل دراسة منصّة بلا بطاقة).
    القفل على **نص docx المستخرَج فعلاً** لا على JSON وحده."""
    import io

    import pytest
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx غير متاحة")
    from silk_platform.engine_bridge import _fake_result
    from silk_storage import save_analysis
    aid = save_analysis(_fake_result("تمور سكري", "080410"), None)
    s = _mk_study(env["cl"], env["tok"])
    _complete_study(env, s["id"], analysis_id=int(aid))
    r = env["cl"].get(f"/platform/studies/{s['id']}/report.docx",
                      headers=hdr(env["tok"]))
    assert r.status_code == 200, r.text
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    assert "product_card" not in text, "معرّف شيفرة داخلي في مستند العميل"
    # وJSON العرض كذلك (الملاحظة تصل خاماً في competitive_position.note).
    rj = env["cl"].get(f"/platform/studies/{s['id']}/report",
                       headers=hdr(env["tok"]))
    import json as _json
    assert "product_card" not in _json.dumps(rj.json(), ensure_ascii=False)


def test_report_pdf_is_throttled_per_account(env, monkeypatch):
    """§58 M5: تحويل soffice عملية كاملة — جلسة واحدة كانت تستطيع خنق
    الحاوية بتنزيلات متوازية. عدّاد مسمى لكل حساب برمز معلن."""
    monkeypatch.setenv("SILK_PLATFORM_PDF_MAX_REQUESTS", "1")
    s = _mk_study(env["cl"], env["tok"])
    # الخنق يسبق فحص الاكتمال عمداً (يحمي المورد لا المسار السعيد فقط) —
    # مسودّة تكفي لإثباته: الأولى 409 (مرّت من الخانق)، الثانية 429.
    r1 = env["cl"].get(f"/platform/studies/{s['id']}/report.pdf",
                       headers=hdr(env["tok"]))
    assert r1.status_code == 409
    r2 = env["cl"].get(f"/platform/studies/{s['id']}/report.pdf",
                       headers=hdr(env["tok"]))
    assert r2.status_code == 429
    assert r2.json()["detail"]["error"] == "pdf_throttled"


# ── ٨ب) التقرير PDF — طلب المالك الحرفي: «بالنسبة للتقرير اريده pdf» ─────────
def test_report_pdf_409_before_completion(env):
    s = _mk_study(env["cl"], env["tok"])
    r = env["cl"].get(f"/platform/studies/{s['id']}/report.pdf",
                      headers=hdr(env["tok"]))
    assert r.status_code == 409
    assert r.json()["detail"]["error"] == "no_report_yet"


def test_report_pdf_real_conversion_end_to_end(env):
    """المسار الحقيقي كاملاً: حفظ نتيجة عبر silk_storage الحقيقي (قاعدة مؤقتة)
    ← build_view ← render_docx ← docx_to_pdf — الجسم يبدأ بـ%PDF فعلاً.

    بلا soffice تُتخطّى (نفس تساهل test_report_output_overhaul) — لا ادعاء
    فحص تحويلٍ لم يجرِ.
    """
    import os
    import tempfile

    import pytest
    import silk_reports
    from tests.pdf_gate import pdf_gate       # R8 (TEST-5): البوّابةُ الواحدة
    pdf_gate("تحويل PDF على سطح المصنع")
    # مسبار قدرة (اتفاق test_report_output_overhaul): soffice موجود لكن بلا
    # مرشّح Writer (بيئة core-فقط) لا يحمّل docx أصلاً — تخطٍّ معلن لا حمرة
    # بيئية، والفشل الحقيقي بعد المسبار يبقى فشلاً.
    _pt = os.path.join(tempfile.mkdtemp(), "probe.docx")
    from silk_platform.engine_bridge import _fake_result
    from silk_render import build_view as _bv
    silk_reports.render_docx(_bv(_fake_result("مسبار", "080410")), _pt)
    try:
        silk_reports.docx_to_pdf(_pt)
    except RuntimeError as _exc:
        from tests.pdf_gate import pdf_engine_broken
        pdf_engine_broken(str(_exc))
    from silk_storage import save_analysis
    aid = save_analysis(_fake_result("تمور سكري", "080410"), None)
    assert aid
    s = _mk_study(env["cl"], env["tok"])
    _complete_study(env, s["id"], analysis_id=int(aid))
    import tempfile
    tmp_root = tempfile.gettempdir()

    def _silk_tmp():
        return {d for d in os.listdir(tmp_root)
                if d.startswith(("silk_pdf_", "silk_lo_", "silk_pdfsrc_",
                                 "silk_acad_"))}

    before_tmp = _silk_tmp()
    r = env["cl"].get(f"/platform/studies/{s['id']}/report.pdf",
                      headers=hdr(env["tok"]))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/pdf")
    assert r.content[:5] == b"%PDF-"
    assert len(r.content) > 1000
    # §58 M5: لا مجلد مؤقت متروك (كان يبقى docx كامل من تقرير العميل +
    # بروفايل LibreOffice في /tmp حتى إعادة التشغيل).
    leaked = _silk_tmp() - before_tmp
    assert not leaked, f"مجلدات مؤقتة متروكة بعد التحويل: {leaked}"


def test_report_pdf_declared_503_when_converter_unavailable(env, monkeypatch):
    """غياب محرّك التحويل = 503 برمز `pdf_unavailable` معلَن — لا docx بديل
    صامت ولا 500 غامضة (نفس عقد §3 الجذري)."""
    import silk_reports
    from silk_platform.engine_bridge import _fake_result
    from silk_storage import save_analysis
    aid = save_analysis(_fake_result("تمور سكري", "080410"), None)
    s = _mk_study(env["cl"], env["tok"])
    _complete_study(env, s["id"], analysis_id=int(aid))
    monkeypatch.setattr(silk_reports, "_find_soffice", lambda: None)
    r = env["cl"].get(f"/platform/studies/{s['id']}/report.pdf",
                      headers=hdr(env["tok"]))
    assert r.status_code == 503
    detail = r.json()["detail"]
    assert detail["error"] == "pdf_unavailable"
    # الرسالة معلنة بلا تسريب مسارات/أسرار (عائلة test_report_output_overhaul).
    assert "PDF" in detail["message"]


@pytest.mark.parametrize("exc_name,code", [
    ("PdfBracketGateError", "pdf_rejected"),
    ("PdfConversionFailed", "pdf_failed"),
])
def test_report_pdf_names_the_failure_and_leaves_a_trace(env, monkeypatch,
                                                         caplog, exc_name, code):
    """البند ٢٨٤ — بلاغ المالك: «توليد ملف PDF معطَّل على الخادم» يتكرّر بينما
    المحرّكُ سليم؛ الذي رفض الملفَّ فحصُ الأقواس بعد تحويلٍ ناجح. رفضُ الفحص
    وفشلُ التحويل رمزان غيرُ `pdf_unavailable`، وكلاهما يترك سطراً في السجلّ
    وصفّاً في `/ops/last-errors` (كان المسارُ صامتاً تماماً)."""
    import logging

    import silk_ops_log
    import silk_reports
    from silk_platform.engine_bridge import _fake_result
    from silk_storage import save_analysis
    aid = save_analysis(_fake_result("تمور سكري", "080410"), None)
    s = _mk_study(env["cl"], env["tok"])
    _complete_study(env, s["id"], analysis_id=int(aid))
    exc = getattr(silk_reports, exc_name)

    def _boom(view, path):
        raise exc("سبب داخلي")

    monkeypatch.setattr(silk_reports, "render_research_pdf", _boom)
    monkeypatch.setattr(silk_reports, "render_client_pdf", _boom)
    recorded = []
    monkeypatch.setattr(silk_ops_log, "record_error",
                        lambda kind, reason, context=None, path=None:
                        recorded.append((kind, context)))
    with caplog.at_level(logging.WARNING):
        r = env["cl"].get(f"/platform/studies/{s['id']}/report.pdf",
                          headers=hdr(env["tok"]))
    assert r.status_code == 503
    assert r.json()["detail"]["error"] == code
    assert any(code in rec.getMessage() for rec in caplog.records)
    assert recorded and recorded[0][0] == "pdf_export_failure"
    assert recorded[0][1]["code"] == code


# ── ٩) المقعد الاختباري خامل بلا env — the fake seam is inert unset ──────────
def test_fake_engine_seam_inert_without_env(env, monkeypatch):
    import silk_platform.engine_bridge as eb
    monkeypatch.delenv("SILK_PLATFORM_FAKE_ENGINE", raising=False)
    assert eb.fake_engine_enabled() is False


def test_fake_engine_persists_through_real_storage(env, monkeypatch, tmp_path):
    import silk_platform.engine_bridge as eb
    monkeypatch.setenv("SILK_PLATFORM_FAKE_ENGINE", "1")
    monkeypatch.setenv("SILK_DB", str(tmp_path / "silk.db"))
    s = _mk_study(env["cl"], env["tok"], hs_code="080410")
    r = env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    assert r.status_code == 200
    assert eb.wait_idle(20)
    row = _study_row(s["id"])
    assert row["state"] == "completed" and row["analysis_id"]
    import silk_storage
    found = silk_storage.get_analysis(int(row["analysis_id"]),
                                      str(tmp_path / "silk.db"))
    assert found is not None
    assert "عيّنة اختبار" in (found.get("note") or "")   # موسومة، لا تُقدَّم حيّة


# ── ١٠) صورة → منتج/رمز — honest declared gaps, never fabricated ─────────────
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def _upload_png(env):
    r = env["cl"].post("/platform/images",
                       files={"file": ("p.png", _TINY_PNG, "image/png")},
                       headers=hdr(env["tok"]))
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_classify_image_valve_off_is_declared(env, monkeypatch):
    monkeypatch.delenv("SILK_IMAGE_INTAKE", raising=False)
    iid = _upload_png(env)
    r = env["cl"].post("/platform/classify-image", json={"image_id": iid},
                       headers=hdr(env["tok"]))
    assert r.status_code == 200
    out = r.json()
    assert out["ok"] is False and out["status"] == "intake_disabled"
    assert "SILK_IMAGE_INTAKE" in out["message"]


def test_classify_image_no_key_reads_failed_honestly(env, monkeypatch):
    """بلا مفتاح كلود: «تعذّرت القراءة» الصادقة — صفر نداء شبكة، صفر اختلاق."""
    monkeypatch.setenv("SILK_IMAGE_INTAKE", "1")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    iid = _upload_png(env)
    r = env["cl"].post("/platform/classify-image", json={"image_id": iid},
                       headers=hdr(env["tok"]))
    assert r.status_code == 200
    out = r.json()
    assert out["ok"] is False and out["status"] == "read_failed"
    assert out["product_name"] == ""            # لا اسم مختلَق
    assert "hs" not in out                       # ولا رمز مقترح من عدم


def test_classify_image_cross_tenant_404(env, monkeypatch):
    other = make_factory("silver", "other-pivot@f.local")
    otok = login(env["cl"], other["email"], other["password"])
    iid = _upload_png(env)
    r = env["cl"].post("/platform/classify-image", json={"image_id": iid},
                       headers=hdr(otok))
    assert r.status_code == 404                  # لا تسريب وجود عبر المستأجرين


# ── ١١) إشراف الأدمِن — admin oversight of all studies ───────────────────────
def test_admin_studies_oversight(env, monkeypatch):
    import silk_platform.engine_bridge as eb
    monkeypatch.setattr(eb, "_run_engine", lambda *a, **_k: _canned_result(880))
    s = _mk_study(env["cl"], env["tok"], market_pref="ARE")
    env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    assert eb.wait_idle(15)
    atok = login(env["cl"], "admin@silk.local", "AdminPass1")
    r = env["cl"].get("/platform/admin/studies", headers=hdr(atok))
    assert r.status_code == 200
    rows = r.json()["studies"]
    mine = [x for x in rows if x["id"] == s["id"]]
    assert mine and mine[0]["product"] == "تمور سكري"
    assert mine[0]["state"] == "completed"
    assert mine[0]["account_name"]
    import json as _json
    assert "@" not in _json.dumps(rows, ensure_ascii=False)   # لا بريد — جدار PII
    # المصنع لا يصل نقطة الأدمِن.
    r = env["cl"].get("/platform/admin/studies", headers=hdr(env["tok"]))
    assert r.status_code == 403


# ── ١٢) حقول الطلب متحقَّق منها — validated research fields ──────────────────
def test_market_pref_and_hs_validation(env):
    r = env["cl"].post("/platform/studies",
                       json={"product": "تمر", "market_pref": "دبي"},
                       headers=hdr(env["tok"]))
    assert r.status_code == 422
    r = env["cl"].post("/platform/studies",
                       json={"product": "تمر", "hs_code": "abc"},
                       headers=hdr(env["tok"]))
    assert r.status_code == 422
    r = env["cl"].post("/platform/studies",
                       json={"product": "تمر", "market_pref": "are",
                             "hs_code": "080410"},
                       headers=hdr(env["tok"]))
    assert r.status_code == 200
    body = r.json()
    assert body["market_pref"] == "ARE" and body["hs_code"] == "080410"
    assert body["title_ar"] == "تمر"            # عنوان تلقائي من المنتج


def test_image_binding_must_be_owned(env):
    other = make_factory("silver", "img-owner@f.local")
    otok = login(env["cl"], other["email"], other["password"])
    r = env["cl"].post("/platform/images",
                       files={"file": ("p.png", _TINY_PNG, "image/png")},
                       headers=hdr(otok))
    foreign_iid = r.json()["id"]
    r = env["cl"].post("/platform/studies",
                       json={"product": "تمر", "image_id": foreign_iid},
                       headers=hdr(env["tok"]))
    assert r.status_code == 422                  # صورة حسابٍ آخر تُرفَض


# ══════ أقفال إصلاحات المراجعة الذاتية §58 (موجة التحوّل) ═════════════════════
def test_unknown_market_rejected_at_input(env):
    """سوق مجهول يُرفَض 422 معلناً — لا تقرير «كل الأسواق» موسوماً زوراً بسوقٍ
    لم يُستهدف (ملاحظة §58: XYZ كان يُقبل ثم يُسقَط صامتاً في الجسر)."""
    r = env["cl"].post("/platform/studies",
                       json={"product": "تمر", "market_pref": "XYZ"},
                       headers=hdr(env["tok"]))
    assert r.status_code == 422
    assert r.json()["detail"]["error"] == "unknown_market"
    # وسوق عالمي حقيقي خارج قائمة الأسواق-١ يُقبل ويُستهدف فعلاً.
    r = env["cl"].post("/platform/studies",
                       json={"product": "تمر", "market_pref": "JPN"},
                       headers=hdr(env["tok"]))
    assert r.status_code == 200
    from silk_platform.engine_bridge import _target_countries
    hit = _target_countries("JPN")
    assert hit and hit[0]["iso3"] == "JPN" and hit[0]["m49"] == "392"


def test_image_binding_can_be_cleared_with_null(env):
    """`image_id: null` يفكّ الربط فعلاً (§58: كان يُتجاهل صامتاً)."""
    iid = _upload_png(env)
    s = _mk_study(env["cl"], env["tok"], image_id=iid)
    assert s["image_id"] == iid
    r = env["cl"].patch(f"/platform/studies/{s['id']}",
                        json={"image_id": None}, headers=hdr(env["tok"]))
    assert r.status_code == 200
    assert r.json()["image_id"] is None


def test_classify_rejects_invalid_image_without_burning_the_daily_cap(env, monkeypatch):
    """صورة تتجاوز سقف الرؤية تُرفَض **قبل** حجز تفعيلة السقف اليومي (§58)."""
    import silk_platform.engine_bridge as eb
    monkeypatch.setenv("SILK_IMAGE_INTAKE", "1")
    reserved = []
    monkeypatch.setattr(eb, "_vision_allowed",
                        lambda: (reserved.append(1) or (True, "")))
    big = b"\x89PNG\r\n\x1a\n" + b"\x00" * (6 * 1024 * 1024)
    out = eb.classify_image_flow(big, "image/png")
    assert out["ok"] is False and out["status"] == "invalid_image"
    assert not reserved, "حُجزت تفعيلة سقفٍ على صورة مرفوضة"


def test_orphan_sweep_grace_window_skips_fresh_runs(env):
    """تشغيلة أحدث من نافذة السماح لا تُحكَم يتيمةً (§58: النشر المتداخل)."""
    import silk_platform.engine_bridge as eb
    from silk_platform import db as pdb
    from silk_platform.db import now_iso
    conn = pdb.connect()
    try:
        conn.execute(
            "INSERT INTO studies (owner_id, state, launched_at, "
            "launched_by_user_id, run_started_at, created_at, updated_at) "
            "VALUES (?, 'in_progress', ?, ?, ?, '', '')",
            (env["fac"]["account_id"], now_iso(), env["fac"]["user_id"],
             now_iso()))
        sid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        conn.commit()
        swept = eb.sweep_orphans(conn)
    finally:
        conn.close()
    assert swept["skipped"] >= 1
    assert _study_row(sid)["state"] == "in_progress"


# ═══════════ رمز المحاولة — لا كتابةَ خيطٍ زومبي فوق إطلاقةٍ جديدة ═══════════
def _force_in_progress(sid: int, token: str) -> None:
    """اضبط الصفَّ «قيد الإعداد» برمز محاولةٍ بعينه — تثبيت حالة، بلا خيوط."""
    from silk_platform import db as pdb
    conn = pdb.connect()
    try:
        conn.execute("UPDATE studies SET state = 'in_progress', "
                     "run_token = ?, analysis_id = NULL, run_error = NULL "
                     "WHERE id = ?", (token, sid))
        conn.commit()
    finally:
        conn.close()


def test_a_swept_runs_thread_cannot_complete_a_newer_launch(env):
    """LESSONS ٧٩ (مراجعة ذاتية §58) — الإنهاء مملوكٌ لمحاولته وحدها.

    السيناريو الذي كان يكسر: تشغيلةٌ تجاوزت مهلتها فكُنست وعادت الدراسة
    مسودّةً، فأعاد المصنع الإطلاق — ثم استيقظ خيطُ التشغيلة القديمة فوجد
    الصفّ `in_progress` ثانيةً وكتب عليه **معرّف تحليله القديم**: تقريرٌ خاطئ
    لدراسةٍ جارية. الرمز يجعل الكتابة مشروطةً بصاحبها.
    """
    import silk_platform.engine_bridge as eb
    s = _mk_study(env["cl"], env["tok"])
    _force_in_progress(s["id"], "NEW-TOKEN")

    eb._finish_success(s["id"], env["fac"]["account_id"], 4242, run_token="OLD-TOKEN")
    row = _study_row(s["id"])
    assert row["state"] == "in_progress", "الزومبي وسم دراسةً جارية «مكتملة»"
    assert row["analysis_id"] is None, "الزومبي كتب معرّف تحليله القديم"

    eb._finish_success(s["id"], env["fac"]["account_id"], 4242, run_token="NEW-TOKEN")
    row = _study_row(s["id"])
    assert row["state"] == "completed" and row["analysis_id"] == 4242, row


def test_a_swept_runs_thread_cannot_fail_a_newer_launch(env):
    """والوجه الآخر: الزومبي لا يُفشل تشغيلةً سليمة ولا يُرجِع حصّتها."""
    import silk_platform.engine_bridge as eb
    s = _mk_study(env["cl"], env["tok"])
    _force_in_progress(s["id"], "NEW-TOKEN")

    eb._finish_failure(s["id"], env["fac"]["account_id"], "سبب قديم",
                       run_token="OLD-TOKEN")
    row = _study_row(s["id"])
    assert row["state"] == "in_progress", "الزومبي أعاد دراسةً جارية مسودّةً"
    assert row["run_error"] is None, row["run_error"]

    eb._finish_failure(s["id"], env["fac"]["account_id"], "سبب حقيقي",
                       run_token="NEW-TOKEN")
    row = _study_row(s["id"])
    assert row["state"] == "draft" and row["run_error"] == "سبب حقيقي", row


# ═══════════ جرس الأدمِن لم يعد فارغاً بنيوياً ═══════════════════════════════
def test_admin_bell_mirrors_factory_study_outcomes(env, monkeypatch):
    """قرار 2026-08-19: إشرافٌ بلا إشعار = اكتشافُ التعثّر بالمصادفة.

    الجسر كان يكتب الإشعار لحساب المصنع وحده والقراءة مقيّدة بالحساب، فجرسُ
    الأدمِن يفتح على «لا إشعارات بعد» دائماً مهما جرى في المنصّة.
    """
    from tests.platform_helpers import mock_engine
    from silk_platform import db as pdb
    eb = mock_engine(monkeypatch, analysis_id=991)
    s = _mk_study(env["cl"], env["tok"])
    env["cl"].post(f"/platform/studies/{s['id']}/launch", headers=hdr(env["tok"]))
    assert eb.wait_idle(15)

    conn = pdb.connect()
    try:
        vault = conn.execute(
            "SELECT id FROM accounts WHERE is_vault = 1").fetchone()["id"]
        rows = [dict(r) for r in conn.execute(
            "SELECT account_id, title, study_id FROM platform_notifications "
            "ORDER BY id")]
    finally:
        conn.close()
    factory_notes = [r for r in rows if r["account_id"] == env["fac"]["account_id"]]
    admin_notes = [r for r in rows if r["account_id"] == vault]
    assert len(factory_notes) == 1, rows
    assert len(admin_notes) == 1, rows          # المرآة موجودة
    assert admin_notes[0]["study_id"] == s["id"]
    # ويُعرَف صاحبُ الخبر من العنوان مباشرةً (اسم الحساب مُلحَق).
    assert env["fac"]["account_id"] and "Factory-silver" in admin_notes[0]["title"]
