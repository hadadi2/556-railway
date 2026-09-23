"""تقرير سِلك ٧ — حساسيةُ أقصى سعر المصنع وتمويلُ مخزون التجربة (§3.5، الدرس ٢٨٣).

The maximum ex-works price is shown against a ±10% shelf price, ±5 freight
points and ±10% exchange rate (only when an observed rate is used); trial
inventory financing is its own measure — entry cost × annual rate × days to
payment ÷ 365 — from two factory inputs wired end to end, and a named gap
without them.

كلُّ الحالات مصطنعةٌ لاختبارٍ هرمتيّ.
"""
import pathlib

import pytest
from docx import Document

import silk_economics as E
import silk_reports as SR

_ROOT = pathlib.Path(__file__).resolve().parents[1]


def test_sensitivity_follows_the_same_formula():
    rs = E.reverse_solve_max_exw(30.0, tariff_pct=0, vat_pct=None)
    rows = E.sensitivity_rows(rs)
    mx = rs["max_exw"]
    shelf = [r["max_exw"] for r in rows if r["factor"] == "shelf"]
    assert shelf == [pytest.approx(mx * 0.9, rel=1e-4),
                     pytest.approx(mx * 1.1, rel=1e-4)]
    fr = [r for r in rows if r["factor"] == "freight"]
    assert fr[0]["max_exw"] < mx < fr[1]["max_exw"]
    assert not [r for r in rows if r["factor"] == "fx"], "لا صرفَ بلا مرصود"
    rs["fx"] = {"rate": 4.5, "pair": "MYR/USD"}
    fx = [r for r in E.sensitivity_rows(rs) if r["factor"] == "fx"]
    assert fx and fx[0]["currency"] == "USD" and fx[0]["pair"] == "MYR/USD"


def test_the_client_report_prints_the_sensitivity_table():
    rs = E.reverse_solve_max_exw(30.0, tariff_pct=0, vat_pct=None)
    rs.update({"currency": "MYR", "unit": "كجم"})
    rs["sensitivity"] = E.sensitivity_rows(rs)
    doc = Document()
    SR._client_economics_section(doc, {"economics": {"reverse_solve": rs,
                                                     "anchor_price": {}}}, "ar")
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "حساسية أقصى سعر المصنع" in heads


def test_financing_is_a_separate_measure_from_two_inputs():
    by = lambda **kw: {e["name"]: e for e in E.build_decision_numbers(
        category="حليب", cost_per_unit=2.0, cost_currency="SAR", **kw)}
    gap = by()[E.FINANCING_NAME]
    assert gap["tier"] == "gap" and "معدل التمويل" in gap["missing"] \
        and "مدة التحصيل" in gap["missing"]
    d = by(financing_rate_pct=8, collection_days=90)
    fin, entry = d[E.FINANCING_NAME], d["كلفة الدخول الكلية حتى أول شحنة"]
    assert fin["tier"] == "estimated"
    assert fin["range"]["low"] == pytest.approx(
        entry["range"]["low"] * 0.08 * 90 / 365, rel=1e-3)


def test_financing_inputs_are_wired_end_to_end():
    mig = (_ROOT / "migrations/platform/026_product_financing.sql").read_text()
    assert "financing_rate_pct" in mig and "collection_days" in mig
    from silk_platform import engine_bridge as B
    card = B.product_card_from_row({"cost_per_unit": 2.0,
                                    "financing_rate_pct": 8,
                                    "collection_days": 90})
    assert card["financing_rate_pct"] == 8 and card["collection_days"] == 90
    html = (_ROOT / "web/platform.html").read_text()
    assert 'name="pfinrate"' in html and "collection_days: num(" in html


def test_missing_components_are_one_short_condition():
    import silk_render as R
    conds = R._flip_conditions("conditional", False, {}, "ماليزيا",
                               missing_components=[f"c{n}" for n in range(10)])
    comp = [c for c in conds if "إعادة تقييم القرار" in c["condition"]]
    assert len(comp) == 1 and "7 مكوّنات أخرى" in comp[0]["condition"]
    assert comp[0]["condition"].count("إعادة تقييم") == 1
    line = SR._T("gap_flip_condition", "ar", condition=comp[0]["condition"],
                 closes_via=comp[0]["closes_via"])
    assert len(line) < 200


# ── مراجعة §58 ──────────────────────────────────────────────────────────────

def test_no_sensitivity_for_a_ceiling_withheld_as_a_negotiating_basis():
    rs = E.reverse_solve_max_exw(30.0, tariff_pct=0, vat_pct=None)
    rs.update({"currency": "MYR", "unit": "كجم"})
    rs["sensitivity"] = E.sensitivity_rows(rs)
    doc = Document()
    SR._client_economics_section(doc, {"economics": {
        "reverse_solve": rs, "anchor_price": {},
        "pricing_contradiction": {"shortfall_pct": 97.9, "max_exw_usd": 0.3,
                                  "reference_import_price_usd_kg": 14.0}}},
        "ar")
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "حساسية أقصى سعر المصنع" not in heads


def test_sensitivity_cells_carry_the_basis_and_arabic_units():
    rs = E.reverse_solve_max_exw(30.0, tariff_pct=0, vat_pct=None)
    rs.update({"currency": "MYR", "unit": "كجم"})
    rs["sensitivity"] = E.sensitivity_rows(rs)
    doc = Document()
    SR._client_economics_section(doc, {"economics": {"reverse_solve": rs,
                                                     "anchor_price": {}}}, "ar")
    t = next(t for t in doc.tables if t.rows[0].cells[0].text == "العامل")
    text = " ".join(c.text for r in t.rows[1:] for c in r.cells)
    assert "MYR/كجم" in text and "نقاط" in text and " pt" not in text


def test_the_financing_gap_names_the_real_blocker():
    fin = {e["name"]: e for e in E.build_decision_numbers(
        category="فئة بلا ثابت", cost_per_unit=2.0, financing_rate_pct=8,
        collection_days=90)}[E.FINANCING_NAME]
    assert fin["tier"] == "gap" and "حجم شحنة" in fin["missing"]
