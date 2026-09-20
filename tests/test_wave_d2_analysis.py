"""الموجة د-٢ — التحليل التجاري: ميزةُ التكلفة، طبيعةُ المنافسين، التمايز،
توثيقُ الأسعار، ربطُ الاشتراط بدرجة التصنيع، الضريبةُ من صفٍّ رسمي، وادّعاءاتُ
القنوات بلا مصدر.

بلاغُ المالك (2026-09-19، الموجة د): «كل دراسة تجيب أسئلة المصدّر الحقيقية
وتستنتج من البيانات… القواعد عامة مدفوعة بالبيانات وبفئة المنتج من فصل HS…
كل نتيجة تدخل سجل الحقائق بدرجتها… لا تقدير بلا افتراض معلن، ولا قاعدة خاصة
بمنتج أو سوق».

test-first: كُتب قبل التنفيذ. هرمتي: صفر شبكة (كومتريد يُرقَّع)، صفر مفتاح،
صفر إنفاق. Run: python3 -m pytest tests/test_wave_d2_analysis.py -q
"""
import importlib
import os
import re
import sys
import tempfile

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for _p in (_ROOT, os.path.join(_ROOT, "tools")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import silk_commercial_analysis as CA                         # noqa: E402
import silk_economics as E                                    # noqa: E402
import silk_fact_ledger as L                                  # noqa: E402
import silk_render as R                                       # noqa: E402
import silk_reports as SR                                     # noqa: E402
from conftest import block_network                            # noqa: E402
from gen_verdict_baseline import CANONICAL_BLOBS              # noqa: E402
from silk_agents import AgentReport                           # noqa: E402
from silk_data_layer import DataPoint                         # noqa: E402
from silk_requirements_agent import applies_to                # noqa: E402


def _blob(key):
    mod, fn = CANONICAL_BLOBS[key]
    return getattr(importlib.import_module(mod), fn)()


def _view(key):
    return R.build_view(_blob(key))


def _signal_of(findings, kind):
    for dp in findings:
        v = getattr(dp, "value", None)
        if isinstance(v, dict) and v.get("kind") == kind:
            return v
    return None


class _Market:
    iso3, m49, name_en = "TUR", "792", "Türkiye"


@pytest.fixture
def no_store(monkeypatch):
    """مخزنٌ فارغ + ميزانيةٌ متاحة — كلُّ نداءٍ يذهب إلى كومتريد المُرقَّع."""
    import silk_collectors
    import silk_store
    monkeypatch.setattr(silk_store, "get_trade_flow", lambda *a, **k: None)
    monkeypatch.setattr(silk_store, "upsert_trade_flows", lambda rows: 0)
    monkeypatch.setattr(silk_collectors, "comtrade_budget_left", lambda: 100)
    monkeypatch.delenv(CA.MAX_CALLS_ENV, raising=False)


def _fake_comtrade(table):
    """`comtrade_trade(hs, reporter, year, flow=, partner=)` من جدولٍ ثابت:
    المفتاح (hs, reporter, flow, partner) → قائمةُ سجلات أو None (تعذُّر)."""
    calls = []

    def _fn(hs, reporter, year, flow="M", partner=0, **kw):
        calls.append((hs, reporter, flow, partner))
        return table.get((hs, reporter, flow, partner), [])
    _fn.calls = calls
    return _fn


def _rec(value, **extra):
    return {"primaryValue": value, "netWgt": None, **extra}


# ── ١) ميزةُ التكلفة — صافي تجارة المدخلات الخام ─────────────────────────────

def test_raw_inputs_come_from_the_data_file_by_chapter():
    spec = CA.raw_input_codes("390210")
    assert spec and spec["codes"] == ["2901", "2711", "2902"]
    assert CA.raw_input_codes("080410") is None          # فصلٌ بلا صفّ = فجوة
    assert CA.raw_input_codes("") is None
    for ch, spec in CA.raw_input_rows().items():
        assert len(spec["codes"]) <= 3 and spec["source"], ch


def test_saudi_net_trade_signal_is_stored_as_a_structured_finding(no_store, monkeypatch):
    import silk_data_layer
    table = {}
    for code, x, m in (("2901", 3.9e9, 1.2e8), ("2711", 9.8e9, 2.1e8),
                       ("2902", 2.6e9, 3.4e8)):
        table[(code, "682", "X", 0)] = [_rec(x)]
        table[(code, "682", "M", 0)] = [_rec(m)]
    fake = _fake_comtrade(table)
    monkeypatch.setattr(silk_data_layer, "comtrade_trade", fake)
    rep = AgentReport("LLMMissionAgent", [], summary="")
    with block_network():
        CA.augment_raw_input_trade(rep, "390210", year=2024)
    sig = _signal_of(rep.findings, "raw_input_trade")
    assert sig and sig["net_usd"] == pytest.approx(16.3e9 - 0.67e9)
    assert sig["observed_codes"] == 3 and sig["calls_used"] == 6
    assert CA.MARK in rep.findings[-1].note
    # idempotent — تشغيلةٌ ثانية لا تكرّر ولا تنادي.
    CA.augment_raw_input_trade(rep, "390210", year=2024)
    assert len(rep.findings) == 1 and len(fake.calls) == 6


def test_fetch_failure_is_a_declared_gap_not_a_zero(no_store, monkeypatch):
    import silk_data_layer
    monkeypatch.setattr(silk_data_layer, "comtrade_trade",
                        _fake_comtrade({("2901", "682", "X", 0): None}))
    rep = AgentReport("LLMMissionAgent", [], summary="")
    with block_network():
        CA.augment_raw_input_trade(rep, "390210", year=2024)
    sig = _signal_of(rep.findings, "raw_input_trade")
    assert sig["net_usd"] is None and sig["exports_usd"] is None
    assert rep.findings[-1].confidence == 0.0
    assert "fetch_failed" in sig["rows"][0]["status"]
    assert "no_record" in sig["rows"][1]["status"]


def test_no_budget_means_no_call_and_a_named_gap(monkeypatch):
    import silk_collectors
    import silk_data_layer
    monkeypatch.setattr(silk_collectors, "comtrade_budget_left", lambda: 0)
    fake = _fake_comtrade({})
    monkeypatch.setattr(silk_data_layer, "comtrade_trade", fake)
    rep = AgentReport("LLMMissionAgent", [], summary="")
    CA.augment_raw_input_trade(rep, "390210", year=2024)
    assert fake.calls == []
    assert rep.findings[-1].value is None
    assert "ميزانية" in rep.findings[-1].note


def test_call_cap_is_owner_bounded_at_fifteen(monkeypatch):
    monkeypatch.setenv(CA.MAX_CALLS_ENV, "40")
    assert CA.max_calls() == 15
    monkeypatch.setenv(CA.MAX_CALLS_ENV, "x")
    assert CA.max_calls() == 12


# ── ٢) طبيعةُ المنافسين — منتجٌ أم معيدُ تصدير ─────────────────────────────

def _competitor_summary(*tops):
    return DataPoint({"year": 2024, "hhi": 1560, "supplier_count": 3,
                      "top_suppliers": [{"partner": p, "share": s}
                                        for p, s in tops]},
                     "UN Comtrade", 0.9, "HHI", "2026-09-01", data_year=2024)


def test_reexporter_rule_is_declared_and_names_the_actual_source(no_store, monkeypatch):
    import silk_data_layer
    table = {
        # صادراتُ العالم: كوريا وألمانيا من أعلى ٥؛ بلجيكا لا.
        ("390210", None, "X", 0): [_rec(5e9, reporterISO="KOR", reporterCode="410"),
                                   _rec(4e9, reporterISO="DEU", reporterCode="276"),
                                   _rec(3e9, reporterISO="USA", reporterCode="840"),
                                   _rec(2e9, reporterISO="SAU", reporterCode="682"),
                                   _rec(1e9, reporterISO="SGP", reporterCode="702"),
                                   _rec(2e8, reporterISO="BEL", reporterCode="056")],
        ("390210", "410", "X", 0): [_rec(5e9)],
        ("390210", "410", "M", 0): [_rec(3e8)],
        ("390210", "056", "X", 0): [_rec(2.1e8)],
        ("390210", "056", "M", 0): [_rec(2.6e8)],
        ("390210", "056", "M", "all"): [_rec(9e7, partnerCode="276"),
                                        _rec(4e7, partnerCode="528"),
                                        _rec(2.6e8, partnerCode="0")],
    }
    fake = _fake_comtrade(table)
    monkeypatch.setattr(silk_data_layer, "comtrade_trade", fake)
    rep = AgentReport("LLMMissionAgent",
                      [_competitor_summary(("Saudi Arabia", 31.0),
                                           ("Rep. of Korea", 14.0),
                                           ("Belgium", 9.0))], summary="")
    with block_network():
        CA.augment_supplier_nature(rep, "390210", _Market(), year=2024)
    sig = _signal_of(rep.findings, "supplier_nature")
    assert sig and [r["partner"] for r in sig["rows"]] == ["Rep. of Korea", "Belgium"]
    kor, bel = sig["rows"]
    assert kor["world_top_exporter"] is True and kor["reexporter_likely"] is False
    assert bel["world_top_exporter"] is False and bel["reexporter_likely"] is True
    assert bel["actual_source"]["iso3"] == "DEU"
    assert bel["actual_source"]["share"] == pytest.approx(100 * 9e7 / 2.6e8, abs=0.1)
    assert "أعلى 5" in rep.findings[-1].note and "ووارداته ≥ صادراته" in rep.findings[-1].note
    # ≤ ١ + ٣×٣ نداءات (البند ٢) — هنا ٦.
    assert sig["calls_used"] == len(fake.calls) == 6 <= CA.max_calls()


def test_unknown_partner_is_skipped_not_guessed(no_store, monkeypatch):
    import silk_data_layer
    fake = _fake_comtrade({("390210", None, "X", 0): [
        _rec(5e9, reporterISO="KOR", reporterCode="410")]})
    monkeypatch.setattr(silk_data_layer, "comtrade_trade", fake)
    rep = AgentReport("LLMMissionAgent",
                      [_competitor_summary(("Areas, nes", 40.0))], summary="")
    with block_network():
        CA.augment_supplier_nature(rep, "390210", _Market(), year=2024)
    sig = _signal_of(rep.findings, "supplier_nature")
    row = sig["rows"][0]
    assert row["iso3"] is None and row["status"] == "unknown_country"
    assert row["reexporter_likely"] is None
    assert rep.findings[-1].confidence == 0.0
    assert len(fake.calls) == 1               # لا نداءَ لمورّدٍ مجهول الرمز


def test_without_a_structured_summary_the_nature_is_a_declared_gap(no_store):
    rep = AgentReport("LLMMissionAgent", [], summary="")
    CA.augment_supplier_nature(rep, "390210", _Market())
    assert rep.findings[-1].value is None
    assert "لا ملخّص مورّدين مهيكل" in rep.findings[-1].note


# ── الاستنتاجاتُ تدخل السجلَّ بدرجتها (نقيّة، من الاكتشافات المخزَّنة) ────────

def test_cost_advantage_is_an_inference_with_assumption_and_flip():
    led = _view("turkey_polymers")["ledger"]["entries"]
    assert led["saudi_raw_input_balance"]["status"] == L.OBSERVED
    ca = led["cost_advantage"]
    assert ca["status"] == L.INFERENCE and ca["value"] == "ميزة"
    assert "saudi_raw_input_balance" in ca["basis"]
    assert "مؤشّرٌ للإنتاج لا قياسٌ له" in ca["assumption"]
    assert ca["flip_if"]
    txt = L.render_value(ca, "ar")
    assert "يشير إليه" in txt and "ميزة" in txt


def test_a_net_importer_of_inputs_is_never_given_the_advantage():
    led = _view("netherlands_honey")["ledger"]["entries"]
    assert led["saudi_raw_input_balance"]["value"] < 0
    assert led["cost_advantage"]["value"] == "عيب"


def test_a_reexporter_is_excluded_from_direct_competitors():
    led = _view("netherlands_honey")["ledger"]["entries"]
    sn = led["supplier_nature"]
    assert sn["status"] == L.INFERENCE and "1 معيد تصدير" in sn["value"]
    bel = next(r for r in sn["items"] if r["iso3"] == "BEL")
    assert bel["reexporter_likely"] is True and bel["actual_source"]["iso3"] == "UKR"
    direct = led["direct_competitor_share_pct"]
    assert direct["status"] == L.INFERENCE and direct["value"] == pytest.approx(34.0)
    assert "على مستوى الدول" in direct["note"]


def test_without_signals_the_insights_are_declared_gaps_not_zeros():
    led = _view("netherlands")["ledger"]["entries"]
    for key in ("saudi_raw_input_balance", "cost_advantage", "supplier_nature",
                "direct_competitor_share_pct"):
        assert led[key]["status"] == L.MISSING and led[key]["value"] is None, key
        assert "لم تُشغَّل إشارات التحليل التجاري" in led[key]["note"]


def test_missing_insight_keys_are_hidden_from_the_writer_and_the_gaps_table():
    view = _view("netherlands")
    block = L.facts_block(view["ledger"])
    for key in ("cost_advantage", "supplier_nature", "differentiation"):
        assert key not in block, key
        assert view["ledger"]["entries"][key]["status"] == L.MISSING
    assert not any(r["key"] == "cost_advantage" for r in view["ledger"]["gaps"])
    # وحين تُحسَب تظهر للكاتب بصيغتها الطبيعية.
    block = L.facts_block(_view("turkey_polymers")["ledger"])
    assert "ميزة التكلفة" in block and "يشير إليه" in block


# ── ٤) التمايزُ الحقيقي ──────────────────────────────────────────────────────

def test_an_attribute_a_rival_owns_is_not_declared_a_differentiator():
    d = _view("netherlands_honey")["ledger"]["entries"]["differentiation"]
    assert d["status"] == L.INFERENCE
    assert "عسل سدر" in d["value"] and "premium" in d["value"]
    assert "organic" not in d["value"] and "organic" in d["note"]
    assert d["basis"] == ["competitor_prices"]


def test_without_card_attributes_differentiation_is_a_declared_gap():
    d = _view("turkey_polymers")["ledger"]["entries"]["differentiation"]
    assert d["status"] == L.MISSING and "بطاقة المنتج" in d["note"]


def test_all_attributes_owned_by_rivals_reads_unclear_not_advantage():
    put_log = {}
    dr = {"missions": {"pricing_scout": {"findings": [
        {"value": 5.0, "note": "Carrefour سعر رف عبوة 1 kg halal organic 2026-08-01 €5"}]}}}
    CA.fill_insights(dr, lambda e: put_log.__setitem__(e["key"], e),
                     hs_code="200811", market_iso3="KWT",
                     product_card={"certifications": ["HALAL", "organic"]})
    d = put_log["differentiation"]
    assert d["value"] == "غير واضح بعد" and d["status"] == L.INFERENCE


# ── ٥) توثيقُ أسعار المنافسين ────────────────────────────────────────────────

def test_price_lines_need_store_date_pack_and_currency():
    dr = {"missions": {"pricing_scout": {"findings": [
        {"value": 4.99, "note": "Albert Heijn عبوة 450 g علامة X 2026-08-14 €4.99"},
        {"value": 8.49, "note": "Jumbo عبوة 350 g 2026-08-14 €8.49"},
        {"value": 9.0, "note": "سعر رف 9 يورو"},                 # بلا متجر/تاريخ/عبوة
        {"value": "لا سعر", "note": "تعذّر"},                    # لا رقم ⇒ لا يُعدّ
    ]}}}
    assert CA.price_evidence(dr["missions"]) == (2, 3)


def test_fewer_than_three_complete_lines_is_weak_three_is_observed():
    weak = _view("turkey_polymers")["ledger"]["entries"]["price_evidence_quality"]
    assert weak["status"] == L.WEAK and weak["value"] == "0 من 1"
    ok = _view("netherlands_honey")["ledger"]["entries"]["price_evidence_quality"]
    assert ok["status"] == L.OBSERVED and ok["value"] == "3 من 3"
    assert "غير رسمية" in L.render_value(weak, "ar")


def test_graded_entry_accepts_only_observed_or_weak():
    with pytest.raises(ValueError):
        L.graded_entry("price_evidence_quality", 1, L.ESTIMATE)


# ── ٦) الاشتراطُ بدرجة التصنيع، والضريبةُ من صفٍّ رسمي ──────────────────────

@pytest.mark.parametrize("spec,hs,expected", [
    ("processing:raw", "080410", True),        # فواكه — خام
    ("processing:processed", "080410", False),
    ("processing:semi,processed", "390210", True),
    ("processing:raw", "990000", None),        # فصلٌ غيرُ مصنَّف ⇒ لا يُخمَّن
])
def test_requirement_rows_bind_to_the_processing_level(spec, hs, expected):
    assert applies_to({"applies_when": spec}, hs, "food") is expected


def test_official_certification_rows_require_an_official_url(tmp_path, monkeypatch):
    p = tmp_path / "certs.csv"
    p.write_text("market,category,kind,scheme,rate_pct,basis,authority,source_url,applies_when,note\n"
                 "KWT,all,certification,Ghost,,official,X,,,\n"
                 "KWT,all,certification,Real,,official,X,https://x.gov,,\n"
                 "KWT,all,certification,Trade,,commercial,X,,,\n", encoding="utf-8")
    monkeypatch.setattr(CA, "CERTS_PATH", str(p))
    CA.certification_rows.cache_clear()
    try:
        rows = CA.certification_rows()
        assert [r["scheme"] for r in rows] == ["Real", "Trade"]
    finally:
        CA.certification_rows.cache_clear()


def test_official_vat_prefers_the_category_row_and_declares_absence():
    assert CA.official_vat("NLD", "080410")[0] == 9.0      # صفُّ الأغذية أخصّ
    assert CA.official_vat("NLD", "390210")[0] == 21.0
    assert CA.official_vat("YEM", "080410") == (None, None)  # صفٌّ تجاريٌّ لا يُعتمد
    assert CA.official_vat("", "080410") == (None, None)
    for r in CA.certification_rows():
        if r["basis"] == "official":
            assert r["source_url"].startswith("http"), r["scheme"]


def test_economics_reads_the_official_vat_when_the_mission_is_silent():
    dr = {"missions": {"pricing_scout": {"findings": [
        {"value": 7.49, "source": "Google Maps", "confidence": 0.8,
         "note": "Albert Heijn سعر رف عبوة 1 كجم بسعر €7.49"}]}}}
    with_row = E.economics_view(dr, category="تمور", market_iso3="NLD",
                                hs_code="080410")
    without = E.economics_view(dr, category="تمور", market_iso3="YEM",
                               hs_code="080410")
    assert not any("ضريب" in g for g in with_row.get("gaps") or [])
    p_with = " ".join((with_row.get("reverse_solve") or {}).get("parameters") or [])
    p_without = " ".join((without.get("reverse_solve") or {}).get("parameters") or [])
    assert "الضريبة (اعتُمد 0%" not in p_with
    assert "الضريبة (اعتُمد 0% — تحتاج تحققاً)" in p_without


def test_a_missing_vat_says_needs_verification_never_a_silent_zero():
    rows = E.margin_waterfall(1.0, freight=0.1, tariff_pct=None, vat_pct=None,
                              distributor_margin_pct=10, retailer_margin_pct=20)
    shelf = rows[-1]
    assert "تحتاج تحققاً" in str(shelf)
    src = open(os.path.join(_ROOT, "silk_economics.py"), encoding="utf-8").read()
    assert "الضريبة (اعتُمد 0% — تحتاج تحققاً)" in src
    assert "الضريبة غير متاحة — صفر معلن" not in src


def test_mandatory_certifications_come_only_from_official_rows():
    tur = _view("turkey_polymers")["ledger"]["entries"]["mandatory_certifications"]
    assert tur["status"] == L.OBSERVED and "TSE" in tur["value"]
    assert all(i["source_url"].startswith("http") for i in tur["items"])
    dza = _view("dza_peanut_butter")["ledger"]["entries"]["mandatory_certifications"]
    assert dza["status"] == L.MISSING and "مطلوبة تجارياً" in dza["note"]


# ── ٧) ادّعاءاتُ القنوات بلا مصدر ─────────────────────────────────────────

def test_an_unsourced_channel_claim_is_returned_to_the_writer():
    led = {"entries": {}}
    text = ("هامش القناة في هذا السوق مرتفع. الموزّعون يقبلون عقوداً سنوية "
            "بسهولة. وفق تقرير الغرفة التجارية هامش القناة 25%. نقدّر هامش "
            "القناة بين 15% و20% بافتراض قناة تجزئة واحدة. "
            # معلمتا المحرّك بصيغة الكاتب المأمور بها — جملةٌ واحدةٌ بفواصل منقوطة
            "بافتراض شحن 12%؛ هامش الموزّع 20%؛ هامش التجزئة 25%.")
    issues = L.draft_issues(text, led)
    claims = [i for i in issues if "ادّعاءٌ عن القناة" in i]
    assert len(claims) == 2
    assert any("هامش القناة في هذا السوق" in c for c in claims)
    assert not any("وفق تقرير" in c or "نقدّر" in c or "هامش الموزّع" in c
                   for c in claims)


def test_repair_deletes_the_unsourced_claim_and_keeps_the_sourced_one():
    led = {"entries": {}}
    text = ("هامش القناة في هذا السوق مرتفع. وفق تقرير الغرفة هامش القناة 25%.\n"
            "بافتراض شحن 12%؛ هامش الموزّع 20%؛ هامش التجزئة 25%.\n"
            "| القناة | هامش القناة |\n")
    out, repairs = L.repair(text, led)
    kinds = [r["kind"] for r in repairs]
    assert kinds == ["channel_claim"]
    assert "هامش القناة في هذا السوق مرتفع" not in out
    assert "وفق تقرير الغرفة" in out
    assert "هامش الموزّع 20%؛ هامش التجزئة 25%" in out   # معلمتا المحرّك لا تُمَسّان
    assert "| القناة | هامش القناة |" in out             # صفوفُ الجداول لا تُمَسّ


# ── العيّنة: المدوّناتُ الثماني عشرة كلُّها ────────────────────────────────────

@pytest.mark.parametrize("key", sorted(CANONICAL_BLOBS))
def test_no_false_alarm_and_no_file_term_on_any_codex(key):
    view = _view(key)
    assert L.check(view, SR.render_markdown(view)) == []
    from docx import Document
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "c.docx")
        SR.render_client_docx(view, p)
        d = Document(p)
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.extend(c.text for c in r.cells)
    txt = "\n".join(parts)
    assert SR._client_forbidden_hits(txt, "ar") == []
    assert "{{" not in txt
    for needle in ("raw_input_trade", "supplier_nature", "[commercial]",
                   "reexporter_likely", "world_top_exporter"):
        assert needle not in txt, (key, needle)


def test_an_industrial_codex_carries_zero_religious_mention():
    view = _view("turkey_polymers")
    md = SR.render_markdown(view)
    assert not re.search(r"حلال|رمضان|مسلم|ديني|halal|ramadan|muslim", md, re.I)


def test_the_commercial_layer_is_stdlib_only():
    import ast
    src = open(os.path.join(_ROOT, "silk_commercial_analysis.py"), encoding="utf-8").read()
    top = {n.names[0].name.split(".")[0] for n in ast.parse(src).body
           if isinstance(n, (ast.Import, ast.ImportFrom)) and not isinstance(n, ast.ImportFrom)}
    top |= {n.module.split(".")[0] for n in ast.parse(src).body
            if isinstance(n, ast.ImportFrom) and n.module}
    assert top <= {"__future__", "csv", "functools", "logging", "os", "re"}, top


def test_missions_wire_both_signals_behind_the_bounded_augment():
    src = open(os.path.join(_ROOT, "silk_missions.py"), encoding="utf-8").read()
    assert '_bounded_augment("commercial_supplier_nature"' in src
    assert '_bounded_augment("commercial_raw_inputs"' in src
    assert src.index('_bounded_augment("competitors_structured"') < \
        src.index('_bounded_augment("commercial_supplier_nature"')


def test_product_card_accepts_origin_claim_only():
    """الوصفُ التجاري (مصطلحات البحث) مؤجَّلٌ إلى د-٣ — لا حقلَ بلا قارئ."""
    src = open(os.path.join(_ROOT, "api.py"), encoding="utf-8").read()
    assert "origin_claim: str | None = None" in src
    assert "trade_description" not in src


def test_cost_advantage_never_speaks_of_rivals_without_their_signal():
    log = {}
    dr = {"missions": {"trade_flow": {"findings": [
        {"value": {"kind": "raw_input_trade", "year": 2024, "net_usd": 5e9,
                   "label": "x", "observed_codes": 1},
         "note": "[commercial] صافي تجارة السعودية في المدخلات الخام"}]}}}
    CA.fill_insights(dr, lambda e: log.__setitem__(e["key"], e), hs_code="390210",
                     market_iso3="TUR")
    ca = log["cost_advantage"]
    assert ca["value"] == "محايد" and "لم تُقيَّم بعد" in ca["note"]
    assert "مورّدي السوق ليسوا" not in ca["note"]
    assert ca["basis"] == ["saudi_raw_input_balance"]


def test_a_declared_gap_is_idempotent_on_resume(monkeypatch):
    import silk_collectors
    monkeypatch.setattr(silk_collectors, "comtrade_budget_left", lambda: 0)
    rep = AgentReport("LLMMissionAgent", [], summary="")
    CA.augment_raw_input_trade(rep, "390210", year=2024)
    CA.augment_raw_input_trade(rep, "390210", year=2024)
    assert len(rep.findings) == 1


def test_stored_values_keep_their_original_date_and_say_so(monkeypatch):
    import silk_collectors
    import silk_store
    monkeypatch.setattr(silk_store, "get_trade_flow",
                        lambda *a, **k: {"value_usd": 1.0, "retrieved_at": "2026-03-01T00:00:00"})
    monkeypatch.setattr(silk_collectors, "comtrade_budget_left", lambda: 100)
    rep = AgentReport("LLMMissionAgent", [], summary="")
    CA.augment_raw_input_trade(rep, "390210", year=2024)
    dp = rep.findings[-1]
    assert dp.retrieved_at == "2026-03-01" and "من المخزن" in dp.note
    assert dp.value["calls_used"] == 0


def test_world_ranking_failure_declares_the_gap_before_spending(no_store, monkeypatch):
    import silk_data_layer
    fake = _fake_comtrade({("390210", None, "X", 0): None})
    monkeypatch.setattr(silk_data_layer, "comtrade_trade", fake)
    rep = AgentReport("LLMMissionAgent",
                      [_competitor_summary(("Belgium", 9.0))], summary="")
    with block_network():
        CA.augment_supplier_nature(rep, "390210", _Market(), year=2024)
    assert rep.findings[-1].value is None and "كبار مصدّري العالم" in rep.findings[-1].note
    assert len(fake.calls) == 1


def test_spent_calls_of_the_first_signal_bound_the_second(no_store, monkeypatch):
    import silk_data_layer
    fake = _fake_comtrade({})
    monkeypatch.setattr(silk_data_layer, "comtrade_trade", fake)
    rep = AgentReport("LLMMissionAgent",
                      [_competitor_summary(("Belgium", 9.0))], summary="")
    CA.augment_supplier_nature(rep, "390210", _Market(), year=2024, spent=CA.max_calls())
    assert fake.calls == [] and "ميزانية" in rep.findings[-1].note


def test_the_official_vat_carries_its_source_next_to_the_number():
    dr = {"missions": {"pricing_scout": {"findings": [
        {"value": 7.49, "source": "Google Maps", "confidence": 0.8,
         "note": "Albert Heijn سعر رف عبوة 1 كجم بسعر €7.49"}]}}}
    eco = E.economics_view(dr, product_card={"cost_per_unit": 2.0, "cost_currency": "EUR"},
                           category="تمور", market_iso3="NLD", hs_code="080410")
    params = " ".join(eco["reverse_solve"]["parameters"])
    assert "مصدر رسمي" in params and "Belastingdienst" in params
    shelf = eco["waterfall"][-1]
    assert "Belastingdienst" in str(shelf)
