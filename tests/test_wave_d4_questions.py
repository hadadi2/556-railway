"""الموجة د-٤ — أسئلةُ المصدّر الحاسمة، والمقدّراتُ بنطاق، والكتابةُ التحليلية.

بلاغُ المالك (2026-09-19، الموجة د، البنود ١٠–١٢): «قسمٌ ثابتٌ بأسئلة المصدّر
الحقيقية وجوابُها من السجلّ أو كيف يعرفه»؛ «الاستنتاجُ مسموحٌ ومطلوب، يُحسَب
في الكود، والتقديرُ لا يُقدَّم متحققاً منه»؛ «الكتابةُ تحليلية: الخلاصة ثمّ
ما يعنيه الرقم ثمّ الدليل».

test-first: كُتب قبل التنفيذ. هرمتي: صفر شبكة، صفر مفتاح، صفر إنفاق.
Run: python3 -m pytest tests/test_wave_d4_questions.py -q
"""
import importlib
import os
import re
import sys
import tempfile

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for _p in (_ROOT, os.path.join(_ROOT, "tools")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import silk_fact_ledger as L                                  # noqa: E402
import silk_render as R                                       # noqa: E402
import silk_reports as SR                                     # noqa: E402
from gen_verdict_baseline import CANONICAL_BLOBS              # noqa: E402


def _view(key):
    mod, fn = CANONICAL_BLOBS[key]
    return R.build_view(getattr(importlib.import_module(mod), fn)())


def _questions(key):
    return _view(key)["ledger"]["critical_questions"]


# ── ١٠) القسمُ الحتميّ: أسئلةُ المصدّر ──────────────────────────────────────

def test_every_question_names_the_ledger_keys_that_answer_it():
    """صفٌّ يسمّي مفتاحاً غيرَ موجودٍ في السجلّ = سؤالٌ لا يُجاب أبداً."""
    keys = {r.key for r in L.KEYS}
    rows = L._question_rows()
    assert rows, "ملفُّ الأسئلة لم يُقرأ"
    for r in rows:
        assert r["question_ar"] and r["how_to_learn_ar"], r["q_key"]
        for k in r["answer_keys"].split(";"):
            assert k.strip() in keys, (r["q_key"], k)


@pytest.mark.parametrize("key", sorted(CANONICAL_BLOBS))
def test_every_codex_gets_a_question_section_with_answers_or_how_to_learn(key):
    rows = _questions(key)
    assert 3 <= len(rows) <= 9, (key, len(rows))
    for r in rows:
        assert r["question"] and r["answer"]
        if not r["answered"]:
            assert r["answer"].startswith("لم نعرفه بعد")
            assert r["how"] and r["how"] != "—", (key, r["q_key"])


def test_the_questions_follow_the_product_category_not_the_product():
    """قاعدةٌ عامّةٌ بالفئة: الصناعيُّ يُسأل عن المواصفة، والغذائيُّ عن الرف."""
    industrial = {r["q_key"] for r in _questions("turkey_polymers")}
    food = {r["q_key"] for r in _questions("netherlands")}
    assert "spec_match" in industrial and "spec_match" not in food
    assert "shelf_gap" in food and "shelf_gap" not in industrial
    assert {"market_size", "cost_edge", "entry_blocker"} <= industrial & food


def test_one_fact_never_answers_two_questions():
    """إعادةُ المعطى جواباً لسؤالين حشوٌ يُقرأ مرّتين (الدرس ٢٦٤)."""
    for key in sorted(CANONICAL_BLOBS):
        answers = [r["answer"] for r in _questions(key) if r["answered"]]
        assert len(answers) == len(set(answers)), key


def test_the_writer_does_not_write_this_section():
    """القسمُ محسوبٌ بالكود — لا يُطلَب من النموذج ولا يُقبَل منه."""
    src = open(os.path.join(_ROOT, "silk_ai_judge.py"), encoding="utf-8").read()
    assert "critical_questions" not in src


def test_the_same_rows_reach_markdown_word_and_the_dashboard():
    view = _view("netherlands")
    rows = view["ledger"]["critical_questions"]
    md = SR.render_markdown(view)
    assert SR.QUESTIONS_TITLE["ar"] in md
    for r in rows[:3]:
        assert r["question"] in md
    from docx import Document
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "r.docx")
        SR.render_docx(view, p)
        txt = "\n".join([x.text for x in Document(p).paragraphs]
                        + [c.text for t in Document(p).tables
                           for row in t.rows for c in row.cells])
    assert SR.QUESTIONS_TITLE["ar"] in txt and rows[0]["question"] in txt
    html = open(os.path.join(_ROOT, "web", "index.html"), encoding="utf-8").read()
    assert "critical_questions" in html and "أسئلتك الحاسمة" in html


def test_the_client_report_carries_the_questions_and_stays_clean():
    from docx import Document
    view = _view("netherlands")
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "c.docx")
        SR.render_client_docx(view, p)
        d = Document(p)
        txt = "\n".join([x.text for x in d.paragraphs]
                        + [c.text for t in d.tables
                           for r in t.rows for c in r.cells])
    assert SR.QUESTIONS_TITLE["ar"] in txt
    assert SR._client_forbidden_hits(txt, "ar") == []
    assert "{{" not in txt


# ── ١١) المقدّراتُ بنطاقٍ وافتراض ───────────────────────────────────────────

def test_an_estimate_carries_its_range_and_assumption_or_is_not_written():
    view = _view("netherlands")
    e = view["ledger"]["entries"]["max_exw_estimate"]
    assert e["status"] == L.ESTIMATE
    lo, hi = e["range"]
    assert lo < e["value"] < hi and e["assumption"]
    assert "سيناريو المتوسط" in e["assumption"]
    txt = L.render_value(e, "ar")
    assert "بين" in txt and "إذا افترضنا" in txt


def test_a_missing_input_yields_no_estimate_at_all():
    """مدخلٌ ناقصٌ = لا تقدير — لا رقمٌ مخترَعٌ ليملأ الفراغ."""
    view = _view("turkey_polymers")
    e = view["ledger"]["entries"]["border_to_shelf_multiple"]
    assert e["status"] == L.MISSING and e["value"] is None
    assert not any(r["key"] == "border_to_shelf_multiple"
                   for r in view["ledger"]["gaps"])


def test_estimates_never_move_the_verdict():
    """قيدُ المالك بنيوياً: التقديرُ يدخل السجلَّ ولا يدخل أعمدةَ القرار."""
    import silk_deep_pillars
    blob = getattr(importlib.import_module(CANONICAL_BLOBS["netherlands"][0]),
                   CANONICAL_BLOBS["netherlands"][1])()
    before = silk_deep_pillars.decide_for_deep(blob.get("deep_research") or {},
                                               product_card=None, regulatory=None)
    view = R.build_view(blob)
    L.fill_estimates(view)
    after = silk_deep_pillars.decide_for_deep(blob.get("deep_research") or {},
                                              product_card=None, regulatory=None)
    assert (before or {}).get("verdict") == (after or {}).get("verdict")
    src = open(os.path.join(_ROOT, "silk_deep_pillars.py"), encoding="utf-8").read()
    assert "critical_questions" not in src and "max_exw_estimate" not in src


def test_a_trial_shipment_estimate_states_its_method_and_range():
    e = _view("netherlands")["ledger"]["entries"]["trial_shipment_units"]
    assert e["status"] == L.ESTIMATE
    lo, hi = e["range"]
    assert lo < hi and "حاوية" in e["assumption"]
    assert e["how_to_close"]


# ── ١٢) الكتابةُ التحليلية ──────────────────────────────────────────────────

def _pre_writer_ledger(key: str):
    """السجلُّ **بالشكل الذي يصل الكاتبَ فعلاً** — `silk_research_pipeline:664`.

    اختبارُ الكتلة على سجلّ ما-بعد-`build_view` كان يمرّ وهو لا يصف المسارَ
    الحيّ: ذاك يحمل قسمَ الاقتصاد وهذا لا يحمله (§58).
    """
    import importlib
    mod, fn = CANONICAL_BLOBS[key]
    blob = getattr(importlib.import_module(mod), fn)()
    dr = blob.get("deep_research") or blob
    return L.build_ledger(
        {"deep_research": {"missions": (dr.get("missions") or {})},
         "markets": [{"decision": (dr.get("decision") or {})}],
         "regulatory": {}}, "ar")


def test_the_insights_block_gives_the_writer_tokens_not_values():
    block = L.insights_block(_pre_writer_ledger("netherlands_honey"))
    assert "{{cost_advantage}}" in block
    assert "اكتب الرمز لا القيمة" in block
    # القيمةُ نفسُها لا تُسلَّم للنموذج في هذه الكتلة
    assert "ميزة (يشير" not in block


def test_an_absent_insight_is_not_listed_to_the_writer():
    block = L.insights_block(_pre_writer_ledger("netherlands"))
    assert "{{cost_advantage}}" not in block          # لا إشارةَ محسوبة


def test_the_estimates_reach_the_reader_by_code_not_by_the_writer():
    """حدٌّ بنيويٌّ معلَن: مدخلاتُ التقديرات في قسم الاقتصاد، وهو يُبنى بعد
    الكاتب — فلا تصل كتلةَ `[INSIGHTS]`، وتصل القارئَ من قسم الأسئلة."""
    pre = _pre_writer_ledger("netherlands")
    assert "{{max_exw_estimate}}" not in L.insights_block(pre)
    view = _view("netherlands")
    assert "{{max_exw_estimate}}" in L.insights_block(view["ledger"])
    answers = " ".join(r["answer"] for r in L.critical_questions(view, "ar"))
    assert "بين" in answers and "إذا افترضنا" in answers


def test_the_insights_block_follows_the_report_language():
    block = L.insights_block(_pre_writer_ledger("netherlands_honey"), "en")
    assert "write the token, not the value" in block
    assert "اكتب الرمز" not in block and "استنتاجٌ" not in block


def test_the_writing_rule_reaches_the_prompt_with_the_block():
    src = open(os.path.join(_ROOT, "silk_ai_judge.py"), encoding="utf-8").read()
    assert "[INSIGHTS]" in src and "INSIGHT_WRITING_RULE" in src
    assert "الخلاصة ← ما يعنيه الرقم للمصدّر ← الدليل" in L.INSIGHT_WRITING_RULE


# ── منتجو عبارات الغياب ─────────────────────────────────────────────────────

def test_a_known_hs_code_never_claims_the_product_could_not_be_classified():
    """حدٌّ كاذبٌ كان في ١٦/١٦ تقرير: «تعذّر التصنيف» فوق تقريرٍ يعرض الرمز."""
    for key in sorted(CANONICAL_BLOBS):
        view = _view(key)
        assert view.get("hs_code")
        assert not any("تعذّر التصنيف" in str(x) for x in view.get("limits") or []), key
        assert "تعذّر التصنيف" not in SR.render_markdown(view), key


def test_the_reader_phrase_is_one_across_surfaces():
    """«— الناقص:» مفردةٌ داخلية صارت «— يلزم:» في منشئها لا في الحارس."""
    for mod in ("silk_reports.py", "silk_economics.py"):
        src = open(os.path.join(_ROOT, mod), encoding="utf-8").read()
        assert "لا نعرفه بعد — الناقص:" not in src
        assert "غير محسوب — الناقص:" not in src


@pytest.mark.parametrize("key", sorted(CANONICAL_BLOBS))
def test_no_false_alarm_with_the_new_section(key):
    view = _view(key)
    assert L.check(view, SR.render_markdown(view)) == []


def test_no_report_claims_zero_gaps_while_a_question_is_unanswered():
    """قسمٌ ينفي ما يعرضه قسمٌ آخر — بعينه بلاغُ المالك الذي أنشأ الموجة (د).

    إسقاطُ حدِّ «تعذّر التصنيف» الكاذب أفرغ جدولَ النواقص على ثلاث مدوّنات
    (هولندا، ألمانيا، اليابان — قِيس بـ`tools/golden_set.py compare`)، فصار
    التقريرُ يطبع «لا فجوات مرصودة» وفوقَه مباشرةً أسئلةٌ جوابُها «لم نعرفه
    بعد». الحارسُ عامٌّ على المدوّنات كلِّها لا على الثلاث.
    """
    for key in sorted(CANONICAL_BLOBS):
        view = _view(key)
        md = SR.render_markdown(view)
        if "لا فجوات مرصودة" not in md and "لا حدود مسجّلة" not in md:
            continue
        rows = L.critical_questions(view, "ar")
        unanswered = [r["q_key"] for r in rows if not r.get("answered")]
        assert not unanswered, (
            f"{key}: التقرير ينفي الفجوات وفيه أسئلةٌ بلا جواب {unanswered}")


def test_the_pointer_line_replaces_the_denial_when_questions_are_open():
    view = _view("netherlands_dates" if "netherlands_dates" in CANONICAL_BLOBS
                 else sorted(CANONICAL_BLOBS)[0])
    line = SR._no_gaps_line(view, "ar")
    rows = L.critical_questions(view, "ar")
    if any(not r.get("answered") for r in rows):
        assert "أسئلتك الحاسمة" in line and "لا فجوات" not in line
    else:
        assert line == SR._NO_GAPS["ar"]


# ── حصادُ المراجعة الذاتية §58 ───────────────────────────────────────────────

def test_the_border_to_shelf_multiple_never_divides_a_pack_by_a_kilo():
    """`anchor["value_usd"]` سعرُ **عبوة** (`per_unit / fx`) والمقسومُ عليه
    سعرُ **كيلوغرام** — مضاعفٌ من مقامَين مختلفَين رقمٌ مخترَع."""
    src = open(os.path.join(_ROOT, "silk_fact_ledger.py"), encoding="utf-8").read()
    body = src.split("def fill_estimates(")[1].split("\ndef ")[0]
    head = body.split("# (ب)")[0]
    assert 'anchor.get("value_usd")' not in head
    assert 'anchor.get("raw_value")' not in head
    assert 'anchor.get("per_kg")' in head


def test_a_declared_pricing_contradiction_suppresses_the_ceiling_estimate():
    """تركيا: نقصٌ ٩٧٫٩٪ — المحرّك نفسُه يقول إنّ السقف لا يصلح للتفاوض،
    فلا يُقدَّم جواباً لسؤال «بكم أبيع للمصنع؟»."""
    hit = False
    for key in sorted(CANONICAL_BLOBS):
        view = _view(key)
        eco = (view.get("deep_research") or {}).get("economics") or {}
        if not eco.get("pricing_contradiction"):
            continue
        hit = True
        ents = (view.get("ledger") or {}).get("entries") or {}
        e = ents.get("max_exw_estimate")
        assert not e or e.get("value") is None, key
        row = [r for r in L.critical_questions(view, "ar")
               if r["q_key"] == "max_price"]
        assert row and not row[0]["answered"], key
    assert hit, "لا مدوّنةَ تحمل تناقضَ تسعيرٍ — الحارسُ بلا عيّنة"


def test_every_estimate_carries_its_unit_in_the_rendered_range():
    for key in sorted(CANONICAL_BLOBS):
        view = _view(key)
        for e in ((view.get("ledger") or {}).get("entries") or {}).values():
            if e.get("status") != L.ESTIMATE or e.get("value") is None:
                continue
            unit = str(e.get("unit") or "").strip()
            if not unit or unit in ("%", "USD", "USD/kg"):
                continue
            assert unit in L.render_value(e, "ar"), (key, e["key"])


def test_the_trial_shipment_assumption_matches_the_range_it_prints():
    """ذيلُ «نصفُ حمولةٍ إلى حمولة» لا يُلحَق بنطاقٍ جاء من بند القرار."""
    for key in sorted(CANONICAL_BLOBS):
        e = ((_view(key).get("ledger") or {}).get("entries")
             or {}).get("trial_shipment_units")
        if not e or e.get("value") is None:
            continue
        lo, hi = e["range"]
        widened = abs(lo - round(e["value"] * 0.5, 2)) < 0.01
        assert ("نصف حمولةٍ" in e["assumption"]) == widened, key


def test_a_shared_fact_never_becomes_a_false_gap_in_a_later_question():
    """تركيا كانت تطبع الشهادةَ جواباً للحاجز ثمّ «لم نعرفه بعد» للمواصفة."""
    for key in sorted(CANONICAL_BLOBS):
        view = _view(key)
        rows = L.critical_questions(view, "ar")
        answers = [r["answer"] for r in rows if r.get("answered")]
        # لا نصَّ يتكرّر حرفياً، ولا سؤالَ يُنكِر معلوماً طُبع أعلاه
        body = [a for a in answers if a != "مذكورٌ في جوابٍ أعلاه"]
        assert len(body) == len(set(body)), key
        for r in rows:
            if r.get("answered"):
                continue
            keys = [k.strip() for k in
                    next(q["answer_keys"] for q in L._question_rows()
                         if q["q_key"] == r["q_key"]).split(";") if k.strip()]
            ents = (view.get("ledger") or {}).get("entries") or {}
            assert not any(
                ents.get(k) and ents[k].get("value") is not None
                and ents[k].get("status") != L.MISSING for k in keys), (key, r["q_key"])


def test_the_client_word_questions_do_not_hang_off_the_gaps_branch():
    """مدوّنةٌ بلا فجوةِ قرارٍ حرجة كانت تُسقِط القسمَ كلَّه."""
    from docx import Document
    import tempfile
    for key in ("nadec_yemen_dairy", "netherlands"):
        if key not in CANONICAL_BLOBS:
            continue
        view = _view(key)
        rows = L.critical_questions(view, "ar")
        if not rows:
            continue
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "c.docx")
            SR.render_client_docx(view, p)
            d = Document(p)
            txt = "\n".join([x.text for x in d.paragraphs]
                            + [c.text for t in d.tables
                               for r in t.rows for c in r.cells])
        assert SR.QUESTIONS_TITLE["ar"] in txt, key


def test_a_failing_estimator_never_empties_a_real_gaps_table(monkeypatch):
    import silk_render as SRnd
    view = _view("yemen")
    before = len(view["ledger"]["gaps_table"])
    assert before > 0
    monkeypatch.setattr(L, "fill_estimates",
                        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("x")))
    mod, fn = CANONICAL_BLOBS["yemen"]
    import importlib
    blob = getattr(importlib.import_module(mod), fn)()
    v2 = SRnd.build_view(blob)
    assert len(v2["ledger"]["gaps_table"]) == before


def test_every_question_row_carries_both_language_mirrors():
    for r in L._question_rows():
        assert r["question_ar"] and r["question_en"], r["q_key"]
        assert r["how_to_learn_ar"] and r.get("how_to_learn_en"), r["q_key"]


def test_the_english_client_report_carries_the_questions_and_their_path():
    """كان القسمُ يسقط كلّياً (صفوفٌ عربيةٌ مُصيَّرةٌ سلفاً + `_lang_safe`)،
    ثمّ صار يطبع «not yet known» بلا سبيلِ إغلاق."""
    import tempfile
    from docx import Document
    from tools.gen_client_report_sample import build_sample_view
    view = build_sample_view("en")
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "c.docx")
        SR.render_client_docx(view, p)
        txt = "\n".join(x.text for x in Document(p).paragraphs)
    assert SR.QUESTIONS_TITLE["en"] in txt
    assert "How big is this market and is it growing?" in txt
    assert "How to find out:" in txt
    body = txt[txt.index(SR.QUESTIONS_TITLE["en"]):]
    head = body.split(SR.QUESTIONS_TITLE["en"])[1][:1500]
    assert "لم نعرفه" not in head and "اطلب" not in head
