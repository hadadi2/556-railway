"""تقرير سِلك ٧ — المسؤولُ وخطةُ التسعين يوماً وسجلُّ القناة وخطوةُ كلّ جهة (§4.1، §4.2، §4.4، §6، الدرس ٢٨٢).

Every open condition carries a suggested owner by its kind; the 90-day plan
frame is built from the same list (action, owner, window, output, continue or
stop — data completion leads to re-evaluation, never an automatic "enter");
the entry channel record carries the target consumer, the dominant qualified
buyer, the product nature and the reason; every listed contact carries its
channel, reason and next step, and a long address is printed in full below;
a survey share without sample size or population is flagged.

كلُّ الحالات مصطنعةٌ لاختبارٍ هرمتيّ.
"""
from docx import Document

import silk_quality_gate as Q
import silk_render as R
import silk_reports as SR


def _ed():
    return {"schema": "silk.decision/v1", "verdict": "PRELIMINARY GO",
            "conditions": ["a", "b", "c"],
            "condition_items": [
                {"id": "C1", "kind": "eligibility_gate", "pillar": "regulatory",
                 "status": "open"},
                {"id": "C2", "kind": "pillar_missing", "pillar": "profit",
                 "missing": ["margin_at_border_pct"], "status": "open"},
                {"id": "C3", "kind": "pillar_weak", "pillar": "market",
                 "pct": 31, "status": "open"}],
            "pillars": {"market": {"value": 0.31, "missing": []},
                        "profit": {"value": None,
                                   "missing": ["margin_at_border_pct"]},
                        "regulatory": {"value": 0.3, "missing": [],
                                       "eligibility_gate": True}}}


def test_each_condition_has_a_suggested_owner_by_kind():
    rows = R.condition_texts(_ed(), "ar")
    owners = [r["owner"] for r in rows]
    assert "الجهة الرقابية" in owners[0] and "المصنع" in owners[1]
    assert all("مقترح" in o for o in owners)


def test_the_plan_frame_comes_from_the_same_conditions():
    plan = R.plan_90(R.condition_texts(_ed(), "ar"), "ar")
    assert [p["condition"] for p in plan] == ["الشرط 1", "الشرط 2", "الشرط 3"]
    assert "يسبق كل ما بعده" in plan[0]["window"]
    assert "لا يتحوّل الحكم" in plan[1]["gate"]
    assert all(p["action"] and p["owner"] and p["output"] for p in plan)
    en = R.plan_90(R.condition_texts(_ed(), "en"), "en")
    assert en[0]["condition"] == "Condition 1"


def test_the_basis_carries_the_plan_and_the_client_report_prints_it():
    basis = R.decision_basis(_ed())
    assert len(basis["plan"]) == 3 and len(basis["plan_cols"]) == 6
    doc = Document()
    SR._client_decision_basis(doc, {"decision": {"basis": basis},
                                    "markets": [{"entry_decision": _ed()}]},
                              "ar")
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "إطار خطة التسعين يوماً" in heads
    assert doc.tables[-1].rows[0].cells[5].text == "الاستمرار أو التوقف"


def test_the_channel_record_is_complete_from_its_evidence():
    ch = R.entry_channel({
        "entry_door": [{"value": "موزّعو القهوة المتخصصون", "confidence": 0.7,
                        "note": "هوامش مستقرة ووصول إلى المقاهي"}],
        "demand": [{"value": "أسر حضرية تستهلك القهوة في المنزل",
                    "confidence": 0.6}]}, "090121")
    assert ch["target_consumer"].startswith("أسر حضرية")
    assert ch["reasons"].startswith("هوامش مستقرة")
    assert "مصنّع" in ch["product_nature"]
    assert R._dominant_buyer([
        {"category": "تاجر جملة قهوة", "evidence_status": "specialist"},
        {"category": "تاجر جملة قهوة", "evidence_status": "specialist"},
        {"category": "سوق مركزي", "evidence_status": "named_unverified"}]) \
        == "تاجر جملة قهوة (2)"


def test_the_client_report_prints_the_channel_record():
    doc = Document()
    SR._client_entry_channel(doc, {"entry_channel": {
        "primary": "موزّعو القهوة المتخصصون", "alternative": "الضيافة",
        "target_consumer": "أسر حضرية", "commercial_buyer": "",
        "product_nature": "منتج غذائي/زراعي — مصنّع", "reasons": "هوامش"}},
        "ar")
    cells = [r.cells[0].text for r in doc.tables[-1].rows[1:]]
    assert "المستهلك المستهدف (من تحليل الطلب)" in cells
    assert "المشتري التجاري الغالب بين الجهات المؤهّلة" not in cells
    assert doc.tables[-1].rows[-1].cells[1].text.startswith("مرشّحة")
    en = Document()
    SR._client_entry_channel(en, {"entry_channel": {"primary": "x"}}, "en")
    assert not en.tables and en.paragraphs


def test_each_contact_has_channel_reason_and_next_step():
    long_addr = "Lot 12, Jalan Perusahaan 3, Kawasan Perindustrian Bukit " \
                "Beruntung, 48300 Rawang, Selangor Darul Ehsan, Malaysia"
    doc = Document()
    SR._docx_leads(doc, {"importer_leads": {"leads": [
        {"name": "Kopi Hub", "category": "coffee wholesaler",
         "evidence_status": "specialist", "phone": "+60 3 1",
         "address": long_addr}]},
        "market": {"iso3": "MYS"}, "product": "قهوة", "hs_code": "090121"},
        lang="ar")
    main, actions, addrs = doc.tables[-3], doc.tables[-2], doc.tables[-1]
    assert "(كاملاً أدناه)" in main.rows[1].cells[1].text
    assert actions.rows[1].cells[1].text == "توزيع وجملة"
    assert "قائمة أسعار الشراء" in actions.rows[1].cells[3].text
    assert addrs.rows[1].cells[1].text == long_addr


def test_a_survey_share_without_scope_is_flagged():
    bad = "أظهر استبيان أن 62% من الماليزيين يشربون القهوة في المنزل."
    ok = "أظهر استبيان شمل 400 مستجيب في كوالالمبور أن 62% يشربونها في المنزل."
    assert Q._check_survey_share_generalized(bad)
    assert not Q._check_survey_share_generalized(ok)
    assert not Q._check_survey_share_generalized("نمت الواردات 19% في 2024.")


# ── مراجعة §58 ──────────────────────────────────────────────────────────────

def test_a_tag_or_status_note_is_not_a_reason_and_a_statistic_not_a_consumer():
    ch = R.entry_channel({
        "entry_door": [{"value": "موزّعون", "confidence": 0.7,
                        "note": "[entry_door] مرشّح يحتاج تحققاً"}],
        "demand": [{"value": "نمت الواردات 19% إلى 120 مليون دولار",
                    "confidence": 0.8}]})
    assert ch["reasons"] == "" and ch["target_consumer"] == ""
    ch = R.entry_channel({"entry_door": [
        {"value": "موزّعون", "confidence": 0.7,
         "note": "[entry_door] هوامش مستقرة ووصول إلى المقاهي"}]})
    assert ch["reasons"] == "هوامش مستقرة ووصول إلى المقاهي"


def test_a_single_or_untranslated_buyer_is_not_dominant():
    assert R._dominant_buyer([{"category": "تاجر جملة قهوة",
                               "evidence_status": "specialist"}]) == ""
    assert R._dominant_buyer([{"category": "Cafe"}, {"category": "Cafe"}]) == ""


def test_the_plan_lists_only_the_conditions_shown_with_their_text():
    import silk_render as RR
    ed = _ed()
    ed["condition_items"] += [
        {"id": f"C{n}", "kind": "pillar_weak", "pillar": "risk", "pct": 30,
         "status": "open"} for n in range(4, 9)]
    basis = RR.decision_basis(ed)
    shown = [c for c in basis["conditions"] if c.startswith("الشرط ")]
    assert len(basis["plan"]) == len(shown)


def test_the_view_carries_channel_and_next_step_per_contact():
    ld = {"category": "food manufacturer", "evidence_status": "general_trader"}
    assert R.lead_channel(ld, "ar").startswith("أخرى (")
    assert R.lead_channel({"category": "coffee wholesaler"}, "ar") == "توزيع وجملة"
    assert R.lead_channel({}, "ar") == "غير محدد"
    assert "تحقق من نشاطها" in R.lead_next_step({}, "ar")


def test_common_words_do_not_trip_the_survey_check():
    for ok in ("تستحوذ علامات معينة على 40% من السوق.",
               "ترتفع الرسوم على فئات معيّنة إلى 10%.",
               "Supplier share 62% per UN Comtrade sample year.",
               "مسح ميداني شمل 300 متجر أظهر أن 12% تبيع البن المحمص."):
        assert not Q._check_survey_share_generalized(ok), ok
    assert Q._check_survey_share_generalized(
        "وفق عينة أكاديمية فإن 70% يفضّلون القهوة المحمصة.")
