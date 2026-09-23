"""تقرير سِلك ٧ — اشتراطاتُ ماليزيا ودرجةُ التصنيع والسنةُ الجزئية (§3.2، §4.3، الدرس ٢٨٠).

Malaysia — the report-7 market — had no reference rows, so a coffee study
showed a single gap row; roasted coffee (0901.21) was read "raw" by chapter,
so raw-bean rules applied to it; a zero tariff sat next to an unmeasured
sales tax silently; and the current (incomplete) year entered growth.

كلُّ الحالات مصطنعةٌ لاختبارٍ هرمتيّ.
"""
import csv
import datetime
import pathlib

from docx import Document

import silk_ai_judge as AJ
import silk_deep_pillars as P
import silk_reports as SR
import silk_requirements_agent as RA

_ROOT = pathlib.Path(__file__).resolve().parents[1]


def _entry(market, hs):
    return [r for r in RA.requirement_rows(market, hs) if r["direction"] == "entry"]


def test_roasted_coffee_is_processed_not_raw():
    assert AJ.processing_level("090121") == "processed"
    assert AJ.processing_level("090122") == "processed"
    assert AJ.processing_level("090111") == "raw"
    # بلا صفِّ بند: تصنيفُ الفصل كما كان.
    assert AJ.processing_level("020130") == "raw"


def test_malaysia_coffee_has_real_rows_with_type_and_check_date():
    rows = _entry("MYS", "090121")
    assert rows and not any(r.get("gap") for r in rows)
    assert all(r["status"] and r["verified_at"] for r in rows)
    items = " ".join(r["item"] for r in rows)
    assert "FoSIM" in items and "الملايوية" in items
    # البنُّ المحمّص خارج تصريح MAQIS للبنّ الخام.
    assert "MAQIS" not in items
    assert "MAQIS" in " ".join(r["item"] for r in _entry("MYS", "090111"))


def test_halal_is_a_buyer_requirement_for_coffee_and_a_gate_for_meat():
    coffee = [r for r in _entry("MYS", "090121") if "حلال" in r["item"]]
    assert len(coffee) == 1 and coffee[0]["status"] == "buyer_requirement"
    assert "التجزئة" in coffee[0]["item"] and "الإلكتروني" in coffee[0]["item"]
    meat = _entry("MYS", "020130")
    assert "حلال" in meat[0]["item"] and meat[0]["status"] == "legal_mandatory"
    assert all(r["conditional_on_gate"] for r in meat[1:])
    assert not any("لغير اللحوم" in r["item"] for r in meat)


def test_a_filtered_out_gate_does_not_mark_the_rest_conditional():
    dairy = _entry("MYS", "040120")
    assert dairy and not any(r["conditional_on_gate"] for r in dairy)


def test_sales_tax_is_a_banded_requirement_not_a_guessed_rate():
    rows = _entry("MYS", "090121")
    sst = [r for r in rows if "SST" in r["item"]]
    assert sst and "0% أو 5% أو 10%" in sst[0]["item"]
    import silk_commercial_analysis as CA
    assert CA.official_vat("MYS", "090121") == (None, None)


def test_every_reference_row_has_the_verified_at_column():
    rows = list(csv.DictReader(open(_ROOT / "data/requirements_l1.csv",
                                    encoding="utf-8")))
    assert all("verified_at" in r for r in rows)
    mys = [r for r in rows if r["market"] == "MYS"]
    assert mys and all(r["verified_at"] for r in mys)


def test_the_check_date_is_shown_in_the_client_table():
    doc = Document()
    SR._client_requirements_table(doc, {"requirements": RA.requirement_rows(
        "MYS", "090121")}, "ar")
    auth = [r.cells[2].text for r in doc.tables[-1].rows[1:]]
    assert any("آخر تحقق: 2026-09-23" in a for a in auth)


def test_a_zero_tariff_names_the_unmeasured_sales_tax():
    import silk_economics as E
    dr = {"market": {"iso3": "MYS"}, "hs_code": "090121",
          "missions": {"tariffs_agreements": {"findings": [
              {"value": 0.0, "source": "WITS", "confidence": 0.9,
               "note": "التعرفة المطبقة 0% HS090121"}]}}}
    gaps = " ".join(E.economics_view(dr).get("gaps") or [])
    assert "لا يعني انعدام بقية الرسوم" in gaps


def test_the_current_year_is_partial_and_stays_out_of_growth():
    y = datetime.date.today().year
    missions = {"trade_flow": {"findings": [
        {"value": 50.0e6, "source": "UN Comtrade", "confidence": 0.9,
         "data_year": y - 2, "note": f"إجمالي استيراد {y - 2}"},
        {"value": 60.0e6, "source": "UN Comtrade", "confidence": 0.9,
         "data_year": y - 1, "note": f"إجمالي استيراد {y - 1}"},
        {"value": 10.0e6, "source": "UN Comtrade", "confidence": 0.9,
         "data_year": y, "note": f"إجمالي استيراد {y}"}]}}
    s = P.import_series(missions)
    assert [p["partial"] for p in s["series"]] == [False, False, True]
    assert s["growth_pct"] is not None and s["growth_pct"] > 0
    doc = Document()
    SR._client_imports_section(doc, {"imports": {
        "value_line": "واردات", "head": "الواردات", "series": s["series"]}},
        "ar")
    assert "جزئي" in doc.tables[-1].rows[3].cells[1].text


def test_a_year_dropped_from_claims_is_recovered_from_raw_evidence():
    """سنةٌ حُفظت في لقطة الأداة ثمّ سقطت من ادّعاء البعثة تُستعاد منها."""
    missions = {"trade_flow": {"findings": [
        {"value": 89.4e6, "source": "UN Comtrade", "confidence": 0.9,
         "data_year": 2024, "note": "إجمالي استيراد 2024",
         "raw_evidence": [{"value": 74.87e6, "source": "UN Comtrade",
                           "confidence": 0.9, "data_year": 2023,
                           "note": "إجمالي استيراد 2023"}]}]}}
    years = [p["year"] for p in P.import_series(missions)["series"]]
    assert years == [2023, 2024]


# ── مراجعة §58 ──────────────────────────────────────────────────────────────

def test_the_agent_sees_the_same_filtered_items_as_the_table():
    rep = RA.RequirementsAgent().run({"market_iso3": "MYS", "hs_code": "090121"})
    items = [f.value["item"] for f in rep.findings
             if f.value and f.value.get("direction") == "entry"]
    assert items and not any("MAQIS" in i for i in items)
    meat = RA.RequirementsAgent().run({"market_iso3": "MYS", "hs_code": "020130"})
    vals = [f.value for f in meat.findings if f.value]
    assert vals[0]["eligibility_gate"] and "حلال" in vals[0]["item"]
    assert not any("لغير اللحوم" in v["item"] for v in vals)
    dairy = RA.RequirementsAgent().run({"market_iso3": "MYS", "hs_code": "040120"})
    assert "الأهلية أولاً" not in dairy.summary


def test_a_non_eu_gate_is_named_by_its_item_not_as_eu_listing():
    import silk_decision as D
    import silk_render as R
    assert "2017/625" in D.gate_label_ar("")
    label = D.gate_label_ar("شهادة حلال من جهة أجنبية تعترف بها JAKIM")
    assert "JAKIM" in label and "2017/625" not in label
    rows = R.condition_texts({"condition_items": [
        {"id": "C1", "kind": "eligibility_gate", "pillar": "regulatory",
         "item": "شهادة حلال من جهة أجنبية تعترف بها JAKIM"}]}, "ar")
    assert "JAKIM" in rows[0]["text"] and "2017/625" not in rows[0]["text"]


def test_green_coffee_permit_is_not_a_hard_blocker():
    st = RA.regulatory_state("MYS", "090111")
    assert not st["open_hard"]


def _partial_missions():
    y = datetime.date.today().year
    return {"trade_flow": {"findings": [
        {"value": 60.0e6, "source": "UN Comtrade", "confidence": 0.9,
         "data_year": y - 1, "note": f"إجمالي استيراد {y - 1}"},
        {"value": 10.0e6, "source": "UN Comtrade", "confidence": 0.9,
         "data_year": y, "note": f"إجمالي استيراد {y}"}]}}


def test_market_size_and_ledger_skip_the_partial_year():
    y = datetime.date.today().year
    m = _partial_missions()
    val, _f = P._numeric_with_source(P._metric_findings(m, "trade_flow"),
                                     "tam_usd")
    assert val == 60.0e6
    import silk_fact_ledger as FL
    e = FL.build_ledger({"deep_research": {"missions": m}})["entries"]
    assert e["imports_latest_year"]["value"] == y - 1


def test_the_chart_leaves_out_the_partial_year_and_says_so():
    import silk_render as R
    y = datetime.date.today().year
    imports = {"series": [
        {"year": y - 2, "value": 5e7}, {"year": y - 1, "value": 6e7},
        {"year": y, "value": 1e7, "partial": True}]}
    ch = R._chart_imports_trend(imports, "ar")
    assert [s["year"] for s in ch["series"]] == [y - 2, y - 1]
    assert str(y) in ch["note"]


def test_an_unmeasured_tax_is_named_whatever_the_tariff():
    import silk_economics as E
    dr = {"market": {"iso3": "MYS"}, "hs_code": "090121",
          "missions": {"tariffs_agreements": {"findings": [
              {"value": 5.0, "source": "WITS", "confidence": 0.9,
               "note": "التعرفة المطبقة 5% HS090121"}]}}}
    gaps = " ".join(E.economics_view(dr).get("gaps") or [])
    assert "غير مرصودة" in gaps and "لا يعني انعدام" not in gaps
