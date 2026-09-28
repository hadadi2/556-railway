"""تصدير نمط «دراسة السوق» — Markdown/Word/PDF من محرك القوالب (P2-1/P2-5).

`?style=study` على مسارات التصدير: الحالة من `silk_study_case.build_case` (نتيجة
`/research` المخزَّنة)، والمعرفة من `data/product_knowledge/<hs>_<iso2>.yaml` إن كانت
approved، والنص من `silk_study_render`. Word عبر python-docx بهوية سِلك RTL نفسها
(`_apply_rtl`/`_add_table`)، وPDF عبر `docx_to_pdf` الموجود. لا نداء نموذج هنا:
فراغات (ب)/(ج) تُملأ مرّة واحدة داخل خط `/research` (`fill_slots`، ضمن حارس
الميزانية) وتُخزَّن في `deep_research.study_slots`؛ التصدير يقرأ المخزَّن فقط.
"""
from __future__ import annotations

import re

_BOLD = re.compile(r"\*\*(.+?)\*\*")


def study_case(found: dict) -> dict:
    from silk_study_case import build_case
    return build_case(found)


def _stored_fill(found: dict):
    """llm_fill يقرأ الفراغات المخزَّنة فقط — غيابها فجوة معلنة، لا نداء."""
    slots = ((found.get("deep_research") or {}).get("study_slots") or {})
    return (lambda sid, _brief: slots.get(sid)) if slots else None


def fill_slots(found: dict, call, guard=None, max_calls: int | None = None) -> dict:
    """يملأ فراغات (ب)/(ج) الناقصة بنداء `call(sid, prompt)` ويعيد {sid: نص} لما
    اجتاز فحص المحرك (P2-3/P2-4). يُستدعى من خط `/research` وحده.
    `guard()` يُسأل **قبل كل نداء** (سقف التكلفة/المهلة/الإلغاء)؛ `max_calls` سقفٌ
    صلب لعدد النداءات (SILK_STUDY_SLOTS_MAX_CALLS، افتراضياً عدد الموجزات × محاولتين).
    ما نجح قبل أي عطل يُعاد ولا يضيع (الكلفة دُفعت)."""
    import os as _os
    from silk_study_render import Renderer, load_knowledge, load_templates
    if max_calls is None:
        default = 2 * len(load_templates().get("llm_briefs") or {})
        try:
            max_calls = int(_os.environ.get("SILK_STUDY_SLOTS_MAX_CALLS", str(default)))
        except ValueError:
            max_calls = default
    n = {"calls": 0}

    def gated(sid, prompt):
        if n["calls"] >= max_calls or (guard is not None and not guard()):
            return None
        n["calls"] += 1
        return call(sid, prompt)
    case = found.get("study_case") or study_case(found)
    kn = load_knowledge(case["product"]["hs"], case["market"].get("iso2") or "")
    r = Renderer(case, kn, llm_fill=gated)
    try:
        r.render()
    except Exception as e:  # noqa: BLE001 — ما دُفع ثمنه لا يُرمى
        import logging
        logging.getLogger(__name__).warning("fill_slots render failed after %d calls: %s",
                                            n["calls"], e)
    return {sid: txt for sid, txt in r._llm_cache.items() if txt}


def study_markdown(found: dict, llm_fill=None) -> tuple[str, dict]:
    """(نص Markdown، تقرير التعبئة {gaps, llm_slots, missing})."""
    from silk_study_render import Renderer, load_knowledge
    case = found.get("study_case") or study_case(found)
    kn = load_knowledge(case["product"]["hs"], case["market"].get("iso2") or "")
    stored = llm_fill is None
    llm_fill = llm_fill or _stored_fill(found)
    # الفراغات المخزَّنة اجتازت الفحوص عند ملئها؛ لا يُعاد فحصها بقواعد لاحقة فتُسقط.
    r = Renderer(case, kn, llm_fill=llm_fill, recheck=not stored)
    md = r.render()        # بلا وسم مراجعة (review_marks=False افتراضياً، الدرس 285)
    from silk_quality_gate import study_style_violations   # استشاري: لا يُسقط التصدير
    from silk_study_claims import build_claims
    return md, {"gaps": list(r.gaps), "llm_slots": list(r.llm_slots), "missing": list(r.missing),
                "lint": study_style_violations(md, build_claims(case, kn),
                                               case["product"].get("exporter_type"))}


def _add_runs(par, text: str) -> None:
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        par.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def markdown_to_docx(md: str, path: str, lang: str = "ar") -> str:
    """Markdown الدراسة → Word: العناوين والتسميات العريضة والجداول الخمسة حرفياً."""
    try:
        from docx import Document
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("python-docx غير مثبتة — pip install python-docx") from e
    import silk_reports as R
    doc = Document()
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip() or ln.strip() == "---":
            i += 1
            continue
        if ln.startswith("# "):
            doc.add_heading(ln[2:].strip(), level=0)
        elif ln.startswith("## "):
            doc.add_heading(ln[3:].strip(), level=1)
        elif ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in block
                     if not re.match(r"^\|?\s*-{3,}", row.replace("|", "|").strip())]
            cells = [c for c in cells if not all(re.fullmatch(r"-+", x or "-") for x in c)]
            if cells:
                R._add_table(doc, cells[0], cells[1:])
            continue
        else:
            p = doc.add_paragraph()
            _add_runs(p, ln.strip())
        i += 1
    R._apply_rtl(doc)
    # الفقرات **والخلايا**: عمود «المصدر» في جدول الرف كان خارج الفحص.
    blob = "\n".join([p.text for p in doc.paragraphs]
                     + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    hits = R._client_forbidden_hits(blob, lang)
    if hits:
        raise R.ReportGateError("تصدير الدراسة يحوي مصطلحاً ممنوعاً: " + "؛ ".join(hits[:5]))
    doc.save(path)
    return path


def study_docx(found: dict, path: str) -> tuple[str, dict]:
    md, meta = study_markdown(found)
    return markdown_to_docx(md, path), meta
